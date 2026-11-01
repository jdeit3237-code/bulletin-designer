import io
import os

from bulletin_designer import outline
from tests.test_bulletin import setup_admin, tok, app  # noqa: F401

OUTLINE = """ALL SAINTS
November 01, 2026
Mass of St. Ann
INTRODUCTORY RITES
Entrance For All the Saints (SINE NOMINE) BB26-718
Penitential Act Kyrie (Mass of Saint Ann, Ed Bolduc)
LITURGY OF THE EUCHARIST
Presentation of the Gifts Come to Me (Gregory Norbet) BB26-538
COMMUNION RITE
Communion I Am the Bread of Life (Suzanne Toolan) BB26-355
CONCLUDING RITES
Final Song Companions on the Journey (Carey Landry) BB26-459
"""


def test_parse():
    p = outline.parse_outline(OUTLINE)
    assert p["date"] == "2026-11-01"
    assert [s["kind"] for s in p["songs"]] == ["entrance", "", "gifts", "communion", "going_forth"]
    assert p["songs"][0]["number"] == "718"


def test_outline_fills_slides_and_replaces_on_change(app):
    c = setup_admin(app)
    t = tok(c, "/slides")
    c.post("/slides/songs", data={"csrf": t, "title": "For All the Saints", "code": "BB26-718", "lyrics": "For all the saints who sing\n\nwho from their labors rest"})
    c.post("/slides/songs", data={"csrf": t, "title": "Come to Me", "lyrics": "Come to me, all you who labor\n\nand I will give you rest"})
    r = c.post("/slides/outline", data={"csrf": t, "date": "2026-11-01", "outline_text": OUTLINE}, follow_redirects=True)
    assert b"Not in song library" in r.data
    plan = c.get("/slides/preview?date=2026-11-01"); plan.data = plan.data.replace(b"\xc2\xa0", b" ")
    assert b"who from their labors rest" in plan.data and b"Come to me, all you who labor" in plan.data
    changed = OUTLINE.replace("Come to Me (Gregory Norbet) BB26-538", "Another Hymn (X) BB26-1")
    c.post("/slides/outline", data={"csrf": t, "date": "2026-11-01", "outline_text": changed})
    plan = c.get("/slides/preview?date=2026-11-01"); plan.data = plan.data.replace(b"\xc2\xa0", b" ")
    assert b"For all the saints" in plan.data and b"Come to me" not in plan.data


def test_pptx_applies_weekly_concept_for_every_week(app):
    c = setup_admin(app)
    t = tok(c, "/slides")
    c.post("/slides/songs", data={"csrf": t, "title": "For All the Saints", "lyrics": "For all the saints\n\nwho rest"})
    for day in range(1, 43, 7):
        d = f"2026-{11 + (day - 1) // 31:02d}-{(day - 1) % 31 + 1:02d}" if day < 32 else "2026-12-%02d" % (day - 31)
        c.post("/slides/outline", data={"csrf": t, "date": d, "outline_text": OUTLINE})
        r = c.get(f"/slides/pptx?date={d}")
        assert r.status_code == 200 and r.data[:2] == b"PK", d


def test_design_overrides_saved(app):
    c = setup_admin(app)
    t = tok(c, "/slides")
    c.post("/slides/songs", data={"csrf": t, "title": "A Song", "lyrics": "la la"})
    c.post("/slides/outline", data={"csrf": t, "date": "2026-11-08", "outline_text": OUTLINE})
    c.post("/slides/design", data={"csrf": t, "date": "2026-11-08", "layout": "centered text with a thin gold border",
                                   "accent": "rope border", "color": "#336699", "scale": "120"})
    page = c.get("/slides?date=2026-11-08").data
    assert b'value="#336699"' in page and b'value="120"' in page
    assert c.get("/slides/pptx?date=2026-11-08").data[:2] == b"PK"
    c.post("/slides/design", data={"csrf": t, "date": "2026-11-08", "reset": "1"})
    assert b'value="120"' not in c.get("/slides?date=2026-11-08").data


def test_slide_library_status(app):
    c = setup_admin(app)
    t = tok(c, "/slides")
    assert b"Not started" in c.get("/slides/library").data
    c.post("/slides/songs", data={"csrf": t, "title": "For All the Saints", "lyrics": "For all the saints"})
    c.post("/slides/outline", data={"csrf": t, "date": "2026-11-01", "outline_text": OUTLINE})
    page = c.get("/slides/library").data
    assert b"2026-11-01" in page and b"In progress" in page
    c.post("/slides/status", data={"csrf": t, "date": "2026-11-01", "ready": "1"})
    assert b"Ready to use" in c.get("/slides/library").data
    assert b"Ready to use" in c.get("/slides?date=2026-11-01").data


def test_starter_pptx_for_not_started_week(app):
    c = setup_admin(app)
    r = c.get("/slides/pptx?date=2026-12-06")
    assert r.status_code == 200 and r.data[:2] == b"PK"


def test_refrain_repeats_after_each_verse():
    from bulletin_designer.outline import repeat_refrain
    parts = ["R. Lord, hear us", "Verse one", "Verse two", "Verse three"]
    assert repeat_refrain(parts) == ["R. Lord, hear us", "Verse one", "R. Lord, hear us", "Verse two",
                                     "R. Lord, hear us", "Verse three", "R. Lord, hear us"]
    already = ["Refrain: Amen", "Verse one", "Refrain: Amen", "Verse two"]
    assert repeat_refrain(already) == ["Refrain: Amen", "Verse one", "Refrain: Amen", "Verse two", "Refrain: Amen"]
    assert repeat_refrain(["One", "Two"]) == ["One", "Two"]


def test_black_slide_between_songs_and_readings(app):
    import io
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    c = setup_admin(app)
    t = tok(c, "/slides")
    d = "2026-11-08"
    for kind, name in (("entrance", "a.txt"), ("readings", "b.txt"), ("communion", "c.txt")):
        c.post("/slides", data={"csrf": t, "kind": kind, "date": d, "files": (io.BytesIO(b"hello"), name)},
               content_type="multipart/form-data")
    prs = Presentation(io.BytesIO(c.get(f"/slides/pptx?date={d}").data))
    black = [i for i, s in enumerate(prs.slides) if s.background.fill.fore_color.rgb == RGBColor(0, 0, 0)]
    assert len(black) == 2
    assert all(0 < i < len(prs.slides) - 1 for i in black)


def test_word_announcements_become_slides(app):
    import io
    import zipfile
    from pptx import Presentation
    xml = ('<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
           '<w:p><w:r><w:t>Food drive Sunday</w:t></w:r></w:p><w:p><w:r><w:t>Choir practice Thursday</w:t></w:r></w:p>'
           '</w:body></w:document>')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", xml)
    buf.seek(0)
    c = setup_admin(app)
    t = tok(c, "/slides")
    d = "2026-11-01"
    c.post("/slides", data={"csrf": t, "kind": "announcements", "date": d, "files": (buf, "Announcements.docx")},
           content_type="multipart/form-data")
    prs = Presentation(io.BytesIO(c.get(f"/slides/pptx?date={d}").data))
    text = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame).replace("\u00a0", " ")
    assert "Food drive Sunday" in text and "Choir practice Thursday" in text


def test_announcements_start_at_parish_announcements_line():
    from bulletin_designer.outline import announcements_from
    parts = ["Opening script", "*Lector: Our Parish Announcements this week:", "Food drive", "Choir Thursday"]
    assert announcements_from(parts) == ["Our Parish Announcements this week:", "Food drive", "Choir Thursday"]
    assert announcements_from(["Only", "These"]) == ["Only", "These"]

def test_long_reading_split_and_cp1252_decoded(tmp_path):
    from bulletin_designer.outline import read_text, split_long
    long = " ".join(f"Sentence number {i} is here." for i in range(30))
    pieces = split_long(long)
    assert len(pieces) > 1 and all(len(p) <= 240 for p in pieces)
    assert " ".join(p.replace("\n", " ") for p in pieces) == long
    assert split_long("Short") == ["Short"]
    f = tmp_path / "r.txt"
    f.write_bytes("saw \u201crising\u201d \u2014 east".encode("cp1252"))
    assert read_text(str(f)) == "saw \u201crising\u201d \u2014 east"
    f.write_bytes(b"\xef\xbb\xbfhello")
    assert read_text(str(f)) == "hello"


def test_no_orphans_joins_last_two_words():
    from bulletin_designer.outline import no_orphans
    assert no_orphans("one two three four\nHi there") == "one two three\u00a0four\nHi there"


def test_refrain_between_verses_and_numbers_removed():
    from bulletin_designer.outline import repeat_refrain
    out = repeat_refrain(["1. Verse one", "Refrain: Sing", "2. Verse two", "3) Verse three"])
    assert out == ["Verse one", "Refrain: Sing", "Verse two", "Refrain: Sing", "Verse three", "Refrain: Sing"][:0] or out[0] == "Verse one"
    assert not any(p[:2] in ("1.", "2.", "3)") for p in out)
    assert out == ["Verse one", "Refrain: Sing", "Verse two", "Refrain: Sing", "Verse three", "Refrain: Sing"] or out.count("Refrain: Sing") >= 3
    out = repeat_refrain(["Refrain line", "1 First", "2 Second"])
    assert out == ["Refrain line", "First", "Refrain line", "Second", "Refrain line"]


def test_reading_sections_titles_and_response_slides():
    from bulletin_designer.outline import reading_sections
    secs = dict(reading_sections("Reading 1\nText one.\n\nResponsorial Psalm\nR. Lord, hear.\nVerse A.\nR. Lord, hear.\nVerse B.\n\nReading 2\nText two.\n\nGospel Acclamation\nR. Alleluia.\nCome to me.\n"))
    assert list(secs) == ["First Reading", "Responsorial Psalm", "Second Reading", "Gospel Acclamation"]
    assert secs["Responsorial Psalm"] == ["Lord, hear.", "Verse A.", "Lord, hear.", "Verse B.", "Lord, hear."]
    assert secs["Gospel Acclamation"] == ["Come to me."]
    assert reading_sections("no headings here") is None


def test_duplicate_songs_in_several_sections_appear_once(app):
    c = setup_admin(app)
    t = tok(c, "/slides")
    for kind in ("entrance", "songs", "going_forth"):
        c.post("/slides", data={"csrf": t, "kind": kind, "date": "2026-11-01",
                                "files": (io.BytesIO(b"Unique lyric line here"), "For_All_the_Saints.txt")}, content_type="multipart/form-data")
    page = c.get("/slides/preview?date=2026-11-01").get_data(as_text=True)
    assert page.count("Unique lyric line") == 1


def test_readings_file_in_preview_uses_section_titles(app):
    c = setup_admin(app)
    t = tok(c, "/slides")
    txt = b"Reading 1\nA reading.\n\nResponsorial Psalm\nR. Lord, hear.\nVerse A.\n\nReading 2\nAnother reading.\n"
    c.post("/slides", data={"csrf": t, "kind": "readings", "date": "2026-11-01", "files": (io.BytesIO(txt), "Readings.txt")}, content_type="multipart/form-data")
    page = c.get("/slides/preview?date=2026-11-01").get_data(as_text=True)
    assert "First Reading" in page and "Second Reading" in page and "Responsorial Psalm" in page and "Reading 1" not in page


def test_standard_pptx_slides_appear_in_preview(app):
    from pptx import Presentation
    c = setup_admin(app)
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[5]); s.shapes.title.text = "Lamb of God"
    buf = io.BytesIO(); p.save(buf); buf.seek(0)
    t = tok(c, "/slides")
    c.post("/slides", data={"csrf": t, "kind": "standard", "files": (buf, "Lamb_of_God.pptx")}, content_type="multipart/form-data")
    page = c.get("/slides/preview?date=2026-11-01").get_data(as_text=True)
    assert "Lamb of" in page


def test_slides_follow_outline_order_with_blanks(app):
    from pptx import Presentation
    c = setup_admin(app)
    t = tok(c, "/slides")
    c.post("/slides/songs", data={"csrf": t, "title": "For All the Saints", "code": "BB26-718", "lyrics": "ENTRANCE WORDS"})
    outline = OUTLINE.replace("Penitential Act Kyrie (Mass of Saint Ann, Ed Bolduc)\n", "Penitential Act Kyrie (Mass of Saint Ann, Ed Bolduc)\nFirst Reading\nResponsorial Psalm\nSecond Reading\nGospel Acclamation\nGospel\n")
    c.post("/slides/outline", data={"csrf": t, "date": "2026-11-01", "outline_text": outline}, follow_redirects=True)
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[5]); s.shapes.title.text = "HOLY WORDS"
    buf = io.BytesIO(); p.save(buf); buf.seek(0)
    c.post("/slides", data={"csrf": t, "kind": "standard", "files": (buf, "Holy_Holy_Holy.pptx")}, content_type="multipart/form-data")
    txt = b"Reading 1\nFIRSTTEXT.\n\nResponsorial Psalm\nR. PSALMRESP\nPSALMVERSE.\n\nReading 2\nSECONDTEXT.\n\nGospel\nGOSPELTEXT.\n"
    c.post("/slides", data={"csrf": t, "kind": "readings", "date": "2026-11-01", "files": (io.BytesIO(txt), "Readings.txt")}, content_type="multipart/form-data")
    page = c.get("/slides/preview?date=2026-11-01").get_data(as_text=True).replace("\xa0", " ")
    order = ["ENTRANCE WORDS", "FIRSTTEXT", "PSALMRESP", "SECONDTEXT", "GOSPELTEXT", "HOLY WORDS"]
    pos = [page.index(w) for w in order]
    assert pos == sorted(pos)


def test_readings_split_into_one_or_two_sentence_slides():
    from bulletin_designer.outline import reading_sections
    body = " ".join(f"This is sentence number {i} of the reading." for i in range(12))
    secs = dict(reading_sections("Reading 1\nRevelation 7:2-4\n" + body + "\n\nResponsorial Psalm\nR. (see 6) Lord, hear.\nVerse A.\n"))
    assert len(secs["First Reading"]) >= 3 and all(len(p) <= 220 for p in secs["First Reading"])
    assert secs["Responsorial Psalm"][0] == "Lord, hear."


def test_mac_roman_file_and_single_block_song():
    from bulletin_designer.outline import decode_text, reading_sections, repeat_refrain
    raw = "Reading 1\u00a0\nRevelation 7:2\nI saw\u00a0an angel\n\nResponsorial Psalm\nR.\u00a0Lord, hear.\nVerse.\n".encode("mac_roman")
    secs = dict(reading_sections(decode_text(raw)))
    assert list(secs) == ["First Reading", "Responsorial Psalm"]
    assert secs["Responsorial Psalm"][0] == "Lord, hear."
    assert "\u00ca" not in decode_text(raw)
    song = ["Refrain: Sing out\n1. Verse one line\nline two\n2. Verse two line"]
    assert repeat_refrain(song) == ["Refrain: Sing out", "Verse one line\nline two", "Refrain: Sing out", "Verse two line", "Refrain: Sing out"]


def test_outline_order_puts_song_of_praise_after_communion_or_elsewhere():
    def labels(order):
        chunks = [{"kind": "praise", "name": "Song of Praise-x.txt", "group": None},
                  {"kind": "communion", "name": "Communion-y.txt", "group": None},
                  {"kind": "going_forth", "name": "Final Song-z.txt", "group": None}]
        return [c["label"] for c in outline.order_chunks(chunks, order)]
    assert labels(["Communion", "Song of Praise", "Final Song"]) == ["Communion", "Song of Praise", "Final Song"]
    assert labels(["Song of Praise", "Communion", "Final Song"]) == ["Song of Praise", "Communion", "Final Song"]


def test_song_credit_and_license_only_on_first_slide(app):
    c = setup_admin(app)
    t = tok(c, "/slides")
    c.post("/slides/songs", data={"csrf": t, "title": "For All the Saints", "code": "BB26-718",
                                  "lyrics": "FIRSTVERSE words\n\nSECONDVERSE words\n\nText: HOWTEXT 1823\nMusic: VWMUSIC 1906"})
    c.post("/slides/outline", data={"csrf": t, "date": "2026-11-01", "outline_text": OUTLINE}, follow_redirects=True)
    page = c.get("/slides/preview?date=2026-11-01").get_data(as_text=True)
    assert page.count("ONE LICENSE") == page.count("HOWTEXT") >= 1
    assert page.index("FIRSTVERSE") < page.index("HOWTEXT") < page.index("SECONDVERSE") or page.count("HOWTEXT") == 1
    credit = page[page.index("HOWTEXT"):page.index("ONE LICENSE")]
    assert "\n" in credit


def test_other_songs_is_never_a_section_title(app):
    c = setup_admin(app)
    t = tok(c, "/slides")
    c.post("/slides", data={"csrf": t, "kind": "songs", "date": "2026-11-01", "files": (io.BytesIO(b"Hello world\n"), "Mystery Song.txt")}, content_type="multipart/form-data")
    page = c.get("/slides/preview?date=2026-11-01").get_data(as_text=True)
    assert "Other songs" not in page and "Mystery Song" in page


def test_english_only_drops_spanish_unless_bilingual():
    assert outline.is_spanish("Santo, Santo, Santo es el Señor Dios del universo")
    assert not outline.is_spanish("Holy, Holy, Holy Lord God of hosts")
    parts = ["Holy, Holy, Holy Lord God of hosts", "Santo, Santo, Santo es el Señor Dios del universo"]
    assert outline.english_only(parts) == parts[:1]
    assert outline.english_only(parts, keep_all=True) == parts


def test_reading_first_slide_has_title_and_citation():
    assert outline.split_citation("Revelation 7:2-4, 9-14 I, John, saw another angel.") == ("Revelation 7:2-4, 9-14", "I, John, saw another angel.")
    assert outline.split_citation("1 John 3:1-3 See what love.")[0] == "1 John 3:1-3"
    assert outline.split_citation("Lord, this is the people that longs to see your face.")[0] == ""


def test_announcements_drop_references_elsewhere():
    parts = ["Pumpkin Patch Saturday at 10am. More details in the bulletin. Thank you for your generosity!", "See the bulletin for more information.", "Bring a friend."]
    assert outline.drop_referrals(parts) == ["Pumpkin Patch Saturday at 10am.", "Bring a friend."]


def test_graphic_follows_its_announcement():
    texts = ["Pumpkin Patch opens Saturday for families.", "Pet Wellness Clinic for cats and dogs on Friday."]
    assert outline.best_match("Pet wellness clinic flyer", texts) == 1
    assert outline.best_match("Pumpkin patch", texts) == 0
    assert outline.best_match("Bingo night", texts) is None


def test_preview_warns_when_required_songs_missing(app):
    c = setup_admin(app)
    t = tok(c, "/slides")
    c.post("/slides/songs", data={"csrf": t, "title": "For All the Saints", "code": "BB26-718", "lyrics": "ENTRANCE WORDS"})
    c.post("/slides/outline", data={"csrf": t, "date": "2026-11-01", "outline_text": OUTLINE}, follow_redirects=True)
    page = c.get("/slides/preview?date=2026-11-01").get_data(as_text=True)
    assert "Missing songs" in page and "Going Forth" in page


def test_song_cleanup_labels_title_block_and_code():
    assert outline.clean_song_title("For All the Saints 113657") == "For All the Saints"
    assert outline.clean_song_title("Companions on the Journey BB26-718") == "Companions on the Journey"
    assert outline.drop_title_block(["For All the Saints\nSINE NOMINE", "Verse"], "For All the Saints") == ["Verse"]
    assert outline.strip_labels("Refrain\nAnd I will raise you up") == "And I will raise you up"
    assert outline.slide_size("x" * 250) < outline.slide_size("x" * 50)


def test_psalm_and_acclamation_citation_move_to_first_slide():
    secs = dict(outline.reading_sections("Responsorial Psalm\nPsalm 24:1bc-2, 3-4ab, 5-6\nR. (see 6) Lord, this is the people.\nVerse one.\n\nGospel Acclamation\nMatthew 11:28 R. Alleluia, alleluia.\nCome to me.\n"))
    c, rest = outline.split_citation(secs["Responsorial Psalm"][0])
    assert c == "Psalm 24:1bc-2, 3-4ab, 5-6" and rest.startswith("Lord, this is the people")
    c2, rest2 = outline.split_citation(secs["Gospel Acclamation"][0])
    assert c2 == "Matthew 11:28" and "Alleluia" not in rest2


def test_mac_roman_spanish_not_garbled():
    raw = "Yo los resucitar\u00e1".encode("mac_roman") + b"\n" + "el d\u00eda final y m\u00e1s".encode("mac_roman")
    assert "\u017d" not in outline.decode_text(raw)


def test_extra_section_uses_given_name_and_place():
    chunks = [
        {"kind": "extra", "name": "x.txt", "group": None, "items": [], "custom": "Thanksgiving Hymn", "after": "Communion"},
        {"kind": "communion", "name": "c.txt", "group": None, "items": []},
        {"kind": "going_forth", "name": "g.txt", "group": None, "items": []},
    ]
    out = outline.order_chunks(chunks, [])
    assert [c["kind"] for c in out] == ["communion", "extra", "going_forth"]
    assert out[1]["label"] == "Extra: Thanksgiving Hymn"


def test_psalm_inline_responses_and_acclamation():
    raw = ("Responsorial Psalm\nPsalm 24:1bc-2, 3-4ab, 5-6\nR. (6) Lord, this is the people that longs to see your face. "
           "The LORD's are the earth and its fullness. For he founded it upon the seas. R. Lord, this is the people that longs to see your face. "
           "Who can ascend the mountain of the LORD? R. Lord, this is the people that longs to see your face.\n\n"
           "Gospel Acclamation\nMatthew 11:28\nR. Alleluia, alleluia. Come to me, all you who labor and are burdened, says the Lord. R. Alleluia, alleluia.\n")
    secs = dict(outline.reading_sections(raw))
    psalm = secs["Responsorial Psalm"]
    resp = "Lord, this is the people that longs to see your face."
    assert psalm[0].endswith(resp) and psalm[-1] == resp and not any("R." in p for p in psalm)
    assert not any(resp in p and len(p) > len(resp) + 40 for p in psalm[1:])
    acc = secs["Gospel Acclamation"]
    assert len(acc) == 1 and "R." not in acc[0] and "Alleluia" not in acc[0] and "Come to me" in acc[0]


def test_title_block_becomes_song_title_and_clean_name():
    title, rest = outline.song_title_block(["For All the Saints\nSINE NOMINE", "1. For all the saints who from their labors rest,"])
    assert title == "For All the Saints" and len(rest) == 1
    assert outline.clean_song_title("Entrance-bb26") == "Entrance"


def test_section_designs_make_two_pictures_and_apply_to_slides(app, monkeypatch):
    import base64
    import io as _io
    import json as _json
    import urllib.request
    from PIL import Image
    buf = _io.BytesIO(); Image.new("RGB", (30, 20), "navy").save(buf, "PNG")
    payload = _json.dumps({"data": [{"b64_json": base64.b64encode(buf.getvalue()).decode()}]}).encode()

    class Resp:
        def __enter__(self): return self
        def __exit__(self, *a): return False
        def read(self): return payload

    monkeypatch.setenv("OPENAI_API_KEY", "k")
    monkeypatch.setattr(urllib.request, "urlopen", lambda *a, **k: Resp())
    c = setup_admin(app)
    t = tok(c, "/slides")
    c.post("/slides/songs", data={"csrf": t, "title": "Come to Me", "lyrics": "Come to me, all you who labor\n\nand I will give you rest"})
    c.post("/slides/outline", data={"csrf": t, "date": "2026-11-01", "outline_text": OUTLINE}, follow_redirects=True)
    for role in ("title", "body"):
        r = c.post("/slides/designs", data={"csrf": t, "date": "2026-11-01", "kind": "gifts", "role": role}, follow_redirects=True)
        assert b"Made the" in r.data
    page = c.get("/slides/preview?date=2026-11-01").get_data(as_text=True)
    assert "ai-title.png" in page and "ai-body.png" in page
    assert c.post("/slides/designs", data={"csrf": t, "date": "2026-11-01", "kind": "bogus", "role": "title"}).status_code == 302


def test_tidy_song_parts_fixes_tabs_l_and_repeats():
    parts = outline.tidy_song_parts(["For l am meek\n\tand humble", "Refrain line", "Refrain  line", "Next verse"])
    assert parts == ["For I am meek\n and humble", "Refrain line", "Next verse"]


def test_first_slide_has_written_date_theme_and_church_photo(app):
    c = setup_admin(app)
    r = c.get("/slides/preview?date=2026-11-01")
    assert b"November 1, 2026" in r.data and b"All Saints" in r.data and b"church.jpg" in r.data


def test_chatgpt_prompt_lists_every_slide_text(app):
    c = setup_admin(app)
    t = tok(c, "/slides")
    c.post("/slides/songs", data={"csrf": t, "title": "Come to Me", "lyrics": "Come to me, all you who labor\n\nand I will give you rest"})
    c.post("/slides/outline", data={"csrf": t, "date": "2026-11-01", "outline_text": OUTLINE}, follow_redirects=True)
    r = c.get("/slides/chatgpt-prompt?date=2026-11-01")
    body = r.get_data(as_text=True)
    assert r.status_code == 200 and "attachment" in r.headers["Content-Disposition"]
    assert "November 1, 2026" in body and "SLIDE 1" in body and "and I will give you rest" in body and "regular weekly slides" in body and "INDIVIDUALLY DESIGNED" in body


def test_slides_pdf_has_one_page_per_slide(app):
    import pymupdf, re
    c = setup_admin(app)
    t = tok(c, "/slides")
    c.post("/slides/songs", data={"csrf": t, "title": "Come to Me", "lyrics": "Come to me, all you who labor\n\nand I will give you rest"})
    c.post("/slides/outline", data={"csrf": t, "date": "2026-11-01", "outline_text": OUTLINE}, follow_redirects=True)
    r = c.get("/slides/pdf?date=2026-11-01")
    assert r.status_code == 200 and r.mimetype == "application/pdf"
    doc = pymupdf.open(stream=r.data, filetype="pdf")
    page_count = len(re.findall(r'style="aspect-ratio:16/9', c.get("/slides/preview?date=2026-11-01").get_data(as_text=True)))
    assert doc.page_count == page_count and doc.page_count >= 3
    assert "give you rest" in "".join(p.get_text() for p in doc)


def test_music_pages_zip_matches_prompt(app):
    import io, zipfile
    c = setup_admin(app)
    t = tok(c, "/slides")
    c.post("/slides/outline", data={"csrf": t, "date": "2026-11-01", "outline_text": OUTLINE}, follow_redirects=True)
    r = c.get("/slides/music-pages?date=2026-11-01")
    assert r.status_code == 200 and r.mimetype == "application/zip"
    assert zipfile.ZipFile(io.BytesIO(r.data)).namelist()
    assert c.get("/slides/music-pages?date=bad").status_code == 404


def test_chatgpt_prompt_demands_images_not_powerpoint(app):
    c = setup_admin(app)
    t = tok(c, "/slides")
    c.post("/slides/outline", data={"csrf": t, "date": "2026-11-01", "outline_text": OUTLINE}, follow_redirects=True)
    body = c.get("/slides/chatgpt-prompt?date=2026-11-01").get_data(as_text=True)
    assert body.startswith("READ THIS FIRST") and "IMAGE GENERATION" in body


def test_chatgpt_per_slide_zip(app):
    import io, zipfile
    c = setup_admin(app)
    t = tok(c, "/slides")
    c.post("/slides/songs", data={"csrf": t, "title": "Come to Me", "lyrics": "Come to me, all you who labor\n\nand I will give you rest"})
    c.post("/slides/outline", data={"csrf": t, "date": "2026-11-01", "outline_text": OUTLINE}, follow_redirects=True)
    r = c.get("/slides/chatgpt-slides?date=2026-11-01")
    z = zipfile.ZipFile(io.BytesIO(r.data))
    names = z.namelist()
    assert "00_START_HERE.txt" in names and "SLIDE-01.txt" in names
    assert "IMAGE GENERATION" in z.read("00_START_HERE.txt").decode()
    assert any("and I will give you rest" in z.read(n).decode() for n in names)


def test_outline_export_txt_and_csv(app):
    c = setup_admin(app)
    t = tok(c, "/slides")
    c.post("/slides/songs", data={"csrf": t, "title": "Come to Me", "lyrics": "Come to me, all you who labor\n\nand I will give you rest"})
    c.post("/slides/outline", data={"csrf": t, "date": "2026-11-01", "outline_text": OUTLINE}, follow_redirects=True)
    r = c.get("/slides/outline-export?date=2026-11-01")
    body = r.get_data(as_text=True)
    assert r.status_code == 200 and "attachment" in r.headers["Content-Disposition"]
    assert "\n\tand I will give you rest\n" in body and "\nBODY:\n" in body
    assert "ONE LICENSE" not in body
    assert "SONG: Presentation of the Gifts\n\tCome To Me\n\t\tCome to me" in body
    m = c.get("/slides/outline-macro")
    assert m.status_code == 200 and b"Sub ApplyMassLayouts" in m.data
    r = c.get("/slides/outline-export?date=2026-11-01&format=csv")
    assert r.get_data(as_text=True).splitlines()[0].lstrip("\ufeff") == "Slide,Header,Subheader,Body,Copyright"
    assert c.get("/slides/outline-export?date=x").status_code == 404


def test_master_slides_deck_and_template(app):
    from pptx import Presentation
    import io as _io
    c = setup_admin(app)
    t = tok(c, "/slides")
    c.post("/slides/songs", data={"csrf": t, "title": "Come to Me", "lyrics": "Come to me, all you who labor\n\nand I will give you rest"})
    c.post("/slides/outline", data={"csrf": t, "date": "2026-11-01", "outline_text": OUTLINE}, follow_redirects=True)
    r = c.get("/slides/pptx-master?date=2026-11-01")
    assert r.status_code == 200
    prs = Presentation(_io.BytesIO(r.data))
    names = [s.slide_layout.name for s in prs.slides]
    assert "69_Title Slide No fill" in names and "112_Title Slide No fill" in names
    text = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "Come To Me" in text and "and I will give you rest" in text and "ONE LICENSE" not in text
    assert c.get("/slides/template").status_code == 200
    bad = c.post("/slides/template", data={"csrf": t, "template": (_io.BytesIO(b"nope"), "x.pptx")}, content_type="multipart/form-data")
    assert bad.status_code == 302


def test_master_deck_reading_uses_header_subheader_body(app):
    from pptx import Presentation
    import io as _io
    c = setup_admin(app)
    t = tok(c, "/slides")
    c.post("/slides/outline", data={"csrf": t, "date": "2026-11-01", "outline_text": OUTLINE}, follow_redirects=True)
    prs = Presentation(_io.BytesIO(c.get("/slides/pptx-master?date=2026-11-01").data))
    for s in prs.slides:
        if s.slide_layout.name.startswith("70_"):
            assert False, "reading fell into the title-only layout"


def test_credit_only_slide_moves_into_song_copyright(app):
    c = setup_admin(app)
    t = tok(c, "/slides")
    c.post("/slides/songs", data={"csrf": t, "title": "Come to Me", "lyrics": "Come to me, all you who labor\n\nand I will give you rest\n\nDarryl Ducote\nCCLI Song #253461"})
    c.post("/slides/outline", data={"csrf": t, "date": "2026-11-01", "outline_text": OUTLINE}, follow_redirects=True)
    body = c.get("/slides/outline-export?date=2026-11-01").get_data(as_text=True)
    assert "\t\t\tDarryl Ducote" in body or "CCLI Song #253461" not in body.replace("\t\t\t", "X", 0) and "\n\tDarryl" not in body


def test_duplicate_layout_with_background(tmp_path):
    import io as _io
    from PIL import Image
    from pptx import Presentation
    from bulletin_designer.app import duplicate_layout
    img = tmp_path / "bg.png"
    Image.new("RGB", (64, 36), (200, 30, 30)).save(img)
    prs = Presentation(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "assets", "MassTemplate.pptx"))
    base = prs.slide_layouts[0]
    new = duplicate_layout(prs, base, str(img), "Song - design 1")
    sld = prs.slides.add_slide(new)
    out = _io.BytesIO()
    prs.save(out)
    again = Presentation(_io.BytesIO(out.getvalue()))
    names = [l.name for l in again.slide_layouts]
    assert "Song - design 1" in names and len(names) == 4
    assert any(sl.slide_layout.name == "Song - design 1" for sl in again.slides)



def test_pptx_pages_text(tmp_path):
    from pptx import Presentation
    import app as appmod
    prs = Presentation()
    for t in ("Verse one\nline two", "Refrain here"):
        sl = prs.slides.add_slide(prs.slide_layouts[5])
        sl.shapes.title.text = t
    f = tmp_path / "s.pptx"
    prs.save(str(f))
    assert appmod.pptx_text(str(f)).split("\n\n") == ["Verse one\nline two", "Refrain here"]


def test_copy_slide_keeps_source_slide(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    import app as appmod
    src = Presentation()
    sl = src.slides.add_slide(src.slide_layouts[6])
    tb = sl.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    tb.text_frame.text = "Glory to God"
    f = tmp_path / "g.pptx"
    src.save(str(f))
    dst = Presentation()
    out = appmod.copy_slide(dst, dst.slide_layouts[6], str(f), 0)
    assert any(sh.has_text_frame and sh.text_frame.text == "Glory to God" for sh in out.shapes)

"""St. Peter the Fisherman weekly bulletin designer (standalone Flask app)."""
import base64
import hashlib
import io
import time
import json
import os
import itertools
import re
from urllib.parse import urlsplit
import secrets
import sqlite3
import shutil
import smtplib
import subprocess
import tempfile
import urllib.request
import zipfile
from xml.etree import ElementTree
import uuid
from datetime import date, datetime, timedelta
from email.message import EmailMessage
from functools import wraps

from flask import (Flask, abort, flash, g, jsonify, redirect, render_template, request,
                   send_file, send_from_directory, session, url_for)
from werkzeug.security import check_password_hash
from werkzeug.security import generate_password_hash as _gen


def generate_password_hash(pw):
    return _gen(pw, method="pbkdf2:sha256:600000")
from werkzeug.utils import secure_filename

try:
    from . import liturgy, help_content, outline as music_outline
except ImportError:  # run as a script
    import liturgy
    import help_content
    import outline as music_outline

try:
    from . import routes_choir
except ImportError:
    try:
        import routes_choir
    except ImportError:  # choir needs reportlab; keep the rest of the site working
        routes_choir = None

BULLETS = {"disc": "\u2022 Round dot", "circle": "\u25e6 Open circle", "square": "\u25aa Square", "dash": "\u2013 Dash", "check": "\u2713 Check mark", "arrow": "\u27a4 Arrow", "cross": "\u271d Cross", "star": "\u2605 Star"}

ROLES = ("admin", "editor", "contributor", "webmaster", "viewer")
ROLE_HELP = {
    "admin": "Manage users, sections and layout; everything an editor can do.",
    "editor": "Edit every section, approve contributions, publish the bulletin.",
    "contributor": "Edit only assigned sections; changes wait for editor approval.",
    "webmaster": "Sees only the list of published bulletins and can download the web PDF of each (nothing else).",
    "viewer": "Read-only access to drafts and published bulletins.",
}
WEBMASTER_PAGES = {"dashboard", "download", "password", "logout", "static", "login", "forgot", "reset_with_token", "register"}
IMAGE_EXT = {"png", "jpg", "jpeg", "gif", "webp"}
PHOTO_SHAPES = {"3:2": (1.5, "Wide (3 by 2)"), "1:1": (1.0, "Square"), "2:3": (2 / 3, "Tall (2 by 3)"), "3:1": (3.0, "Banner (3 by 1)"), "16:9": (16 / 9, "Widescreen (16 by 9)")}
LIBRARY_EXT = IMAGE_EXT | {"tif", "tiff"}

# Sections found in the Oct 11/18/25 2026 bulletins. weekly=1 means new content is
# needed every week; weekly=0 means "standing" content carried forward automatically.
FONT_STYLES = {
    "condensed": "Condensed - Archivo Narrow (fits more text)",
    "condensed-serif": "Condensed serif - Georgia-style narrow (fits more text)",
    "serif": "Classic serif - Georgia",
    "sans": "Clean sans-serif - Arial",
    "humanist": "Humanist - Verdana (easy to read)",
    "rounded": "Friendly rounded - Trebuchet",
    "source": "Myriad-style sans - Source Sans 3",
    "title-georgia": "Heading only - Georgia bold (body unchanged)",
    "title-condserif": "Heading only - Condensed serif bold (body unchanged)",
    "title-source": "Heading only - Source Sans bold (body unchanged)",
    "title-narrow": "Heading only - Condensed sans bold (body unchanged)",
    "title-playfair": "Heading only - Playfair Display bold (body unchanged)",
    "title-lora": "Heading only - Lora bold (body unchanged)",
    "title-montserrat": "Heading only - Montserrat bold (body unchanged)",
    "title-oswald": "Heading only - Oswald bold (tall, condensed) (body unchanged)",
    "title-raleway": "Heading only - Raleway bold (body unchanged)",
    "title-pacifico": "Fun title - Pacifico script (body unchanged)",
    "title-kaushan": "Fun title - Kaushan brush script (body unchanged)",
    "title-lobster": "Fun title - Lobster bold script (body unchanged)",
    "title-vibes": "Fun title - Great Vibes elegant script (body unchanged)",
}

DEFAULT_SECTIONS = [
    ("cover", "Parish Information & Cover", 0, 1, "Sunday date and title, parish address, office hours"),
    ("mass-intentions", "Mass Intentions", 1, 1, "Weekly form: day, date, time, who the Mass is for, who requested it (usually page 2)"),
    ("staff", "Parish Staff", 0, 1, "Standing list; edit only when staff change (carried forward)"),
    ("mass-times", "Mass & Service Times", 0, 1, "Standing list: title, days, name of Mass, time, notes, web address (blank fields are left out)"),
    ("welcome", "Welcome Message", 0, 1, "Standing invitation and welcome"),
    ("youth-calendar", "Youth Group Calendar", 1, 1, "One event per line: Date | Time | Event (assign to the youth group leaders)"),
    ("calendar", "This Week in Our Parish (Calendar of Events)", 1, 1, "One event per line: Date | Time | Event (shown as a table)"),
    ("mass-schedule", "This Week in Our Parish: Lectors, Sacristans & Ministers", 1, 1, "Weekly form: header, then day, date, lectors, sacristans and Eucharistic ministers (several names allowed)"),
    ("readings", "Readings & Gospel", 1, 1, "usccb.org daily readings link or text"),
    ("pastor", "Pastor's Message", 1, 1, "Weekly reflection from the pastor"),
    ("sick", "Pray for Those Who Are Sick", 1, 1, "Prayer list; review weekly"),
    ("announcements", "Parish Announcements", 1, 1, "Events, bake sales, fundraisers"),
    ("food-pantry", "St. Peter Community Food Pantry", 0, 1, "Hours, tours, donation info"),
    ("ministries", "Ministries & Faith Formation", 0, 1, "Schedules, RCIA, youth group"),
    ("sacraments", "Sacraments & Office Information", 0, 1, "Baptism, marriage, reconciliation"),
    ("giving", "Stewardship & Giving", 1, 1, "Offertory totals and online giving"),
    ("sponsors", "Community Sponsors", 0, 1, "Business advertisements"),
]


AD_W, AD_H, AD_GAP_X, AD_GAP_Y = 7.9, 8.0, 0.25, 0.2


def clean_ad_url(raw):
    """Optional link for an ad card: only http(s) addresses are kept; a bare domain gets https://."""
    u = (raw or "").strip()[:500]
    if not u or re.search(r"\s", u):
        return ""
    if not re.match(r"^[a-z][a-z0-9+.-]*:", u, re.I):
        u = "https://" + u
    p = urlsplit(u)
    return u if p.scheme in ("http", "https") and p.netloc and "." in p.netloc else ""


def pack_ads(cards, scale, gap=AD_GAP_X, pad=(0.0, 0.0)):
    """Shelf-pack cards (long edge 3.5in x scale, true aspect). Horizontal cards first, then vertical. Returns pages of rows."""
    def size(c):
        ratio = min(c["w"], c["h"]) / max(c["w"], c["h"]) if c["w"] and c["h"] else 0.57
        long_edge = 3.5 * scale
        return (long_edge, long_edge * ratio) if c["w"] >= c["h"] else (long_edge * ratio, long_edge)
    ordered = [c for c in cards if c["w"] >= c["h"]] + [c for c in cards if c["w"] < c["h"]]
    rows, row, row_w = [], [], 0.0
    for c in ordered:
        w, h = size(c)
        item = {"name": c["name"], "image": c["image"], "url": c.get("url") or "", "w": round(w, 3), "h": round(h, 3)}
        if row and (row_w + gap + pad[0] + w > AD_W + pad[0] + 0.001 or (row[0]["w"] >= row[0]["h"]) != (c["w"] >= c["h"])):
            rows.append(row)
            row, row_w = [], 0.0
        row_w += (gap + pad[0] if row else 0) + w
        item["ph"] = round(pad[1], 3)
        row.append(item)
    if row:
        rows.append(row)
    pages, cur, used = [], [], 0.0
    for r in rows:
        h = max(i["h"] for i in r) + pad[1]
        if cur and used + gap + h > AD_H:
            pages.append(cur)
            cur, used = [], 0.0
        used += (gap if cur else 0) + h
        cur.append(r)
    if cur:
        pages.append(cur)
    return pages


def ad_pages(cards, want=0, gap=AD_GAP_Y, pad=(0.0, 0.0)):
    """Fit all cards on the last pages. One size scale is used for every card on every page; the page break is chosen where that scale can be largest.
    want = number of pages to spread them over (0 = automatic, 1 or 2)."""
    if not cards:
        return []

    def row_h(r):
        return max(i["h"] for i in r) + pad[1]

    def factor(rows):
        if not rows:
            return 99.0, True
        fw = min((AD_W + pad[0] - (len(r) - 1) * (gap + pad[0])) / sum(i["w"] for i in r) for r in rows)
        fh = (AD_H - gap * (len(rows) - 1)) / sum(row_h(r) for r in rows)
        return min(fw, fh, 2.5), fh >= 0.9999

    best = None
    for scale in [round(1.5 - 0.01 * k, 2) for k in range(0, 101)]:
        rows = [r for pg in pack_ads(cards, scale, gap, pad) for r in pg]
        parts = []
        for n in ([want] if want else [1, 2]):
            n = max(1, min(n, len(rows)))
            for cuts in itertools.combinations(range(1, len(rows)), n - 1):
                edges = (0,) + cuts + (len(rows),)
                parts.append([rows[x:y] for x, y in zip(edges, edges[1:])])
        if want != 1 and len(rows) == 1:
            parts.append([rows, []])
        for part in parts:
            fs = [factor(pg) for pg in part]
            if not all(ok for _, ok in fs):
                continue
            common = min(f for f, _ in fs)
            score = scale * common
            if best is None or score > best[0] + 1e-9:
                best = (score, common, part)
    if best:
        _, common, pages = best
    else:
        pages, common = pack_ads(cards, 0.5, gap, pad), 1.0
    k = common - 0.003 if common > 1.003 else 1.0
    out = []
    for rows in pages:
        rows = [[dict(i, w=round(i["w"] * k, 3), h=round(i["h"] * k, 3)) for i in r] for r in rows]
        used = sum(row_h(r) for r in rows)
        top = max(gap, (AD_H - used - gap * max(len(rows) - 1, 0)) / 2) if rows else 0
        out.append({"rows": rows, "gap": round(gap, 3), "top": round(top, 3)})
    return out


W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def docx_text(fileobj):
    """Extract paragraph text from a .docx (blank line between paragraphs, tabs kept as ' | ')."""
    with zipfile.ZipFile(fileobj) as z:
        info = z.getinfo("word/document.xml")
        if info.file_size > 20 * 1024 * 1024:
            raise ValueError("Document is too large.")
        root = ElementTree.fromstring(z.read("word/document.xml"))
    paras = []
    for p in root.iter(W_NS + "p"):
        parts = []
        for node in p.iter():
            if node.tag == W_NS + "t":
                parts.append(node.text or "")
            elif node.tag == W_NS + "tab":
                parts.append(" | ")
            elif node.tag in (W_NS + "br", W_NS + "cr"):
                parts.append("\n")
        paras.append("".join(parts).strip())
    return "\n\n".join(x for x in paras if x)


def doc_text(path):
    """Legacy .doc via antiword or macOS textutil when installed."""
    for cmd in (["antiword", path], ["textutil", "-convert", "txt", "-stdout", path]):
        if shutil.which(cmd[0]):
            out = subprocess.run(cmd, capture_output=True, timeout=30, check=True)
            return out.stdout.decode("utf-8", "replace").strip()
    raise ValueError("Old .doc files cannot be read on this server. Please save the file as .docx and upload again.")


BULLETIN_CARRY = ("footer", "hdr_image", "hdr_alt", "hdr_h", "hdr_zoom", "hdr_x", "hdr_y", "hdr_style",
                  "hdr2_image", "hdr2_alt", "hdr2_zoom", "hdr2_x", "hdr2_y", "hdr_split",
                  "toc_title", "toc_before", "toc_title_gap", "toc_item_gap", "toc_start", "toc_show", "toc_pt", "toc_style",
                  "toc_place", "toc_cols", "no_stragglers", "no_underline")


def gap_value(raw):
    """Blank = default (stored 0); a typed 0 = no space (stored as a tiny value)."""
    raw = (raw or "").strip()
    if not raw:
        return 0.0
    try:
        v = max(0.0, min(1.0, float(raw)))
    except ValueError:
        return 0.0
    return v if v > 0 else 0.0001


def gap_display(v):
    if not v:
        return ""
    return "0" if v < 0.001 else ("%g" % v)


def read_photo_meta(path):
    """Pull size, EXIF, IPTC and XMP details from an image. Never raises; GPS location is deliberately not kept."""
    out = {"keywords": [], "title": "", "description": "", "creator": "", "copyright": "", "taken": "", "camera": "", "size": ""}
    try:
        from PIL import Image
        im = Image.open(path)
        out["size"] = "%d x %d px" % im.size
        try:
            ex = im.getexif()
            exif = dict(ex)
            exif.update(dict(ex.get_ifd(0x8769)))
        except Exception:
            exif = {}
        out["taken"] = str(exif.get(36867) or exif.get(306) or "").strip()[:19].replace(":", "-", 2)
        out["camera"] = " ".join(str(exif.get(k, "")).strip() for k in (271, 272) if exif.get(k)).strip()
        out["description"] = str(exif.get(270) or "").strip()
        out["creator"] = str(exif.get(315) or "").strip()
        out["copyright"] = str(exif.get(33432) or "").strip()
        try:
            from PIL import IptcImagePlugin
            iptc = IptcImagePlugin.getiptcinfo(im) or {}

            def txt(v):
                v = v[0] if isinstance(v, list) else v
                return v.decode("utf-8", "ignore").strip() if isinstance(v, bytes) else str(v or "").strip()
            kw = iptc.get((2, 25))
            for k in (kw if isinstance(kw, list) else [kw] if kw else []):
                out["keywords"].append(txt(k))
            out["title"] = out["title"] or txt(iptc.get((2, 5), ""))
            out["description"] = out["description"] or txt(iptc.get((2, 120), ""))
            out["creator"] = out["creator"] or txt(iptc.get((2, 80), ""))
        except Exception:
            pass
        raw = b""
        try:
            with open(path, "rb") as fh:
                raw = fh.read(2000000)
        except OSError:
            pass
        i = raw.find(b"<x:xmpmeta")
        if i >= 0:
            xmp = raw[i:raw.find(b"</x:xmpmeta>", i) + 12].decode("utf-8", "ignore")

            def bag(tag):
                m = re.search(r"<%s>(.*?)</%s>" % (tag, tag), xmp, re.S)
                return [re.sub(r"<[^>]+>", "", x).strip() for x in re.findall(r"<rdf:li[^>]*>(.*?)</rdf:li>", m.group(1), re.S)] if m else []
            out["keywords"] += bag("dc:subject")
            out["title"] = out["title"] or " ".join(bag("dc:title")[:1])
            out["description"] = out["description"] or " ".join(bag("dc:description")[:1])
            out["creator"] = out["creator"] or ", ".join(bag("dc:creator"))
    except Exception:
        pass
    seen, kws = set(), []
    for k in out["keywords"]:
        if k and k.lower() not in seen:
            seen.add(k.lower())
            kws.append(k)
    out["keywords"] = kws
    return out


def create_app(config=None):
    app = Flask(__name__)
    base = os.path.dirname(os.path.abspath(__file__))
    app.config.update(
        SECRET_KEY=os.environ.get("SECRET_KEY") or secrets.token_hex(32),
        DATABASE=os.environ.get("BULLETIN_DB", os.path.join(base, "bulletin.db")),
        UPLOAD_DIR=os.environ.get("BULLETIN_UPLOADS", os.path.join(base, "uploads")),
        MAX_CONTENT_LENGTH=100 * 1024 * 1024,
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE="Lax",
    )
    if config:
        app.config.update(config)
    os.makedirs(app.config["UPLOAD_DIR"], exist_ok=True)

    def db():
        if "db" not in g:
            g.db = sqlite3.connect(app.config["DATABASE"])
            g.db.row_factory = sqlite3.Row
        return g.db

    @app.teardown_appcontext
    def close_db(_exc):
        conn = g.pop("db", None)
        if conn:
            conn.close()

    def init_db():
        conn = db()
        conn.executescript("""
        CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY, username TEXT UNIQUE,
            display_name TEXT, pw_hash TEXT, role TEXT, must_reset INTEGER DEFAULT 0);
        CREATE TABLE IF NOT EXISTS sections(id INTEGER PRIMARY KEY, slug TEXT UNIQUE, title TEXT,
            weekly INTEGER, pages INTEGER, position INTEGER, hint TEXT, bulletin_id INTEGER);
        CREATE TABLE IF NOT EXISTS photos(id INTEGER PRIMARY KEY, filename TEXT, title TEXT, alt TEXT,
            tags TEXT DEFAULT '', uploaded_by TEXT, decorative INTEGER DEFAULT 0);
        CREATE TABLE IF NOT EXISTS ads(id INTEGER PRIMARY KEY, name TEXT, image TEXT, w INTEGER, h INTEGER,
            position INTEGER DEFAULT 0, active INTEGER DEFAULT 1);
        CREATE TABLE IF NOT EXISTS questions(id INTEGER PRIMARY KEY, asked_by TEXT, question TEXT, answer TEXT DEFAULT '',
            answered_by TEXT DEFAULT '', created TEXT DEFAULT '');
        CREATE TABLE IF NOT EXISTS tokens(id INTEGER PRIMARY KEY, user_id INTEGER, kind TEXT, token_hash TEXT, expires TEXT);
        CREATE TABLE IF NOT EXISTS libraries(id INTEGER PRIMARY KEY, name TEXT UNIQUE);
        CREATE TABLE IF NOT EXISTS ai_usage(day TEXT PRIMARY KEY, n INTEGER DEFAULT 0);
        CREATE TABLE IF NOT EXISTS custom_fonts(id INTEGER PRIMARY KEY, name TEXT UNIQUE, filename TEXT);
        CREATE TABLE IF NOT EXISTS admin_notes(id INTEGER PRIMARY KEY, from_user TEXT, page TEXT, note TEXT, status TEXT DEFAULT 'open', created TEXT);
        CREATE TABLE IF NOT EXISTS hidden_sections(bulletin_id INTEGER, section_id INTEGER, snap TEXT, PRIMARY KEY(bulletin_id, section_id));
        CREATE TABLE IF NOT EXISTS photo_libs(photo_id INTEGER, lib_id INTEGER, PRIMARY KEY(photo_id, lib_id));
        CREATE TABLE IF NOT EXISTS event_titles(id INTEGER PRIMARY KEY, title TEXT UNIQUE);
        CREATE TABLE IF NOT EXISTS history(id INTEGER PRIMARY KEY, bulletin_id INTEGER, label TEXT, snapshot TEXT, created TEXT, by TEXT);
        CREATE TABLE IF NOT EXISTS user_sections(user_id INTEGER, section_id INTEGER,
            PRIMARY KEY(user_id, section_id));
        CREATE TABLE IF NOT EXISTS meta(k TEXT PRIMARY KEY, v TEXT);
        CREATE TABLE IF NOT EXISTS user_areas(user_id INTEGER, area TEXT, PRIMARY KEY(user_id, area));
        CREATE TABLE IF NOT EXISTS bulletins(id INTEGER PRIMARY KEY, sunday TEXT UNIQUE,
            theme TEXT DEFAULT '', status TEXT DEFAULT 'draft');
        CREATE TABLE IF NOT EXISTS items(id INTEGER PRIMARY KEY, bulletin_id INTEGER, section_id INTEGER,
            body TEXT DEFAULT '', image TEXT DEFAULT '', alt TEXT DEFAULT '', decorative INTEGER DEFAULT 0,
            caption TEXT DEFAULT '', source TEXT DEFAULT '', status TEXT DEFAULT 'approved',
            carried INTEGER DEFAULT 0, updated_by TEXT DEFAULT '', position REAL DEFAULT 0,
            pages INTEGER DEFAULT 1, cols INTEGER DEFAULT 1, UNIQUE(bulletin_id, section_id));
        """)
        if not conn.execute("SELECT 1 FROM sections").fetchone():
            for i, (slug, title, weekly, pages, hint) in enumerate(DEFAULT_SECTIONS):
                conn.execute("INSERT INTO sections(slug,title,weekly,pages,position,hint) VALUES(?,?,?,?,?,?)",
                             (slug, title, weekly, pages, i, hint))
        for slug, title, weekly, pages, hint in DEFAULT_SECTIONS:
            if not conn.execute("SELECT 1 FROM sections WHERE slug=?", (slug,)).fetchone():
                top = conn.execute("SELECT COALESCE(MAX(position),0)+1 m FROM sections").fetchone()["m"]
                after = {"mass-times": "staff", "mass-intentions": "cover"}.get(slug)
                anchor = conn.execute("SELECT position FROM sections WHERE slug=?", (after,)).fetchone() if after else None
                if anchor:
                    top = anchor["position"] + 1
                    conn.execute("UPDATE sections SET position=position+1 WHERE position>=?", (top,))
                conn.execute("INSERT INTO sections(slug,title,weekly,pages,position,hint) VALUES(?,?,?,?,?,?)",
                             (slug, title, weekly, pages, top, hint))
        if "footer" not in {r[1] for r in conn.execute("PRAGMA table_info(bulletins)")}:
            conn.execute("ALTER TABLE bulletins ADD COLUMN footer TEXT DEFAULT ''")
        bcols = {r[1] for r in conn.execute("PRAGMA table_info(bulletins)")}
        for col, ddl in (("hdr2_image", "TEXT DEFAULT ''"), ("hdr2_alt", "TEXT DEFAULT ''"), ("hdr2_zoom", "REAL DEFAULT 1"), ("hdr2_x", "REAL DEFAULT 50"), ("hdr2_y", "REAL DEFAULT 50"), ("hdr_split", "REAL DEFAULT 62"), ("hdr_image", "TEXT DEFAULT ''"), ("hdr_alt", "TEXT DEFAULT ''"), ("hdr_h", "REAL DEFAULT 2.4"),
                         ("hdr_zoom", "REAL DEFAULT 1"), ("hdr_x", "REAL DEFAULT 50"), ("hdr_y", "REAL DEFAULT 50"), ("hdr_style", "TEXT DEFAULT 'waves'"), ("hdr_date", "TEXT DEFAULT ''"), ("extra_pages", "INTEGER DEFAULT 0"), ("blank_at", "TEXT DEFAULT ''"), ("total_pages", "INTEGER DEFAULT 0"), ("ad_page_count", "INTEGER DEFAULT 0"), ("toc_title", "TEXT DEFAULT ''"), ("toc_title_gap", "REAL DEFAULT 0"), ("no_underline", "INTEGER DEFAULT 0"), ("toc_item_gap", "REAL DEFAULT 0"), ("toc_start", "INTEGER DEFAULT 2"), ("toc_show", "INTEGER DEFAULT 1"), ("toc_pt", "INTEGER DEFAULT 0"), ("toc_style", "TEXT DEFAULT ''"), ("toc_place", "TEXT DEFAULT 'left'"), ("toc_cols", "INTEGER DEFAULT 1"), ("no_stragglers", "INTEGER DEFAULT 1")):
            if col not in bcols:
                conn.execute(f"ALTER TABLE bulletins ADD COLUMN {col} {ddl}")
        ucols = {r[1] for r in conn.execute("PRAGMA table_info(users)")}
        if not conn.execute("SELECT 1 FROM libraries").fetchone():
            for nm in ("Church", "Logos", "Religious"):
                conn.execute("INSERT OR IGNORE INTO libraries(name) VALUES(?)", (nm,))
        if "meta" not in {r[1] for r in conn.execute("PRAGMA table_info(photos)")}:
            conn.execute("ALTER TABLE photos ADD COLUMN meta TEXT DEFAULT ''")
        if "active" not in ucols:
            conn.execute("ALTER TABLE users ADD COLUMN active INTEGER DEFAULT 1")
        for col in ("email", "last_login", "last_seen"):
            if col not in ucols:
                conn.execute(f"ALTER TABLE users ADD COLUMN {col} TEXT DEFAULT ''")
        if "is_template" not in bcols:
            conn.execute("ALTER TABLE bulletins ADD COLUMN is_template INTEGER DEFAULT 0")
        if "toc_before" not in bcols:
            conn.execute("ALTER TABLE bulletins ADD COLUMN toc_before REAL DEFAULT 0")
        if "url" not in {r[1] for r in conn.execute("PRAGMA table_info(ads)")}:
            conn.execute("ALTER TABLE ads ADD COLUMN url TEXT DEFAULT ''")
        have = {r[1] for r in conn.execute("PRAGMA table_info(items)")}
        for col, default in (("crop_zoom", 1), ("crop_x", 50), ("crop_y", 50), ("crop_h", 0), ("min_h", 0), ("crop_t", 0), ("crop_b", 0), ("crop_l", 0), ("crop_r", 0)):
            if col not in have:
                conn.execute(f"ALTER TABLE items ADD COLUMN {col} REAL DEFAULT {default}")
        if "img_pos" not in have:
            conn.execute("ALTER TABLE items ADD COLUMN img_pos TEXT DEFAULT ''")
        if "img_shadow" not in have:
            conn.execute("ALTER TABLE items ADD COLUMN img_shadow TEXT DEFAULT ''")
        if "bg_color" not in have:
            conn.execute("ALTER TABLE items ADD COLUMN bg_color TEXT DEFAULT ''")
        if "img_align" not in have:
            conn.execute("ALTER TABLE items ADD COLUMN img_align TEXT DEFAULT ''")
        if "img_w" not in have:
            conn.execute("ALTER TABLE items ADD COLUMN img_w REAL DEFAULT 0")
        if "gallery" not in have:
            conn.execute("ALTER TABLE items ADD COLUMN gallery TEXT DEFAULT ''")
        if "bullet_style" not in have:
            conn.execute("ALTER TABLE items ADD COLUMN bullet_style TEXT DEFAULT ''")
        if "space_before" not in have:
            conn.execute("ALTER TABLE items ADD COLUMN space_before REAL DEFAULT 0")
        if "no_toc" not in have:
            conn.execute("ALTER TABLE items ADD COLUMN no_toc INTEGER DEFAULT 0")
        if "hide_title" not in have:
            conn.execute("ALTER TABLE items ADD COLUMN hide_title INTEGER DEFAULT 0")
        if "para_gap" not in have:
            conn.execute("ALTER TABLE items ADD COLUMN para_gap REAL DEFAULT 0")
        if "text_center" not in have:
            conn.execute("ALTER TABLE items ADD COLUMN text_center INTEGER DEFAULT 0")
        if "text_nudge" not in have:
            conn.execute("ALTER TABLE items ADD COLUMN text_nudge REAL DEFAULT 0")
        if "vcenter" not in have:
            conn.execute("ALTER TABLE items ADD COLUMN vcenter INTEGER DEFAULT 0")
        if "item_gap" not in have:
            conn.execute("ALTER TABLE items ADD COLUMN item_gap REAL DEFAULT 0")
        if "title_gap" not in have:
            conn.execute("ALTER TABLE items ADD COLUMN title_gap REAL DEFAULT 0")
        if "title" not in have:
            conn.execute("ALTER TABLE items ADD COLUMN title TEXT DEFAULT ''")
        if "area" not in have:
            conn.execute("ALTER TABLE items ADD COLUMN area TEXT DEFAULT ''")
        if "font_pt" not in have:
            conn.execute("ALTER TABLE items ADD COLUMN font_pt INTEGER DEFAULT 0")
        if "font_style" not in have:
            conn.execute("ALTER TABLE items ADD COLUMN font_style TEXT DEFAULT ''")
        if "page_no" not in have:
            conn.execute("ALTER TABLE items ADD COLUMN page_no INTEGER DEFAULT 0")
        conn.execute("UPDATE sections SET title=?, hint=? WHERE slug='mass-schedule' AND title=?",
                     ("This Week in Our Parish: Lectors, Sacristans & Ministers",
                      "Weekly form: header, then day, date, lectors, sacristans and Eucharistic ministers (several names allowed)",
                      "Mass Schedule, Intentions & Ministers"))
        if not conn.execute("SELECT 1 FROM meta WHERE k='areas_v1'").fetchone():
            for u in conn.execute("SELECT id FROM users WHERE role NOT IN ('admin','editor')").fetchall():
                for a in ("ads", "library", "browse"):
                    conn.execute("INSERT OR IGNORE INTO user_areas VALUES(?,?)", (u["id"], a))
            conn.execute("INSERT INTO meta(k,v) VALUES('areas_v1','1')")
        for slug, old, new in (("calendar", "This Week in Our Parish", "This Week in Our Parish (Calendar of Events)"),
                               ("mass-schedule", "This Week in Our Parish", "This Week in Our Parish: Lectors, Sacristans & Ministers")):
            conn.execute("UPDATE sections SET title=? WHERE slug=? AND title=?", (new, slug, old))
        if not conn.execute("SELECT 1 FROM users").fetchone():
            pw = os.environ.get("BULLETIN_ADMIN_PASSWORD") or secrets.token_urlsafe(12)
            conn.execute("INSERT INTO users(username,display_name,pw_hash,role,must_reset) VALUES(?,?,?,?,1)",
                         ("admin", "Administrator", generate_password_hash(pw), "admin"))
            if not os.environ.get("BULLETIN_ADMIN_PASSWORD"):
                print(f"[bulletin] One-time admin password: {pw}  (change it at first login)")
        conn.commit()

    with app.app_context():
        init_db()

    # ---------- auth helpers ----------
    HISTORY_LABELS = {"layout": "Moved or resized a section", "reset-layout": "Reset layout", "section": "Saved a section",
                      "add-section": "Added a section", "delete": "Deleted section content", "extra-pages": "Added or removed pages",
                      "assistant": "Assistant changes", "approve": "Approved a section", "copy": "Copied a section", "keep": "Changed keep/weekly",
                      "header": "Changed header", "footer": "Changed footer"}

    def snapshot_bulletin(bid, label):
        conn = db()
        b = conn.execute("SELECT * FROM bulletins WHERE id=?", (bid,)).fetchone()
        if not b:
            return
        snap = {"bulletin": dict(b),
                "items": [dict(r) for r in conn.execute("SELECT * FROM items WHERE bulletin_id=?", (bid,))],
                "sections": [dict(r) for r in conn.execute("SELECT * FROM sections WHERE bulletin_id=?", (bid,))]}
        data = json.dumps(snap)
        last = conn.execute("SELECT snapshot FROM history WHERE bulletin_id=? ORDER BY id DESC LIMIT 1", (bid,)).fetchone()
        if last and last["snapshot"] == data:
            return
        conn.execute("INSERT INTO history(bulletin_id,label,snapshot,created,by) VALUES(?,?,?,?,?)",
                     (bid, label, data, datetime.now().strftime("%Y-%m-%d %H:%M:%S"), g.user["username"] if g.user else ""))
        conn.execute("DELETE FROM history WHERE bulletin_id=? AND id NOT IN (SELECT id FROM history WHERE bulletin_id=? ORDER BY id DESC LIMIT 40)", (bid, bid))
        conn.commit()

    @app.before_request
    def load_user():
        g.user = None
        uid = session.get("uid")
        if uid:
            g.user = db().execute("SELECT * FROM users WHERE id=?", (uid,)).fetchone()
            if g.user is None or g.user["active"] == 0:
                g.user = None
                session.clear()
        if not session.get("csrf"):
            session["csrf"] = secrets.token_hex(16)
        if request.method == "POST":
            if not secrets.compare_digest(request.form.get("csrf", ""), session["csrf"]):
                abort(400, "Invalid form token")
        if g.user and request.endpoint != "static":
            now = datetime.utcnow()
            if not g.user["last_seen"] or g.user["last_seen"] < (now - timedelta(seconds=60)).isoformat():
                db().execute("UPDATE users SET last_seen=? WHERE id=?", (now.isoformat(), g.user["id"]))
                db().commit()
        if (g.user and g.user["must_reset"] and request.endpoint
                not in ("password", "logout", "static", "forgot", "reset_with_token", "register")):
            return redirect(url_for("password"))
        if g.user and g.user["role"] == "webmaster" and request.endpoint not in WEBMASTER_PAGES:
            abort(403)

    app.jinja_env.globals["csrf"] = lambda: session.get("csrf", "")

    def calendar_rows(body):
        rows = []
        for line in (body or "").splitlines():
            if line.strip():
                parts = [p.strip() for p in line.split("|")]
                rows.append((parts + ["", "", ""])[:3] if len(parts) >= 3 else [parts[0], "", " ".join(parts[1:])])
        return rows

    app.jinja_env.globals["calendar_rows"] = calendar_rows

    CAL_MIN = ("lectors", "sacristans", "ministers", "location", "end_date")

    def calendar_entries(body):
        """Calendar rows as dicts, with optional weekend Mass ministers (lectors, sacristans, Eucharistic ministers)."""
        out = []
        for line in (body or "").splitlines():
            if not line.strip():
                continue
            parts = [p.strip() for p in line.split("|")]
            if len(parts) >= 3:
                parts = (parts + [""] * 8)[:8]
                out.append(dict(zip(("date", "time", "event") + CAL_MIN, parts)))
            else:
                out.append(dict(zip(("date", "time", "event") + CAL_MIN, [parts[0], "", " ".join(parts[1:]), "", "", "", "", ""])))
        return out

    app.jinja_env.globals["calendar_entries"] = calendar_entries

    def youth_events(body):
        """Youth calendar rows plus month/day/weekday pieces for the colourful cards."""
        out = []
        for i, ev in enumerate(calendar_entries(body)):
            d = None
            try:
                d = date.fromisoformat(iso_date(ev["date"]))
            except (ValueError, TypeError):
                pass
            month, day, weekday = (f"{d:%b}".upper() if d else ""), (d.day if d else ""), (f"{d:%A}" if d else ev["date"])
            try:
                e = date.fromisoformat(iso_date(ev["end_date"]))
            except (ValueError, TypeError):
                e = None
            if d and e and e > d:
                same = (e.year, e.month) == (d.year, d.month)
                month = month if same else f"{month}-{e:%b}".upper()
                day = f"{d.day}-{e.day}"
                weekday = f"{d:%A} to {e:%A}"
            ev = dict(ev, color=i % 4, month=month, day=day, weekday=weekday)
            out.append(ev)
        return out

    app.jinja_env.globals["youth_events"] = youth_events

    def long_date(iso):
        try:
            d = date.fromisoformat((iso or "").strip())
        except ValueError:
            return ""
        return f"{d:%A, %B} {d.day}, {d.year}"

    def iso_date(text):
        try:
            return datetime.strptime((text or "").strip(), "%A, %B %d, %Y").date().isoformat()
        except ValueError:
            return ""

    def form_vals(prefix, i, fields):
        vals = []
        for k in fields:
            v = request.form.get(f"{prefix}{i}_{k}", "").replace("|", "/").strip()
            if k == "notes":
                v = "\\n".join(l.strip() for l in request.form.get(f"{prefix}{i}_{k}", "").replace("|", "/").splitlines() if l.strip())
            vals.append("" if k == "day" else long_date(v) if k == "date" else v)
        return vals

    app.jinja_env.globals["iso_date"] = iso_date

    def styled(esc):
        import re
        from markupsafe import Markup
        s = str(esc)
        for _ in range(3):
            s = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", s, flags=re.S)
            s = re.sub(r"~~(.+?)~~", r"<em>\1</em>", s, flags=re.S)
            s = re.sub(r"__(.+?)__", r'<span class="ul">\1</span>', s, flags=re.S)
            s = re.sub(r"\{(\d{1,2})pt\|(.+?)\}", lambda m: '<span style="font-size:%dpt">%s</span>' % (min(40, max(8, int(m.group(1)))), m.group(2)), s, flags=re.S)
            s = re.sub(r"\{(#[0-9a-fA-F]{6})\|(.+?)\}", r'<span style="color:\1">\2</span>', s, flags=re.S)
        s = re.sub(r"(?<![\w#])(\d+)(st|nd|rd|th)\b", r"\1<sup>\2</sup>", s)
        return Markup(s)

    def sup_ordinals(html):
        """Superscript 1st/2nd/3rd/4th... in text everywhere (not inside tags, scripts, styles or an existing sup)."""
        import re
        rx = re.compile(r"(?<![\w#.$])(\d+)(st|nd|rd|th)\b(?![-'])", re.I)
        out, skip = [], 0
        for part in re.split(r"(<[^>]*>)", html):
            if part.startswith("<") and part.endswith(">"):
                m = re.match(r"</?\s*([a-zA-Z0-9]+)", part)
                if m and m.group(1).lower() in ("script", "style", "sup", "textarea", "title"):
                    skip += -1 if part.startswith("</") else 1
                    skip = max(skip, 0)
                out.append(part)
            else:
                out.append(part if skip else rx.sub(lambda m: m.group(1) + "<sup>" + m.group(2) + "</sup>", part))
        return "".join(out)
    app.jinja_env.globals["sup_ordinals"] = sup_ordinals

    def page_vcenter(pg):
        """True when any section placed on this page asks for its text to be centered top-to-bottom."""
        def walk(x):
            if hasattr(x, "keys"):
                ks = list(x.keys())
                if "vcenter" in ks and x["vcenter"]:
                    return True
                return any(walk(x[k]) for k in ks if k not in ("body", "image"))
            if isinstance(x, (list, tuple)):
                return any(walk(v) for v in x)
            return False
        return pg["number"] != 1 and not pg.get("photo_h") and walk([pg.get(k) for k in ("row", "more", "duo", "left_more", "right_more")])
    app.jinja_env.globals["page_vcenter"] = page_vcenter

    def ocr_text(path):
        """Read text from a picture with the tesseract program; None when it is not installed or fails."""
        import shutil
        exe = shutil.which("tesseract")
        if not exe:
            return None
        try:
            out = subprocess.run([exe, path, "stdout"], capture_output=True, text=True, timeout=60)
        except (OSError, subprocess.SubprocessError):
            return None
        return " ".join(out.stdout.split()) if out.returncode == 0 else None

    def shadow_parse(raw):
        try:
            x, y, b, o = (float(v) for v in str(raw or "").split(","))
        except ValueError:
            return None
        clamp = lambda v, lo, hi: max(lo, min(hi, v))
        d = {"x": clamp(x, -30, 30), "y": clamp(y, -30, 30), "b": clamp(b, 0, 30), "o": clamp(o, 5, 100)}
        a1 = 1 - (1 - d["o"] / 100) ** (1 / 8)
        n = 8 if d["b"] > 0 else 1
        al = a1 if n == 8 else d["o"] / 100
        d["layers"] = [{"e": round(-d["b"] / 2 + d["b"] * (i + 0.5) / n, 2), "a": round(al, 4)} for i in range(n)]
        d["L"], d["R"] = max(0, d["b"] / 2 - d["x"]), max(0, d["b"] / 2 + d["x"])
        d["T"], d["B"] = max(0, d["b"] / 2 - d["y"]), max(0, d["b"] / 2 + d["y"])
        d["raw"] = f'{d["x"]:g},{d["y"]:g},{d["b"]:g},{d["o"]:g}'
        return d

    def shadow_from_form(prefix, suffix=""):
        if not request.form.get(f"{prefix}on{suffix}"):
            return ""
        vals = []
        for k, dflt in (("x", 3), ("y", 3), ("b", 6), ("o", 45)):
            try:
                vals.append(float(request.form.get(f"{prefix}{k}{suffix}", dflt)))
            except ValueError:
                vals.append(dflt)
        d = shadow_parse(",".join(str(v) for v in vals))
        return d["raw"] if d else ""

    def gallery_data(raw):
        try:
            d = json.loads(raw or "")
        except ValueError:
            d = {}
        if not isinstance(d, dict):
            d = {}
        d.setdefault("h", 1.5)
        d["size"] = d.get("size") if d.get("size") in ("width", "area") else "height"
        d["valign"] = d.get("valign") if d.get("valign") in ("middle", "bottom") else "top"
        d["dir"] = "stack" if d.get("dir") == "stack" else "row"
        d["wrap"] = d.get("wrap") if d.get("wrap") in ("left", "right") else ""
        d["align"] = d.get("align") if d.get("align") in ("left", "right") else "center"
        for k in ("before", "after"):
            d[k] = d.get(k, 0.06) if isinstance(d.get(k), (int, float)) else 0.06
        d["gap"] = d.get("gap", 0.15) if isinstance(d.get("gap"), (int, float)) else 0.15
        d["photos"] = [p for p in d.get("photos", []) if isinstance(p, dict) and p.get("f")]
        for p in d["photos"]:
            try:
                from PIL import Image
                with Image.open(os.path.join(app.config["UPLOAD_DIR"], p["f"])) as im:
                    p["ar"] = max(0.2, min(5.0, im.width / im.height))
            except Exception:
                p["ar"] = 1.5
            p["ar0"] = p["ar"]
            p["shape"] = p.get("shape") if p.get("shape") in PHOTO_SHAPES else ""
            if p["shape"]:
                p["ar"] = PHOTO_SHAPES[p["shape"]][0]
            for k, lo, hi, dflt in (("z", 1, 4, 1), ("x", 0, 100, 50), ("y", 0, 100, 50)):
                try:
                    p[k] = max(lo, min(hi, float(p.get(k, dflt))))
                except (TypeError, ValueError):
                    p[k] = dflt
            p["wrap"] = p.get("wrap") if p.get("wrap") in ("left", "right") else ""
            p["crop"] = bool(p["shape"] or p["z"] > 1)
            p["shd"] = shadow_parse(p.get("sh"))
        return d

    def is_dark(hexcolor):
        try:
            r, g_, b_ = (int(hexcolor[i:i + 2], 16) / 255 for i in (1, 3, 5))
        except (ValueError, TypeError, IndexError):
            return False
        lin = lambda v: v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
        lum = 0.2126 * lin(r) + 0.7152 * lin(g_) + 0.0722 * lin(b_)
        return lum < 0.18

    app.jinja_env.globals["is_dark"] = is_dark
    app.jinja_env.globals["gallery_data"] = gallery_data
    app.jinja_env.globals["shadow_parse"] = shadow_parse

    AD_SHADOW_DEFAULT = "5,5,0,70"

    def ad_shadow_raw():
        r = db().execute("SELECT v FROM meta WHERE k='ad_shadow'").fetchone()
        return r["v"] if r else AD_SHADOW_DEFAULT

    def ad_gap_value():
        r = db().execute("SELECT v FROM meta WHERE k='ad_gap'").fetchone()
        try:
            return max(0.05, min(0.6, float(r["v"]))) if r else AD_GAP_X
        except ValueError:
            return AD_GAP_X

    def ad_layout():
        """(gap in inches, (sideways, downward) room in inches the shadow needs between and under cards)."""
        sd = shadow_parse(ad_shadow_raw())
        if not sd:
            return ad_gap_value(), (0.0, 0.0)
        return ad_gap_value(), ((abs(sd["x"]) + sd["b"] / 2) / 72, (abs(sd["y"]) + sd["b"] / 2) / 72)

    app.jinja_env.globals["ad_gap_value"] = ad_gap_value
    app.jinja_env.globals["ad_shadow_raw"] = ad_shadow_raw
    app.jinja_env.globals["ad_shadow"] = lambda: shadow_parse(ad_shadow_raw())
    app.jinja_env.globals["PHOTO_SHAPES"] = PHOTO_SHAPES

    def edge_values(form):
        out = []
        for k in ("crop_t", "crop_b", "crop_l", "crop_r"):
            try:
                out.append(max(0.0, min(80.0, float(form.get(k) or 0))))
            except ValueError:
                out.append(0.0)
        if out[0] + out[1] > 90:
            out[0] = out[1] = 0.0
        if out[2] + out[3] > 90:
            out[2] = out[3] = 0.0
        return out

    def edge_crop(name, t, b, l, r):
        """CSS for a picture with its edges trimmed (percent of each side), or None."""
        t, b, l, r = (float(v or 0) / 100 for v in (t, b, l, r))
        if not (t or b or l or r) or not name:
            return None
        try:
            from PIL import Image
            with Image.open(os.path.join(app.config["UPLOAD_DIR"], name)) as im:
                ratio = im.size[0] / im.size[1]
        except Exception:
            return None
        w, h = 1 - l - r, 1 - t - b
        return dict(box="position:relative;overflow:hidden;height:0;padding-bottom:%.3f%%" % (100 * h / (ratio * w)),
                    img="position:absolute;max-width:none;width:%.3f%%;height:%.3f%%;left:-%.3f%%;top:-%.3f%%" % (100 / w, 100 / h, l / w * 100, t / h * 100))
    app.jinja_env.globals["edge_crop"] = edge_crop

    def cs_plan_est(r, width_in, text_w_in):
        """Split a prose section's paragraphs into those that sit beside a column-spanning photo and those that flow below it.
        Lines are measured with the real bundled fonts so every column ends level instead of running past the footer."""
        import re
        from PIL import Image, ImageFont
        fdir = os.path.join(app.root_path, "static", "fonts")
        fam = {"sans": ("LiberationSans-Regular.ttf", 1.0), "source": ("SourceSans3-Regular.ttf", 1.0),
               "condensed": ("ArchivoNarrow-Regular.ttf", 1.0), "condensed-serif": ("NotoSerifCondensed-Regular.ttf", 1.0), "humanist": ("LiberationSans-Regular.ttf", 1.1),
               "rounded": ("LiberationSans-Regular.ttf", 1.08)}.get(r["font_style"] or "", ("Gelasio-Regular.ttf", 1.0))
        font = ImageFont.truetype(os.path.join(fdir, fam[0]), 100)
        fs = float(r["font_pt"] or 12.5)
        lh_in = fs * 1.4 / 72
        gap_in = float(r["para_gap"]) if r["para_gap"] else fs / 72
        plain = lambda t: re.sub(r"\*\*|~~|__|[{}]", "", t)

        def wrap(text, w_in, fsz=fs, f=font, k=fam[1]):
            avail = w_in * 72
            words, out, cur = plain(text).split(" "), [], ""
            for wd in words:
                trial = (cur + " " + wd) if cur else wd
                if cur and f.getlength(trial) * fsz / 100 * k > avail:
                    out.append(cur)
                    cur = wd
                else:
                    cur = trial
            out.append(cur)
            return out

        try:
            with Image.open(os.path.join(app.config["UPLOAD_DIR"], r["image"])) as im:
                ratio = im.size[0] / im.size[1]
            t, b_, l, rt = (float(r[k] or 0) / 100 for k in ("crop_t", "crop_b", "crop_l", "crop_r"))
            ratio = ratio * (1 - l - rt) / (1 - t - b_)
            h = float(r["crop_h"]) if r["crop_h"] else width_in / ratio
        except Exception:
            h = 0.0
        if r["caption"]:
            ital = ImageFont.truetype(os.path.join(fdir, "Gelasio-Italic-Regular.ttf"), 100)
            h += len(wrap(r["caption"], width_in, 12, ital, 1.0)) * 12 * 1.2 / 72 + 0.1
        else:
            h += 0.06
        paras = soft_paras(r["body"])
        units = []
        for para in paras:
            if isinstance(para, tuple):
                units.append(("ul", para[1], sum(len(wrap(it, text_w_in - 0.3)) for it in para[1]) * lh_in + gap_in))
            else:
                units.append(("p", para, sum(len(wrap(ln, text_w_in)) for ln in para) * lh_in + gap_in))
        total = sum(u[2] for u in units)
        n = max(1, int(r["cols"] or 2) - 1)
        target = (total + n * h) / (n + 1) - 2.5 * lh_in
        beside, rest, used = [], [], 0.0
        for idx, (kind, content, uh) in enumerate(units):
            if used + uh <= target + gap_in:
                beside.append(("ul", content) if kind == "ul" else content)
                used += uh
                continue
            room = int((target - used) / lh_in)
            if kind == "p" and room >= 2:
                lines = content
                # whole lines first, then a word-level cut inside the first line that does not fit
                taken, k = 0, 0
                while k < len(lines) and taken + len(wrap(lines[k], text_w_in)) <= room:
                    taken += len(wrap(lines[k], text_w_in))
                    k += 1
                if k:
                    beside.append(lines[:k])
                if k < len(lines):
                    ln = lines[k]
                    if not any(m in ln for m in ("**", "~~", "__", "{")) and room - taken >= 1:
                        ws = ln.split(" ")
                        cut, cnt = 0, 0
                        for q in range(1, len(ws) + 1):
                            if len(wrap(" ".join(ws[:q]), text_w_in)) > room - taken:
                                break
                            cut = q
                        if 0 < cut < len(ws):
                            beside.append([" ".join(ws[:cut])])
                            lines = [" ".join(ws[cut:])] + lines[k + 1:]
                            rest = [lines] + [("ul", u[1]) if u[0] == "ul" else u[1] for u in units[idx + 1:]]
                            return beside, rest
                    rest = [lines[k:]] + [("ul", u[1]) if u[0] == "ul" else u[1] for u in units[idx + 1:]]
                    return beside, rest
                rest = [("ul", u[1]) if u[0] == "ul" else u[1] for u in units[idx + 1:]]
                return beside, rest
            rest = [("ul", u[1]) if u[0] == "ul" else u[1] for u in units[idx:]]
            return beside, rest
        return beside, rest
    cs_cache = {}

    def cs_plan(r, width_in, text_w_in, no_str=True):
        """Decide exactly where the text beside a column-spanning photo stops and the text under the photo begins.
        The real PDF engine lays out the row for each candidate split, and the largest amount of text that keeps the
        beside column no taller than the photo + continued-text block is chosen. Nothing is estimated."""
        import hashlib
        key = hashlib.sha1(repr((r["body"], r["image"], r["caption"], r["cols"], r["font_pt"], r["font_style"], r["para_gap"],
                                 r["crop_t"], r["crop_b"], r["crop_l"], r["crop_r"], r["crop_h"], r["img_w"], r["bullet_style"],
                                 r["slug"], r["img_pos"], round(width_in, 3), round(text_w_in, 3), no_str)).encode()).hexdigest()
        if key in cs_cache:
            return cs_cache[key]
        try:
            res = cs_plan_exact(r, width_in, text_w_in, no_str)
        except Exception:
            app.logger.exception("exact column-wrap layout failed; using estimate")
            res = tuple(cs_plan_est(r, width_in, text_w_in)) + (None,)
        if len(cs_cache) > 200:
            cs_cache.clear()
        cs_cache[key] = res
        return res

    class ContPara(list):
        cont = True

    def cs_plan_exact(r, width_in, text_w_in, no_str):
        from weasyprint import HTML
        from markupsafe import Markup, escape
        from PIL import Image
        n = max(1, int(r["cols"] or 2) - 1)
        units = soft_paras(r["body"])
        marked = lambda t: any(m in t for m in ("**", "~~", "__", "{"))
        # every place the text may be cut: a paragraph boundary, or a word boundary inside a plain line
        cuts = []
        for u, para in enumerate(units):
            cuts.append((u, 0, 0))
            if isinstance(para, tuple):
                continue
            for l, ln in enumerate(para):
                if marked(ln):
                    continue
                for q in range(1, len(ln.split(" "))):
                    cuts.append((u, l, q))
        cuts.append((len(units), 0, 0))

        def split(c):
            u, l, q = c
            if not q:
                return list(units[:u]), list(units[u:])
            ws = units[u][l].split(" ")
            head = list(units[:u]) + [units[u][:l] + [" ".join(ws[:q])]]
            tail = [ContPara([" ".join(ws[q:])] + units[u][l + 1:])] + list(units[u + 1:])
            return head, tail

        def block(paras):
            out = Markup("")
            for para in paras:
                if isinstance(para, tuple):
                    out += Markup('<ul class="bl b-%s">' % escape(r["bullet_style"] or "disc")) + Markup("").join(
                        Markup("<li>") + tight(it, no_str) + Markup("</li>") for it in para[1]) + Markup("</ul>")
                else:
                    out += Markup('<p class="bp">') + Markup("<br>").join(tight(ln, no_str) for ln in para) + Markup("</p>")
            return out

        try:
            with Image.open(os.path.join(app.config["UPLOAD_DIR"], r["image"])) as im:
                ratio = im.size[0] / im.size[1]
            t, b_, l_, rt = (float(r[k] or 0) / 100 for k in ("crop_t", "crop_b", "crop_l", "crop_r"))
            ratio = ratio * (1 - l_ - rt) / (1 - t - b_)
        except Exception:
            ratio = 1.5
        photo_h = float(r["crop_h"]) if r["crop_h"] else width_in / ratio
        fig_style = (' style="width:%d%%;margin-left:auto;margin-right:auto"' % int(r["img_w"])) if r["img_w"] else ""
        cap = Markup("<figcaption>%s</figcaption>") % r["caption"] if r["caption"] else Markup("")
        sty = ""
        if r["font_pt"]:
            sty += "--fs: %spt;" % r["font_pt"]
        if r["para_gap"]:
            sty += "--pg: %sin;" % r["para_gap"]
        cls = "sec sec-%s %s %s %s" % (r["slug"], "has-pg" if r["para_gap"] else "", "custom-size" if r["font_pt"] else "",
                                         "fs-" + r["font_style"] if r["font_style"] else "")
        aw = width_in + 0.35 + text_w_in
        title_h = "" if r["hide_title"] else "<h2>%s</h2>" % escape(r["title"] or "")
        has_rest_area = r["slug"] not in ("staff", "giving")

        def measure(c, ph):
            head, tail = split(c)
            rest = ('<div class="body-cols" style="column-count:%d">%s</div>' % (n, block(tail))) if (tail and has_rest_area) else ""
            photo = ('<td class="cs-photo" style="width:%sin"><div data-m="c"><figure class="photo-%s"%s><div style="height:%sin"></div>%s</figure>%s</div></td>'
                     % (width_in, r["img_pos"], fig_style, ph, cap, rest))
            text = ('<td class="cs-textcell" style="width:%sin"><div data-m="b" class="body-cols cs-text" style="column-count:1">%s</div></td>'
                    % (text_w_in, block(head)))
            cells = (photo + '<td class="cs-gap"></td>' + text) if r["img_pos"] == "cols-left" else (text + '<td class="cs-gap"></td>' + photo)
            html = ('<!doctype html><html><head><meta charset="utf-8"><link rel="stylesheet" href="%sstatic/style.css">'
                    '<style>@page{size:9in 80in;margin:0}.sheet{display:block;width:9in;height:auto;margin:0;padding:0;border:0;overflow:visible}.sheet .pg-body{position:static!important;display:block;width:8in;height:auto!important;overflow:visible;flex:none}</style></head>'
                    '<body><div class="sheet"><div class="pg-body"><div class="%s" style="%s"><div data-m="t">%s<table class="cs-wrap" style="width:%sin"><tr>%s</tr></table></div></div></div></div></body></html>'
                    % (PDF_BASE, cls, sty, title_h, aw, cells))
            doc = HTML(string=html, base_url=PDF_BASE, url_fetcher=pdf_fetcher).render()
            hs = {}
            for page in doc.pages:
                for box in page._page_box.descendants():
                    el = getattr(box, "element", None)
                    m = el.get("data-m") if el is not None else None
                    if m and m not in hs:
                        hs[m] = box.margin_height()
            return hs.get("b", 0.0), hs.get("c", 0.0), hs.get("t", 0.0)

        def solve(ph):
            lo, hi = 0, len(cuts) - 1
            while lo < hi:
                mid = (lo + hi + 1) // 2
                bh, ch, _ = measure(cuts[mid], ph)
                if bh <= ch + 0.5:
                    lo = mid
                else:
                    hi = mid - 1
            return lo, measure(cuts[lo], ph)[2]

        # the page body is 9.58in tall; if the whole row is taller, shorten the photo (cropped, never squashed) until it fits
        avail = (9.58 - 0.05) * 96
        shrinkable = not any(float(r[k] or 0) for k in ("crop_t", "crop_b", "crop_l", "crop_r"))
        lo, total = solve(photo_h)
        used_h = None
        floor = max(1.0, photo_h * 0.4)
        ph = photo_h
        for _ in range(3):
            if not shrinkable or total <= avail or ph <= floor + 0.01:
                break
            ph = max(floor, ph - (total - avail) / 96 * (n + 1) / n - 0.02)
            used_h = round(ph, 3)
            lo, total = solve(ph)
        head, tail = split(cuts[lo])
        conv = lambda ps: [("ul", p[1]) if isinstance(p, tuple) else p for p in ps]
        return conv(head), conv(tail), used_h
    app.jinja_env.globals["cs_plan"] = cs_plan

    def tight(text, on=True):
        """With no-stragglers on, bind the last two words with a no-break space so a paragraph never ends with one word alone on a line."""
        from markupsafe import Markup, escape
        text = str(text or "")
        if on is not None and not on:
            return styled(escape(text))
        words = text.split(" ")
        if len(words) < 3:
            return styled(escape(text))
        out = escape(words[0])
        for i in range(1, len(words)):
            first_of_sentence = i >= 2 and words[i - 2][-1:] in ".!?"
            sep = Markup("&nbsp;") if (i == len(words) - 1 or first_of_sentence) else " "
            out += sep + escape(words[i])
        return styled(out)
    def soft_paras(body):
        """Split text into paragraphs; a line ending in // is a soft return (line break, no paragraph gap)."""
        import re

        def per_line(m):
            op, inner, cl = m.group(1), m.group(2), m.group(3)
            if "\n" not in inner:
                return m.group(0)
            return "\n".join(op + ln + cl if ln.strip() else ln for ln in inner.split("\n"))
        body = body or ""
        for pat in (r"(\{\d{1,2}pt\|)(.+?)(\})", r"(\{#[0-9a-fA-F]{6}\|)(.+?)(\})", r"(\*\*)(.+?)(\*\*)", r"(~~)(.+?)(~~)", r"(__)(.+?)(__)"):
            body = re.sub(pat, per_line, body, flags=re.S)
        out, cur = [], []
        for line in body.split("\n"):
            t = line.strip()
            if t[:2] in ("- ", "* ", "\u2022 "):
                if cur:
                    out.append(cur)
                    cur = []
                if out and isinstance(out[-1], tuple):
                    out[-1][1].append(t[2:].strip())
                else:
                    out.append(("ul", [t[2:].strip()]))
                continue
            if t.endswith("//"):
                t = t[:-2].rstrip()
                if t:
                    cur.append(t)
                continue
            if t:
                cur.append(t)
            if cur:
                out.append(cur)
                cur = []
        if cur:
            out.append(cur)
        return out

    app.jinja_env.globals["soft_paras"] = soft_paras
    app.jinja_env.globals["BULLETS"] = BULLETS
    app.jinja_env.globals["tight"] = tight
    app.jinja_env.globals["gap_display"] = gap_display
    app.jinja_env.globals["long_sunday"] = long_date

    MASS_FIELDS = ("day", "date", "time", "lectors", "sacristans", "ministers", "notes")

    def split_names(text):
        return [n.strip() for n in re.split(r"[;,\n]", text or "") if n.strip()]

    def mass_data(body):
        """Return (header, rows); each row has its name fields split into lists."""
        header, rows = "", []
        for line in (body or "").splitlines():
            if line.startswith("#"):
                header = line[1:].strip()
            elif line.strip():
                parts = [p.strip() for p in line.split("|")]
                row = dict(zip(MASS_FIELDS, (parts + [""] * 7)[:7]))
                row["notes"] = row["notes"].replace("\\n", "\n")
                for k in ("lectors", "sacristans", "ministers"):
                    row[k + "_list"] = split_names(row[k])
                rows.append(row)
        return header, rows

    app.jinja_env.globals["mass_data"] = mass_data

    GIVING_KEYS = ("subtitle", "date", "end_date", "first", "second", "no_second", "purpose", "deposit", "notes", "first_label", "second_label", "deposit_label", "extras")

    def money(text):
        t = (text or "").strip().lstrip("$").replace(",", "")
        try:
            return "${:,.2f}".format(float(t))
        except ValueError:
            return (text or "").strip()

    def giving_data(body):
        """Structured Stewardship & Giving fields; older free text is kept as the notes."""
        d = dict.fromkeys(GIVING_KEYS, "")
        lines = (body or "").split("\n")
        if lines and lines[0].strip() == "#giving":
            for line in lines[1:]:
                k, _, v = line.partition("=")
                if k in d:
                    d[k] = v.replace("\\n", "\n")
        else:
            d["notes"] = (body or "").strip()
        d["subtitle"] = d["subtitle"].strip() or "Donations Received"
        d["hide_second"] = bool(d["no_second"]) or not d["second"].strip()
        MON = {1: "Jan.", 2: "Feb.", 3: "Mar.", 4: "Apr.", 5: "May", 6: "June", 7: "July", 8: "Aug.", 9: "Sept.", 10: "Oct.", 11: "Nov.", 12: "Dec."}

        def parse(v):
            try:
                return datetime.strptime(v, "%Y-%m-%d")
            except ValueError:
                return None
        a, z = parse(d["date"]), parse(d["end_date"])
        if not a:
            d["date_text"] = d["date"]
        elif not z or z == a:
            d["date_text"] = "%s %d" % (MON[a.month], a.day)
        elif (z.year, z.month) == (a.year, a.month):
            d["date_text"] = "%s %d-%d" % (MON[a.month], a.day, z.day)
        else:
            d["date_text"] = "%s %d \u2013 %s %d" % (MON[a.month], a.day, MON[z.month], z.day)
        d["first_text"], d["second_text"], d["deposit_text"] = money(d["first"]), money(d["second"]), money(d["deposit"])
        d["extra_rows"] = []
        for ln in d["extras"].split("\n"):
            parts = [x.strip() for x in (ln.split("|") + ["", "", ""])[:3]]
            if parts[0] or parts[1]:
                d["extra_rows"].append({"label": parts[0], "amount": parts[1], "amount_text": money(parts[1]), "note": parts[2]})
        return d

    app.jinja_env.globals["giving_data"] = giving_data

    ADDR_FIELDS = ("street", "city", "phone", "email", "web")

    def cover_data(body):
        """Parse the cover: @date | date | title, @address | ..., @hours | days | time | notes; other lines are extra text."""
        out = {"date": "", "title": "", "addr": dict.fromkeys(ADDR_FIELDS, ""), "hours": [], "extra": ""}
        extra = []
        for line in (body or "").splitlines():
            parts = [x.strip() for x in line.split("|")]
            tag = parts[0].lower()
            if tag == "@date":
                out["date"], out["title"] = (parts + ["", ""])[1], (parts + ["", ""])[2]
            elif tag == "@address":
                out["addr"] = dict(zip(ADDR_FIELDS, (parts[1:] + [""] * 5)[:5]))
                if not out["addr"]["web"].lower().startswith(("http://", "https://")):
                    out["addr"]["web"] = ""
            elif tag == "@hours":
                out["hours"].append(dict(zip(("days", "time", "notes"), (parts[1:] + ["", "", ""])[:3])))
            else:
                extra.append(line)
        out["extra"] = "\n".join(extra).strip()
        out["has_addr"] = any(out["addr"].values())
        return out

    app.jinja_env.globals["cover_data"] = cover_data

    def cover_prefilled(body, sunday):
        cv = cover_data(body)
        if not cv["date"]:
            cv["date"] = long_date(sunday)
        if not cv["has_addr"]:
            cv["addr"].update(street="31300 Overseas Hwy", city="Big Pine Key, FL 33043", phone="305-872-2537")
        if not cv["hours"]:
            cv["hours"] = [{"days": "Monday-Friday", "time": "9:30am-3:30pm", "notes": ""}]
        return cv

    TIMES_FIELDS = ("title", "days", "name", "time", "notes", "web")

    def times_rows(body):
        rows = []
        for line in (body or "").splitlines():
            if line.strip():
                parts = [p.strip() for p in line.split("|")]
                row = dict(zip(TIMES_FIELDS, (parts + [""] * 6)[:6]))
                if not row["web"].lower().startswith(("http://", "https://")):
                    row["web"] = ""
                rows.append(row)
        return rows

    app.jinja_env.globals["times_rows"] = times_rows

    STAFF_FIELDS = ("name", "role", "phone", "email")

    def staff_rows(body):
        rows = []
        for line in (body or "").splitlines():
            if line.strip():
                parts = [p.strip() for p in line.split("|")]
                rows.append(dict(zip(STAFF_FIELDS, (parts + [""] * 4)[:4])))
        return rows

    app.jinja_env.globals["staff_rows"] = staff_rows

    INTENT_FIELDS = ("day", "date", "time", "for_whom", "by_whom")

    def intention_data(body):
        """Return (header, days) where days groups intentions by (day, date), in entered order."""
        header, days = "", []
        for line in (body or "").splitlines():
            if line.startswith("#"):
                header = line[1:].strip()
            elif line.strip():
                parts = [p.strip() for p in line.split("|")]
                row = dict(zip(INTENT_FIELDS, (parts + [""] * 5)[:5]))
                if not days or (days[-1]["day"], days[-1]["date"]) != (row["day"], row["date"]):
                    days.append({"day": row["day"], "date": row["date"], "items": []})
                days[-1]["items"].append(row)
        return header, days

    app.jinja_env.globals["intention_data"] = intention_data

    def need(*roles):
        def deco(fn):
            @wraps(fn)
            def wrapper(*a, **kw):
                if not g.user:
                    return redirect(url_for("login", next=request.path))
                if roles and g.user["role"] not in roles:
                    abort(403)
                return fn(*a, **kw)
            return wrapper
        return deco

    AREAS = (("ads", "Advertisements"), ("library", "Photo library"), ("browse", "Browse photos"), ("slides", "Mass slides (AV Ministry)"), ("choir", "Choir: find and download music"))

    def has_area(name):
        if not g.user:
            return False
        if g.user["role"] in ("admin", "editor"):
            return True
        return bool(db().execute("SELECT 1 FROM user_areas WHERE user_id=? AND area=?", (g.user["id"], name)).fetchone())

    app.jinja_env.globals["has_area"] = has_area

    def need_area(name):
        def deco(fn):
            @wraps(fn)
            def wrapper(*a, **kw):
                if not has_area(name):
                    abort(403)
                return fn(*a, **kw)
            return wrapper
        return deco

    def save_areas(conn, uid):
        conn.execute("DELETE FROM user_areas WHERE user_id=?", (uid,))
        for a, _ in AREAS:
            if a in request.form.getlist("areas"):
                conn.execute("INSERT OR IGNORE INTO user_areas VALUES(?,?)", (uid, a))

    def can_edit_section(section_id):
        r = g.user["role"]
        if r in ("admin", "editor"):
            return True
        if r == "contributor":
            return bool(db().execute("SELECT 1 FROM user_sections WHERE user_id=? AND section_id=?",
                                     (g.user["id"], section_id)).fetchone())
        return False

    def safe_next(target):
        return target if target and target.startswith("/") and not target.startswith("//") else None

    # ---------- bulletin helpers ----------
    def next_sunday(d=None):
        d = d or date.today()
        return d + timedelta(days=(6 - d.weekday()) % 7)

    def upcoming_events(body, sunday):
        """Keep calendar lines dated on or after the new issue's Sunday; lines with no readable date are kept."""
        keep = []
        for line in (body or "").splitlines():
            cols = [c.strip() for c in line.split("|")]
            iso = (len(cols) > 7 and iso_date(cols[7])) or iso_date(cols[0])
            if iso and iso < sunday:
                continue
            keep.append(line)
        return "\n".join(keep)

    def week_free(body):
        """Drop the week-specific @date line so a carried cover does not repeat last week's date and title."""
        return "\n".join(l for l in (body or "").splitlines() if not l.strip().lower().startswith("@date"))

    def ensure_items(bulletin_id, sunday):
        """Create items: carry standing sections (and their layout) from the previous bulletin;
        weekly sections start empty. Custom sections belong to one bulletin only."""
        conn = db()
        prev = conn.execute("SELECT id FROM bulletins WHERE sunday<? ORDER BY sunday DESC LIMIT 1",
                            (sunday,)).fetchone()
        secs = conn.execute("SELECT * FROM sections WHERE bulletin_id IS NULL OR bulletin_id=? ORDER BY position",
                            (bulletin_id,)).fetchall()
        hidden = {r["section_id"] for r in conn.execute("SELECT section_id FROM hidden_sections WHERE bulletin_id=?", (bulletin_id,))}
        secs = [x for x in secs if x["id"] not in hidden]
        fresh = False
        for s in secs:
            if conn.execute("SELECT 1 FROM items WHERE bulletin_id=? AND section_id=?",
                            (bulletin_id, s["id"])).fetchone():
                continue
            fresh = True
            old = None
            if prev and s["bulletin_id"] is None:
                old = conn.execute("SELECT * FROM items WHERE bulletin_id=? AND section_id=?",
                                   (prev["id"], s["id"])).fetchone()
            pos, pages, cols = (old["position"], old["pages"], old["cols"]) if old else (s["position"], s["pages"], 1)
            carry_events = s["slug"] == "youth-calendar"
            if old and (not s["weekly"] or carry_events):
                conn.execute("""INSERT INTO items(bulletin_id,section_id,body,image,alt,decorative,caption,source,carried,
                    position,pages,cols,crop_zoom,crop_x,crop_y,crop_h,title,title_gap,item_gap,area,page_no,font_pt,font_style,min_h,img_pos,para_gap,hide_title,space_before,bullet_style,no_toc)
                    VALUES(?,?,?,?,?,?,?,?,1,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                             (bulletin_id, s["id"], week_free(old["body"]) if s["slug"] == "cover" else (upcoming_events(old["body"], sunday) if carry_events else old["body"]), old["image"], old["alt"], old["decorative"],
                              old["caption"], old["source"], pos, pages, cols,
                              old["crop_zoom"], old["crop_x"], old["crop_y"], old["crop_h"], old["title"], old["title_gap"], old["item_gap"],
                              old["area"], old["page_no"], old["font_pt"], old["font_style"], old["min_h"], old["img_pos"], old["para_gap"], old["hide_title"], old["space_before"], old["bullet_style"], old["no_toc"]))
                conn.execute("UPDATE items SET gallery=?, img_w=?, img_shadow=?, crop_t=?, crop_b=?, crop_l=?, crop_r=? WHERE bulletin_id=? AND section_id=?", (old["gallery"] or "", old["img_w"] or 0, old["img_shadow"] or "", old["crop_t"] or 0, old["crop_b"] or 0, old["crop_l"] or 0, old["crop_r"] or 0, bulletin_id, s["id"]))
            else:
                conn.execute("INSERT INTO items(bulletin_id,section_id,position,pages,cols,area,page_no,font_pt,font_style,title,title_gap,item_gap,min_h,img_pos,para_gap,hide_title,space_before,bullet_style,no_toc) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                             (bulletin_id, s["id"], pos, pages, cols, old["area"] if old else "", old["page_no"] if old else 0,
                              old["font_pt"] if old else 0, old["font_style"] if old else "", old["title"] if old else "",
                              old["title_gap"] if old else 0, old["item_gap"] if old else 0, old["min_h"] if old else 0, old["img_pos"] if old else "", old["para_gap"] if old else 0, old["hide_title"] if old else 0, old["space_before"] if old else 0, old["bullet_style"] if old else "", old["no_toc"] if old else 0))
        if fresh:
            fill_mass_week(conn, bulletin_id, sunday, prev)
        conn.commit()

    def fill_mass_week(conn, bulletin_id, sunday, prev):
        """Date the weekly Mass schedule: Saturday before the issue through the following Sunday.
        The previous issue's last Saturday and Sunday (with their names and notes) become the first two days."""
        item = conn.execute("""SELECT i.id, i.body FROM items i JOIN sections s ON s.id=i.section_id
                               WHERE i.bulletin_id=? AND s.slug='mass-schedule'""", (bulletin_id,)).fetchone()
        if not item or (item["body"] or "").strip():
            return
        try:
            sun = date.fromisoformat(sunday)
        except ValueError:
            return
        head, old = "", {}
        if prev:
            pb = conn.execute("""SELECT i.body FROM items i JOIN sections s ON s.id=i.section_id
                                 WHERE i.bulletin_id=? AND s.slug='mass-schedule'""", (prev["id"],)).fetchone()
            for line in ((pb["body"] if pb else "") or "").splitlines():
                if line.startswith("#"):
                    head = line[1:].strip()
                elif line.strip():
                    parts = [x.strip() for x in line.split("|")]
                    parts = (parts + [""] * 7)[:7]
                    d = iso_date(parts[1])
                    if d:
                        old[d] = parts
        lines = ["# " + head] if head else []
        for n in range(-1, 8):
            d = sun + timedelta(days=n)
            parts = old.get(d.isoformat()) if n in (-1, 0) else None
            parts = list(parts) if parts else ["", long_date(d.isoformat()), "", "", "", "", ""]
            parts[0] = ""
            lines.append(" | ".join(parts))
        conn.execute("UPDATE items SET body=? WHERE id=?", ("\n".join(lines), item["id"]))

    def area_of(r):
        """Where a section sits: full page width, or the left/right half. Mass Intentions defaults to a left half."""
        a = r["area"] or ("left" if r["slug"] == "mass-intentions" else "full")
        return a if a in ("left", "right", "below", "above", "left3", "mid3", "right3", "left23", "right23") and r["slug"] != "cover" and r["pages"] == 1 else "full"

    def keep_page_blank(bid, sid):
        """Called before a section is deleted: a page that held only that section stays as a blank page."""
        pages = layout(bid, include_pending=True)[0]
        gone = []
        for p in pages:
            if p.get("duo") or p.get("blank") or p.get("ad_rows") is not None:
                continue
            rs = [r for r in rows_on(p) if r is not None]
            if rs and all(r["id"] == sid for r in rs) and p["number"] >= 2:
                gone.append(p["number"])
        if gone:
            cur = db().execute("SELECT blank_at FROM bulletins WHERE id=?", (bid,)).fetchone()["blank_at"] or ""
            keep = [x for x in cur.split(",") if x.strip()]
            db().execute("UPDATE bulletins SET blank_at=? WHERE id=?", (",".join(keep + [str(n) for n in gone]), bid))

    def rows_on(p):
        d = p.get("duo")
        rows = []
        if d:
            for k in (("left3", "mid3", "right3", "left23", "right23") if d.get("third") else ("left", "right")):
                rows += [d[k], *d[k + "_more"]]
        else:
            rows = [p.get("row")]
        return list(p.get("above", [])) + rows + list(p.get("right_more", [])) + list(p.get("p1_below", [])) + list(p.get("more", []))

    def layout(bulletin_id, include_pending=False):
        """Return (pages, toc). Full-width sections take whole pages; left/right half sections share a page
        (placed on their chosen page number, else the first free half). Total padded to a multiple of 4."""
        rows = db().execute("""SELECT s.id, s.slug, COALESCE(NULLIF(i.title, ''), s.title) AS title, s.weekly, s.hint, s.bulletin_id,
            i.position, i.pages, i.cols, i.area, i.page_no, i.font_pt, i.font_style, i.id item_id, i.body, i.image, i.alt, i.decorative, i.caption,
            i.source, i.img_align, i.bg_color, i.status, i.carried, i.crop_zoom, i.crop_x, i.crop_y, i.crop_h, i.crop_t, i.crop_b, i.crop_l, i.crop_r, i.min_h, i.title_gap, i.item_gap, i.img_pos, i.para_gap, i.hide_title, i.space_before, i.bullet_style, i.gallery, i.img_w, i.img_shadow, i.no_toc, i.vcenter, i.text_center, i.text_nudge FROM sections s JOIN items i ON i.section_id=s.id
            WHERE i.bulletin_id=? ORDER BY i.position, s.id""", (bulletin_id,)).fetchall()
        printed = [r for r in rows
                   if (bool((cover_data(r["body"])["extra"] if r["slug"] == "cover" else (r["body"] or "").strip()) or r["image"]))
                   and not (r["status"] == "pending" and not include_pending)]
        pages = []
        def blocks(r):
            return [{"row": r, "part": n + 1, "of": r["pages"]} for n in range(r["pages"])]
        cover_printed = any(r["slug"] == "cover" for r in printed)
        p1_all = [] if cover_printed else sorted((r for r in printed if r["page_no"] == 1 and r["slug"] != "cover"),
                                                 key=lambda r: r["position"])
        brow = db().execute("SELECT toc_place, toc_show FROM bulletins WHERE id=?", (bulletin_id,)).fetchone()
        toc_full = bool(brow and brow["toc_place"] == "full" and (brow["toc_show"] is None or brow["toc_show"]))
        slot_keys = ("left", "right", "left3", "mid3", "right3", "left23", "right23")
        if toc_full and p1_all:
            duo = {"third": None, **{k: None for k in slot_keys}, **{k + "_more": [] for k in slot_keys}}
            under = []
            for r in p1_all:
                a = area_of(r)
                if a not in slot_keys:
                    under.append(r)
                    continue
                third = a.endswith("3")
                if duo["third"] is None:
                    duo["third"] = third
                if duo["third"] != third:
                    under.append(r)
                elif duo[a] is None:
                    duo[a] = r
                else:
                    duo[a + "_more"].append(r)
            first = next((duo[k] for k in slot_keys if duo[k]), None)
            if first is None:
                first, under = under[0], under[1:]
                duo = None
            pages.append({"row": first, "part": 1, "of": first["pages"], "duo": duo, "p1_full": True, "p1_below": under})
        else:
            p1_side = [r for r in p1_all if area_of(r) != "below"]
            p1_under = [r for r in p1_all if area_of(r) == "below"]
            if not p1_side and p1_under:
                p1_side, p1_under = p1_under[:1], p1_under[1:]
            p1 = p1_side[0] if p1_side else None
            if p1 is not None:
                pages.extend(blocks(p1))
                pages[0]["right_more"] = p1_side[1:]
                pages[0]["p1_below"] = p1_under
        fulls = [r for r in printed if area_of(r) == "full" and r not in p1_all]
        host_pages = {r["page_no"] for r in printed if r["page_no"] and r["page_no"] >= 2 and r["slug"] != "cover"
                      and area_of(r) not in ("above", "below") and r not in p1_all}
        orphan_bands = [r for r in printed if area_of(r) in ("above", "below") and r not in p1_all and r["slug"] != "cover"
                        and (r["page_no"] or 0) >= 2 and r["page_no"] not in host_pages]
        fulls = fulls + orphan_bands
        for r in fulls:
            if not (r["page_no"] and r["page_no"] >= 2 and r["slug"] != "cover"):
                pages.extend(blocks(r))
        full_on = {}
        for r in sorted((r for r in fulls if r["page_no"] and r["page_no"] >= 2 and r["slug"] != "cover"),
                        key=lambda r: (r["page_no"], r["position"])):
            host = full_on.get(r["page_no"])
            if host is not None and r["pages"] == 1 and host["row"]["pages"] == 1:
                host.setdefault("more", []).append(r)
                continue
            at = max(1, min(r["page_no"] - 1, len(pages)))
            new_blocks = blocks(r)
            pages[at:at] = new_blocks
            if new_blocks:
                full_on[r["page_no"]] = new_blocks[0]
                new_blocks[0]["pin"] = r["page_no"]
        if not pages:
            pages.append({"row": None})
        halves = [r for r in printed if area_of(r) in ("left", "right", "left3", "mid3", "right3", "left23", "right23") and r not in p1_all]
        def page_wanted(r):
            return r["page_no"] or (2 if r["slug"] == "mass-intentions" else 0)

        made = {}
        for r in sorted(halves, key=lambda r: (page_wanted(r) or 99, r["position"])):
            side, want = area_of(r), max(2, page_wanted(r))
            third = side.endswith("3")
            def fits(d):
                if bool(d.get("third")) != third or d[side] is not None:
                    return False
                if not third:
                    return True
                wide_l, wide_r = d["left23"] is not None, d["right23"] is not None
                clash = {"left23": wide_r or d["left3"] or d["mid3"], "right23": wide_l or d["right3"] or d["mid3"],
                         "left3": wide_l or wide_r and False, "right3": wide_r, "mid3": wide_l or wide_r}[side]
                return not clash
            idx = None
            cand = made.get(want) if page_wanted(r) else None
            if cand is None and page_wanted(r) and 1 <= want - 1 < len(pages):
                cand = pages[want - 1]
            if cand is not None and cand.get("duo") is not None and bool(cand["duo"].get("third")) == third \
                    and cand["duo"][side] is not None:
                cand["duo"][side + "_more"].append(r)
                continue
            first = next((i for i, pg in enumerate(pages) if pg is made.get(want)), None) if page_wanted(r) else None
            for k in range(first if first is not None else (max(1, want - 1) if page_wanted(r) else 1), len(pages)):
                if pages[k].get("duo") is not None and fits(pages[k]["duo"]):
                    idx = k
                    break
            if idx is None:
                new = {"row": None, "duo": {"third": third, **{k: None for k in ("left", "right", "left3", "mid3", "right3", "left23", "right23")},
                                                   **{k + "_more": [] for k in ("left", "right", "left3", "mid3", "right3", "left23", "right23")}}}
                at = max(1, min(want - 1, len(pages))) if page_wanted(r) else len(pages)
                while at < len(pages) and pages[at].get("duo") is not None:
                    at += 1
                pages.insert(at, new)
                idx = at
                if page_wanted(r):
                    made.setdefault(want, new)
                    new["pin"] = want
            pages[idx]["duo"][side] = r
            d = pages[idx]["duo"]
            pages[idx]["row"] = next(d[k] for k in ("left", "right", "left3", "mid3", "right3", "left23", "right23") if d[k])
        def keep_pinned_pages():
            """Halves and multi-page sections insert pages ahead of later ones; put each pinned page back on its chosen page number."""
            if not any(pg.get("pin") for pg in pages[1:]):
                return
            groups = []
            for k, pg in enumerate(pages):
                if groups and pg.get("part", 1) > 1 and pg.get("row") is not None and groups[-1][-1].get("row") is pg["row"] and not pg.get("duo"):
                    groups[-1].append(pg)
                else:
                    groups.append([pg])
            slots, rest = {}, []
            slots[0] = groups[0]
            for g in sorted(range(1, len(groups)), key=lambda i: (bool(groups[i][0].get("blank")), groups[i][0].get("pin") or 999, i)):
                grp, pin = groups[g], groups[g][0].get("pin")
                if grp[0].get("blank"):
                    if pin and (pin - 1) not in slots:
                        slots[pin - 1] = [grp[0]]
                    continue
                start = (pin or 0) - 1
                if pin and start >= 1 and all((start + j) not in slots for j in range(len(grp))) and \
                        all((start + j) not in slots for j in range(len(grp))):
                    for j in range(len(grp)):
                        slots[start + j] = [grp[j]]
                else:
                    rest.append(g)
            at = 0
            for g in rest:
                for pg in groups[g]:
                    while at in slots:
                        at += 1
                    slots[at] = [pg]
            pages[:] = [slots[k][0] for k in sorted(slots)]
        try:
            holes = [int(x) for x in (db().execute("SELECT blank_at FROM bulletins WHERE id=?", (bulletin_id,)).fetchone()["blank_at"] or "").split(",") if x.strip()]
        except Exception:
            holes = []
        wanted = {r["page_no"] for r in printed if r["page_no"]}
        for n in sorted(set(holes)):
            if n >= 2 and n not in wanted:
                pages.append({"row": None, "pin": n, "blank": True})
        keep_pinned_pages()

        def band_pos(pg):
            ps = [x["position"] for x in rows_on(pg) if x is not None]
            return (min(ps), max(ps)) if ps else None
        for r in sorted((r for r in printed if area_of(r) in ("below", "above") and r not in p1_all and r not in orphan_bands), key=lambda r: r["position"]):
            target = None
            if r["page_no"] and 2 <= r["page_no"] <= len(pages) and pages[r["page_no"] - 1].get("row") is not None:
                target = pages[r["page_no"] - 1]
                target.setdefault("more" if area_of(r) == "below" else "above", []).append(r)
                continue
            for k, pg in enumerate(pages):
                bp = band_pos(pg)
                if bp is None or pg.get("row") is None:
                    continue
                if area_of(r) == "below" and bp[1] < r["position"] and (target is None or bp[1] >= band_pos(target)[1]):
                    target = pg
                if area_of(r) == "above" and k > 0 and bp[0] > r["position"] and (target is None or bp[0] < band_pos(target)[0]):
                    target = pg
            if target is None:
                pages.append({"row": r, "part": 1, "of": 1})
            else:
                target.setdefault("more" if area_of(r) == "below" else "above", []).append(r)
        for pg in pages:
            if pg.get("more") and pg.get("row") is not None and not pg.get("duo"):
                share = [pg["row"]] + list(pg["more"])
                def text_height(r):
                    c = r["cols"] if r["cols"] in (2, 3) else 1
                    lines = sum(max(1, -(-len(x) * c // 100)) for x in (r["body"] or "").splitlines() if x.strip())
                    return 0.6 + 0.25 * (-(-lines // c) + (1 if c > 1 else 0))
                text_in = sum(text_height(r) for r in share)
                n_ph = sum(1 for r in share if r["image"]) or 1
                pg["photo_h"] = round(max(1.2, min(5.0, (9.5 - text_in) / n_ph)), 2)
        toc = []
        for n, p in enumerate(pages, 1):
            p["number"] = n
        seen = set()
        for p in pages:
            for r in rows_on(p):
                if r is not None and r["id"] not in seen:
                    seen.add(r["id"])
                    if r["no_toc"]:
                        continue
                    toc.append({"title": r["title"], "page": p["number"], "pos": r["position"]})
        toc.sort(key=lambda t: (t["page"], t["pos"]))
        cur = len(pages) + 1
        brow = db().execute("SELECT extra_pages, total_pages, ad_page_count FROM bulletins WHERE id=?", (bulletin_id,)).fetchone()
        extra = (brow["extra_pages"] or 0) if brow else 0
        total_want = (brow["total_pages"] or 0) if brow else 0
        ads = [dict(r) for r in db().execute("SELECT name,image,w,h,url FROM ads WHERE active=1 ORDER BY position, id").fetchall()]
        ad_gap, ad_pad = ad_layout()
        ad_pg = ad_pages(ads, (brow["ad_page_count"] or 0) if brow else 0, ad_gap, ad_pad)
        blanks = max(0, total_want - len(pages) - len(ad_pg)) if total_want else ((-(len(pages) + len(ad_pg))) % 4 if pages else 4) + extra
        for _ in range(blanks):
            pages.append({"row": None, "number": cur})
            cur += 1
        for k, pg_ in enumerate(ad_pg):
            pages.append({"row": None, "ad_rows": pg_["rows"], "ad_gap": pg_["gap"], "ad_top": pg_["top"], "ad_card_gap": ad_gap, "ad_first": k == 0, "ad_last": k == len(ad_pg) - 1, "number": cur})
            if k == 0:
                toc.append({"title": "Our Advertisers", "page": cur, "pos": 99999})
            cur += 1
        return pages, toc

    def page_spans(rows, bid=None):
        """Map section id -> (first, last, filled) page, from the real layout. Empty sections show where they would start."""
        placed = {}
        for p in layout(bid, include_pending=True)[0]:
            for r in rows_on(p):
                if r is not None:
                    lo, hi = placed.get(r["id"], (p["number"], p["number"]))
                    placed[r["id"]] = (min(lo, p["number"]), max(hi, p["number"]))
        spans, cur = {}, 1
        for r in rows:
            if r["id"] in placed:
                spans[r["id"]] = (*placed[r["id"]], True)
            else:
                spans[r["id"]] = (cur, cur + r["pages"] - 1, False)
        return spans

    PLAIN_SLUGS_EXCLUDED = ("giving", "mass-schedule", "mass-times", "mass-intentions", "calendar", "youth-calendar", "cover", "staff")

    # ---------- routes ----------
    @app.before_request
    def auto_history():
        if request.method != "POST" or not g.get("user") or not request.view_args:
            return
        bid = request.view_args.get("bid")
        path = request.path
        if not bid or "/history" in path or "/check" in path:
            return
        label = "Changed the bulletin"
        for key, text in HISTORY_LABELS.items():
            if f"/{key}" in path:
                label = text
                break
        try:
            snapshot_bulletin(bid, label)
        except Exception:
            pass

    @app.route("/login", methods=["GET", "POST"])
    def login():
        if request.method == "POST":
            who = request.form.get("username", "").strip().lower()
            u = db().execute("SELECT * FROM users WHERE username=? OR (email<>'' AND email=?)", (who, who)).fetchone()
            if u and u["active"] == 0:
                flash("This account is deactivated. Ask an administrator.", "error")
                return render_template("login.html")
            if u and check_password_hash(u["pw_hash"], request.form.get("password", "")):
                session.clear()
                session["uid"] = u["id"]
                session["csrf"] = secrets.token_hex(16)
                db().execute("UPDATE users SET last_login=?, last_seen=? WHERE id=?", ((datetime.utcnow().isoformat(),) * 2) + (u["id"],))
                db().commit()
                return redirect(safe_next(request.args.get("next")) or url_for("dashboard"))
            flash("Incorrect email or password.", "error")
        return render_template("login.html")

    # ---------- email, registration and password recovery ----------
    EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

    def send_mail(to, subject, body):
        """Send through the SMTP account set in environment variables; returns True if sent."""
        hook = app.config.get("MAIL_SENDER")
        if hook:
            return bool(hook(to, subject, body))
        host = os.environ.get("SMTP_HOST")
        app.config["MAIL_ERROR"] = ""
        if not host:
            app.config["MAIL_ERROR"] = "no email settings (SMTP_HOST) found on the server - check the WSGI file and press Reload"
            return False
        msg = EmailMessage()
        msg["To"], msg["Subject"] = to, subject
        msg["From"] = os.environ.get("MAIL_FROM") or os.environ.get("SMTP_USER", "")
        msg.set_content(body)
        try:
            with smtplib.SMTP(host, int(os.environ.get("SMTP_PORT", "587")), timeout=20) as sm:
                sm.starttls()
                if os.environ.get("SMTP_USER"):
                    sm.login(os.environ["SMTP_USER"], os.environ.get("SMTP_PASSWORD", ""))
                sm.send_message(msg)
            return True
        except (OSError, smtplib.SMTPException) as exc:
            app.config["MAIL_ERROR"] = "%s: %s" % (type(exc).__name__, str(exc)[:200])
            return False

    def make_token(uid, kind, hours):
        raw = secrets.token_urlsafe(32)
        conn = db()
        conn.execute("DELETE FROM tokens WHERE user_id=? AND kind=?", (uid, kind))
        conn.execute("INSERT INTO tokens(user_id,kind,token_hash,expires) VALUES(?,?,?,?)",
                     (uid, kind, hashlib.sha256(raw.encode()).hexdigest(), (datetime.utcnow() + timedelta(hours=hours)).isoformat()))
        conn.commit()
        return raw

    def find_token(raw, kinds):
        row = db().execute("SELECT * FROM tokens WHERE token_hash=?", (hashlib.sha256((raw or "").encode()).hexdigest(),)).fetchone()
        if not row or row["kind"] not in kinds or row["expires"] < datetime.utcnow().isoformat():
            return None
        return row

    def eastern(iso):
        if not iso:
            return "never"
        from zoneinfo import ZoneInfo
        try:
            t = datetime.fromisoformat(iso).replace(tzinfo=ZoneInfo("UTC")).astimezone(ZoneInfo("America/New_York"))
        except (ValueError, KeyError):
            return iso[:16].replace("T", " ")
        return t.strftime("%b %-d, %Y %-I:%M %p ") + t.tzname()

    def invite_user(u):
        """Email the registration link. Returns (sent, link)."""
        link = url_for("register", token=make_token(u["id"], "invite", 24 * 7), _external=True)
        body = ("Hello %s,\n\nAn account has been created for you on the St. Peter the Fisherman Bulletin Designer.\n"
                "Open this link to choose your password and finish registering (good for 7 days):\n\n%s\n\n"
                "Your sign-in is this email address.\n" % (u["display_name"] or u["username"], link))
        return send_mail(u["email"], "Finish setting up your bulletin account", body), link

    @app.route("/register/<token>", methods=["GET", "POST"])
    def register(token):
        t = find_token(token, ("invite",))
        if not t:
            flash("That registration link is invalid or has expired. Ask an administrator to send a new one.", "error")
            return redirect(url_for("login"))
        u = db().execute("SELECT * FROM users WHERE id=?", (t["user_id"],)).fetchone()
        if request.method == "POST":
            new = request.form.get("new", "")
            if len(new) < 10 or new != request.form.get("again", ""):
                flash("Type the same password twice, at least 10 characters.", "error")
            else:
                name = request.form.get("display_name", "").strip()
                db().execute("UPDATE users SET pw_hash=?, must_reset=0, display_name=? WHERE id=?",
                             (generate_password_hash(new), name or u["display_name"], u["id"]))
                db().execute("DELETE FROM tokens WHERE user_id=?", (u["id"],))
                db().commit()
                session.clear()
                session["uid"] = u["id"]
                session["csrf"] = secrets.token_hex(16)
                db().execute("UPDATE users SET last_login=?, last_seen=? WHERE id=?", ((datetime.utcnow().isoformat(),) * 2) + (u["id"],))
                db().commit()
                flash("Welcome! Your account is ready.", "ok")
                return redirect(url_for("dashboard"))
        return render_template("register.html", u=u, reset=False)

    @app.route("/forgot", methods=["GET", "POST"])
    def forgot():
        if request.method == "POST":
            who = request.form.get("email", "").strip().lower()
            u = db().execute("SELECT * FROM users WHERE email<>'' AND (email=? OR username=?)", (who, who)).fetchone()
            if u:
                link = url_for("reset_with_token", token=make_token(u["id"], "reset", 2), _external=True)
                send_mail(u["email"], "Reset your bulletin password",
                          "Open this link to choose a new password (good for 2 hours):\n\n%s\n\n"
                          "If you did not ask for this, ignore this email.\n" % link)
            flash("If that email has an account, a reset link has been sent. Check your inbox.", "ok")
            return redirect(url_for("login"))
        return render_template("forgot.html")

    @app.route("/reset/<token>", methods=["GET", "POST"])
    def reset_with_token(token):
        t = find_token(token, ("reset",))
        if not t:
            flash("That reset link is invalid or has expired. Ask for a new one.", "error")
            return redirect(url_for("forgot"))
        u = db().execute("SELECT * FROM users WHERE id=?", (t["user_id"],)).fetchone()
        if request.method == "POST":
            new = request.form.get("new", "")
            if len(new) < 10 or new != request.form.get("again", ""):
                flash("Type the same password twice, at least 10 characters.", "error")
            else:
                db().execute("UPDATE users SET pw_hash=?, must_reset=0 WHERE id=?", (generate_password_hash(new), u["id"]))
                db().execute("DELETE FROM tokens WHERE user_id=?", (u["id"],))
                db().commit()
                flash("Password changed. Please sign in.", "ok")
                return redirect(url_for("login"))
        return render_template("register.html", u=u, reset=True)

    @app.route("/logout", methods=["POST"])
    def logout():
        session.clear()
        return redirect(url_for("login"))

    @app.route("/password", methods=["GET", "POST"])
    @need()
    def password():
        if request.method == "POST":
            new = request.form.get("new", "")
            if len(new) < 10 or not check_password_hash(g.user["pw_hash"], request.form.get("current", "")):
                flash("Current password must be correct and the new one at least 10 characters.", "error")
            else:
                db().execute("UPDATE users SET pw_hash=?, must_reset=0 WHERE id=?",
                             (generate_password_hash(new), g.user["id"]))
                db().commit()
                flash("Password updated.", "ok")
                return redirect(url_for("dashboard"))
        return render_template("password.html")

    @app.route("/")
    @need()
    def dashboard():
        rows = db().execute("SELECT * FROM bulletins ORDER BY sunday DESC").fetchall()
        return render_template("dashboard.html", bulletins=rows, suggested=next_sunday().isoformat())

    @app.route("/bulletins/new", methods=["POST"])
    @need("admin", "editor")
    def new_bulletin():
        try:
            sunday = date.fromisoformat(request.form.get("sunday", "")).isoformat()
        except ValueError:
            flash("Enter a valid date.", "error")
            return redirect(url_for("dashboard"))
        conn = db()
        created = conn.execute("INSERT OR IGNORE INTO bulletins(sunday) VALUES(?)", (sunday,)).rowcount == 1
        b = conn.execute("SELECT id FROM bulletins WHERE sunday=?", (sunday,)).fetchone()
        prev = conn.execute("SELECT * FROM bulletins WHERE sunday<? AND id<>? ORDER BY sunday DESC LIMIT 1", (sunday, b["id"])).fetchone() if created else None
        if prev:
            keep = [c for c in BULLETIN_CARRY if c in prev.keys()]
            conn.execute("UPDATE bulletins SET " + ", ".join(c + "=?" for c in keep) + " WHERE id=?",
                         [prev[c] for c in keep] + [b["id"]])
        if created:
            conn.execute("UPDATE bulletins SET no_underline=1 WHERE id=?", (b["id"],))
        conn.commit()
        ensure_items(b["id"], sunday)
        return redirect(url_for("edit", bid=b["id"]))

    @app.route("/bulletins/<int:bid>/duplicate", methods=["POST"])
    @need("admin", "editor")
    def duplicate_bulletin(bid):
        src = get_bulletin(bid)
        try:
            sunday = date.fromisoformat(request.form.get("sunday", "")).isoformat()
        except ValueError:
            flash("Enter a valid date for the copy.", "error")
            return redirect(url_for("dashboard"))
        conn = db()
        if conn.execute("SELECT 1 FROM bulletins WHERE sunday=?", (sunday,)).fetchone():
            flash("A bulletin already exists for that date. Pick a different date.", "error")
            return redirect(url_for("dashboard"))
        cols = [c for c in src.keys() if c not in ("id", "sunday", "status", "is_template")]
        conn.execute("INSERT INTO bulletins(sunday, status, " + ", ".join(cols) + ") VALUES(?, 'draft', " + ", ".join("?" * len(cols)) + ")",
                     [sunday] + [src[c] for c in cols])
        new = conn.execute("SELECT id FROM bulletins WHERE sunday=?", (sunday,)).fetchone()["id"]
        conn.execute("UPDATE bulletins SET no_underline=1 WHERE id=?", (new,))
        smap = {}
        for sec in conn.execute("SELECT * FROM sections WHERE bulletin_id=?", (bid,)).fetchall():
            slug = "%s-copy%d" % (sec["slug"], new)
            cur = conn.execute("INSERT INTO sections(slug,title,weekly,pages,position,hint,bulletin_id) VALUES(?,?,?,?,?,?,?)",
                               (slug, sec["title"], sec["weekly"], sec["pages"], sec["position"], sec["hint"], new))
            smap[sec["id"]] = cur.lastrowid
        for it in conn.execute("SELECT * FROM items WHERE bulletin_id=?", (bid,)).fetchall():
            ic = [c for c in it.keys() if c not in ("id", "bulletin_id", "section_id")]
            conn.execute("INSERT INTO items(bulletin_id, section_id, " + ", ".join(ic) + ") VALUES(?, ?, " + ", ".join("?" * len(ic)) + ")",
                         [new, smap.get(it["section_id"], it["section_id"])] + [it[c] for c in ic])
        conn.commit()
        flash("Copied. This is a full copy of every section, photo and layout.", "ok")
        return redirect(url_for("edit", bid=new))

    @app.route("/bulletins/<int:bid>/template", methods=["POST"])
    @need("admin", "editor")
    def toggle_template(bid):
        b = get_bulletin(bid)
        db().execute("UPDATE bulletins SET is_template=? WHERE id=?", (0 if b["is_template"] else 1, bid))
        db().commit()
        return redirect(url_for("dashboard"))

    def get_bulletin(bid):
        b = db().execute("SELECT * FROM bulletins WHERE id=?", (bid,)).fetchone()
        if not b:
            abort(404)
        return b

    @app.route("/bulletins/<int:bid>")
    @need()
    def edit(bid):
        b = get_bulletin(bid)
        ensure_items(bid, b["sunday"])
        secs = db().execute("""SELECT s.id, s.slug, COALESCE(NULLIF(i.title, ''), s.title) AS title, s.weekly, s.hint, s.bulletin_id, i.position, i.pages,
            i.cols, i.area, i.page_no, i.font_pt, i.font_style, i.title_gap, i.item_gap, i.space_before, i.hide_title, i.para_gap, i.no_toc, i.vcenter, i.text_center, i.text_nudge, i.status, i.carried, i.body, i.image, i.updated_by
            FROM sections s JOIN items i ON i.section_id=s.id WHERE i.bulletin_id=? ORDER BY i.position, s.id""",
                            (bid,)).fetchall()
        removed = db().execute("""SELECT s.id, s.title FROM hidden_sections h JOIN sections s ON s.id=h.section_id WHERE h.bulletin_id=? ORDER BY s.position""", (bid,)).fetchall()
        pages, _ = layout(bid, include_pending=True)
        editable = {s["id"]: can_edit_section(s["id"]) for s in secs}
        spans = page_spans(secs, bid)
        why = {}
        for s_ in secs:
            want = s_["page_no"] or 0
            if want >= 2 and spans[s_["id"]][2] and spans[s_["id"]][0] != want:
                others = [o["title"] for o in secs if o["id"] != s_["id"] and spans[o["id"]][2] and spans[o["id"]][0] <= want <= spans[o["id"]][1]]
                why[s_["id"]] = (f"Page {want} is already used by " + " and ".join(others) + ". Give that section a different page, or set this one to the same page with 1 page each.") if others else f"Page {want} could not be used, so this went to page {spans[s_['id']][0]}."
        return render_template("edit.html", font_styles=all_fonts(), photos=photo_choices(), b=b, secs=secs, total=len(pages), editable=editable, removed=removed,
                               where=spans, why=why)

    @app.route("/bulletins/<int:bid>/section/<int:sid>", methods=["GET", "POST"])
    @need()
    def section(bid, sid):
        b = get_bulletin(bid)
        ensure_items(bid, b["sunday"])
        s = db().execute("SELECT * FROM sections WHERE id=?", (sid,)).fetchone() or abort(404)
        item = db().execute("SELECT * FROM items WHERE bulletin_id=? AND section_id=?", (bid, sid)).fetchone()
        if request.method == "POST":
            if not can_edit_section(sid) or b["status"] == "published":
                abort(403)
            body = request.form.get("body", "")
            if s["slug"] == "mass-schedule":
                lines = []
                head = request.form.get("header", "").replace("|", "/").replace("\n", " ").strip()
                if head:
                    lines.append("# " + head)
                rows = []
                for i in range(min(int(request.form.get("mass_count", 0)), 40)):
                    if request.form.get(f"m{i}_remove"):
                        continue
                    vals = form_vals("m", i, MASS_FIELDS)
                    if any(vals):
                        rows.append(vals)
                rows.sort(key=lambda v: iso_date(v[1]) or "9999")
                lines += [" | ".join(v) for v in rows]
                body = "\n".join(lines)
            elif s["slug"] == "mass-intentions":
                lines = []
                head = request.form.get("header", "").replace("|", "/").replace("\n", " ").strip()
                if head:
                    lines.append("# " + head)
                for gi in range(min(int(request.form.get("group_count", 0)), 100)):
                    gdate = long_date(request.form.get(f"g{gi}_date", ""))
                    for r in range(min(int(request.form.get(f"g{gi}_count", 0)), 100)):
                        vals = [request.form.get(f"g{gi}_{r}_{k}", "").replace("|", "/").strip() for k in ("time", "for_whom", "by_whom")]
                        if any(vals):
                            lines.append(" | ".join(["", gdate] + vals))
                body = "\n".join(lines)
            elif s["slug"] == "cover":
                clean = lambda k: request.form.get(k, "").replace("|", "/").replace("\n", " ").strip()
                lines = []
                cd, ct = long_date(request.form.get("cover_date", "")) or clean("cover_date_old"), clean("cover_title")
                if cd or ct:
                    lines.append(f"@date | {cd} | {ct}")
                db().execute("UPDATE bulletins SET hdr_date=?, theme=? WHERE id=?", (iso_date(cd), ct, bid))
                av = [clean("addr_" + k) for k in ADDR_FIELDS]
                if any(av):
                    lines.append("@address | " + " | ".join(av))
                for i in range(min(int(request.form.get("hours_count", 0)), 40)):
                    hv = [clean(f"h{i}_{k}") for k in ("days", "time", "notes")]
                    if any(hv):
                        lines.append("@hours | " + " | ".join(hv))
                extra = request.form.get("extra", "").strip()
                if extra:
                    lines.append(extra)
                body = "\n".join(lines)
            elif s["slug"] in ("calendar", "youth-calendar"):
                lines = []
                for i in range(min(int(request.form.get("event_count", 0)), 300)):
                    vals = [request.form.get(f"e{i}_{k}", "").replace("|", "/").strip() for k in ("date", "time", "event") + CAL_MIN]
                    vals[0] = long_date(vals[0]) or request.form.get(f"e{i}_old", "").replace("|", "/").strip()
                    vals[-1] = long_date(vals[-1]) if s["slug"] == "youth-calendar" else ""
                    if any(vals):
                        lines.append(" | ".join(vals).rstrip(" |"))
                body = "\n".join(lines)
            elif s["slug"] == "staff":
                lines = []
                for i in range(min(int(request.form.get("staff_count", 0)), 80)):
                    vals = [request.form.get(f"s{i}_{k}", "").replace("|", "/").strip() for k in STAFF_FIELDS]
                    if any(vals):
                        lines.append(" | ".join(vals))
                body = "\n".join(lines)
            elif s["slug"] == "giving":
                vals = {k: request.form.get("g_" + k, "").replace("\n", "\\n").replace("\r", "").strip() for k in GIVING_KEYS}
                vals["no_second"] = "1" if request.form.get("g_no_second") else ""
                for k in ("first_label", "second_label", "deposit_label"):
                    vals[k] = " ".join(vals[k].split())[:60]
                rows = []
                for i in range(min(int(request.form.get("g_extra_count", 0) or 0), 30)):
                    r3 = [" ".join(request.form.get(f"gx{i}_{k}", "").replace("|", "/").split())[:120] for k in ("label", "amount", "note")]
                    if (r3[0] or r3[1]) and not request.form.get(f"gx{i}_rm"):
                        rows.append(" | ".join(r3))
                vals["extras"] = "\\n".join(rows)
                body = "#giving\n" + "\n".join(f"{k}={vals[k]}" for k in GIVING_KEYS)
            elif s["slug"] == "mass-times":
                lines = []
                for i in range(min(int(request.form.get("times_count", 0)), 40)):
                    vals = [request.form.get(f"t{i}_{k}", "").replace("|", "/").strip() for k in TIMES_FIELDS]
                    if any(vals):
                        lines.append(" | ".join(vals))
                body = "\n".join(lines)
            doc = request.files.get("document")
            doc_added = False
            if doc and doc.filename:
                dext = doc.filename.rsplit(".", 1)[-1].lower()
                if s["slug"] in ("mass-schedule", "mass-times", "mass-intentions", "calendar", "youth-calendar", "cover", "staff", "giving"):
                    flash("Mass schedules use the form; paste details into the fields instead of uploading a document.", "error")
                    return redirect(request.url)
                try:
                    if dext == "docx":
                        text = docx_text(io.BytesIO(doc.stream.read(25 * 1024 * 1024)))
                    elif dext == "doc":
                        with tempfile.NamedTemporaryFile(suffix=".doc") as tmp:
                            doc.save(tmp.name)
                            text = doc_text(tmp.name)
                    elif dext in ("txt", "text", "md"):
                        raw = doc.stream.read(2 * 1024 * 1024)
                        try:
                            text = raw.decode("utf-8-sig").strip()
                        except UnicodeDecodeError:
                            text = raw.decode("cp1252", "replace").strip()
                    else:
                        raise ValueError("Documents must be .doc, .docx or .txt.")
                except (ValueError, zipfile.BadZipFile, KeyError, ElementTree.ParseError, subprocess.SubprocessError):
                    flash("Could not read that document. Please upload a valid .docx file (.docx, .doc or .txt).", "error")
                    return redirect(request.url)
                if not text.strip():
                    flash("No text was found in that document (text inside pictures cannot be read).", "error")
                    return redirect(request.url)
                body = (body.rstrip() + "\n\n" + text).strip() if body.strip() else text
                doc_added = True
            image = item["image"]
            f = request.files.get("image")
            if f and f.filename:
                ext = f.filename.rsplit(".", 1)[-1].lower()
                if ext not in IMAGE_EXT:
                    flash("Images must be PNG, JPG, GIF or WebP.", "error")
                    return redirect(request.url)
                image = f"{uuid.uuid4().hex}_{secure_filename(f.filename)}"
                f.save(os.path.join(app.config["UPLOAD_DIR"], image))
            lib = request.form.get("library_id")
            lib_photo = db().execute("SELECT * FROM photos WHERE id=?", (lib,)).fetchone() if lib else None
            if lib_photo and not (f and f.filename):
                image = lib_photo["filename"]
                if not request.form.get("alt", "").strip():
                    request.form = request.form.copy()
                    request.form["alt"] = lib_photo["alt"]
            if request.form.get("remove_image"):
                image = ""
            decorative = 1 if request.form.get("decorative") else 0
            alt = request.form.get("alt", "").strip()
            if request.form.get("ocr") and image:
                found = ocr_text(os.path.join(app.config["UPLOAD_DIR"], image))
                if found:
                    alt = found[:2000]
                    decorative = 0
                    flash("The text in the photo was read and put in the alternative text. Please check it for mistakes.", "ok")
                elif found is None:
                    flash("Text reading (OCR) is not available on this server yet. Please type the words from the photo into the alternative text.", "error")
                else:
                    flash("No text was found in the photo.", "ok")
            if image and not decorative and not alt:
                flash("Alternative text is required for images (or mark the image decorative).", "error")
                return redirect(request.url)
            gal = gallery_data(item["gallery"])
            keep = []
            for n, p in enumerate(gal["photos"]):
                if request.form.get(f"gp_rm_{n}"):
                    continue
                p["alt"] = request.form.get(f"gp_alt_{n}", p.get("alt", "")).strip()[:300]
                wrap = request.form.get(f"gp_wrap_{n}")
                p["wrap"] = wrap if wrap in ("left", "right") else ""
                p["w"] = "above" if (request.form.get(f"gp_w_{n}") == "above" or p["wrap"]) else "below"
                shape = request.form.get(f"gp_shape_{n}", "")
                p["shape"] = shape if shape in PHOTO_SHAPES else ""
                for k, lo, hi, dflt in (("z", 1, 4, 1), ("x", 0, 100, 50), ("y", 0, 100, 50)):
                    try:
                        p[k] = max(lo, min(hi, float(request.form.get(f"gp_{k}_{n}", p.get(k, dflt)))))
                    except ValueError:
                        p[k] = dflt
                sh = shadow_from_form("gp_sh_", f"_{n}")
                if sh:
                    p["sh"] = sh
                else:
                    p.pop("sh", None)
                for k in ("ar", "ar0", "crop", "shd"):
                    p.pop(k, None)
                keep.append(p)
            for n in range(4):
                gf = request.files.get(f"ng_file_{n}")
                glib = db().execute("SELECT * FROM photos WHERE id=?", (request.form.get(f"ng_lib_{n}"),)).fetchone() if request.form.get(f"ng_lib_{n}") else None
                galt = request.form.get(f"ng_alt_{n}", "").strip()[:300]
                fname = ""
                if gf and gf.filename:
                    if gf.filename.rsplit(".", 1)[-1].lower() not in IMAGE_EXT:
                        flash("Images must be PNG, JPG, GIF or WebP.", "error")
                        return redirect(request.url)
                    fname = f"{uuid.uuid4().hex}_{secure_filename(gf.filename)}"
                    gf.save(os.path.join(app.config["UPLOAD_DIR"], fname))
                elif glib:
                    fname = glib["filename"]
                    galt = galt or glib["alt"] or ""
                if fname:
                    if not galt:
                        flash("Please type a caption for each extra photo (it also describes the photo for screen readers).", "error")
                        return redirect(request.url)
                    keep.append({"f": fname, "alt": galt, "w": "above" if request.form.get(f"ng_w_{n}") == "above" else "below"})
            try:
                gal_h = max(0.5, min(6.0, float(request.form.get("gal_h") or gal["h"])))
            except ValueError:
                gal_h = 1.5
            try:
                gal_gap = max(0.0, min(1.0, float(request.form.get("gal_gap") or 0.15)))
            except ValueError:
                gal_gap = 0.15
            def _sp(key):
                try:
                    return max(0.0, min(1.0, float(request.form.get(key) or 0.06)))
                except ValueError:
                    return 0.06
            gal_align = request.form.get("gal_align") if request.form.get("gal_align") in ("left", "right") else "center"
            gal_json = json.dumps({"h": gal_h, "gap": gal_gap, "before": _sp("gal_before"), "after": _sp("gal_after"), "align": gal_align,
                                   "dir": "stack" if request.form.get("gal_dir") == "stack" else "row",
                                   "size": request.form.get("gal_size") if request.form.get("gal_size") in ("width", "area") else "height",
                                   "valign": request.form.get("gal_valign") if request.form.get("gal_valign") in ("middle", "bottom") else "top",
                                   "wrap": request.form.get("gal_wrap") if request.form.get("gal_wrap") in ("left", "right") else "", "photos": keep[:12]}) if keep else ""
            status = "approved" if g.user["role"] in ("admin", "editor") else "pending"
            def num(key, lo, hi, default):
                try:
                    return max(lo, min(hi, float(request.form.get(key, default))))
                except ValueError:
                    return default
            new_title = request.form.get("section_title", s["title"]).strip()[:120]
            db().execute("""UPDATE items SET body=?, image=?, alt=?, decorative=?, caption=?, source=?,
                crop_zoom=?, crop_x=?, crop_y=?, crop_h=?, crop_t=?, crop_b=?, crop_l=?, crop_r=?, title=?, title_gap=?, item_gap=?, img_pos=?, para_gap=?, hide_title=?, space_before=?, bullet_style=?, gallery=?, img_w=?, img_align=?, bg_color=?, img_shadow=?,
                status=?, carried=0, updated_by=? WHERE id=?""",
                         (body, image, alt, decorative,
                          request.form.get("caption", "").strip(), request.form.get("source", "").strip(),
                          num("crop_zoom", 1, 4, 1), num("crop_x", 0, 100, 50), num("crop_y", 0, 100, 50), num("crop_h", 0, 9, 0),
                          *edge_values(request.form),
                          "" if new_title == s["title"] else new_title,
                          gap_value(request.form.get("title_gap")), gap_value(request.form.get("item_gap")),
                          request.form.get("img_pos") if request.form.get("img_pos") in ("above", "inset-left", "inset-right", "cols-left", "cols-right", "beside-left", "beside-right", "only") else "",
                          gap_value(request.form.get("para_gap")),
                          1 if request.form.get("hide_title") else 0,
                          gap_value(request.form.get("space_before")),
                          request.form.get("bullet_style") if request.form.get("bullet_style") in BULLETS else "",
                          gal_json, num("img_w", 0, 100, 0),
                          request.form.get("img_align") if request.form.get("img_align") in ("left", "right", "center") else "",
                          (request.form.get("bg_color") or "").lower() if re.fullmatch(r"#[0-9a-fA-F]{6}", request.form.get("bg_color") or "") and request.form.get("use_bg") else "",
                          shadow_from_form("sh_"),
                          status, g.user["username"], item["id"]))
            db().execute("UPDATE items SET text_center=?, text_nudge=? WHERE id=?",
                         (1 if request.form.get("text_center") else 0, gap_value(request.form.get("text_nudge")) or 0, item["id"]))
            db().commit()
            if doc_added:
                flash("The document's text is now in the box below. Please review it, then press Save section when it looks right.", "ok")
                return redirect(url_for("section", bid=bid, sid=sid))
            flash("Saved." if status == "approved" else "Saved - waiting for editor approval.", "ok")
            return redirect(url_for("edit", bid=bid))
        rows = db().execute("""SELECT s.id, s.slug, i.pages, i.body, i.image FROM sections s JOIN items i ON i.section_id=s.id
            WHERE i.bulletin_id=? ORDER BY i.position, s.id""", (bid,)).fetchall()
        _aw = {"left": 3.5, "right": 3.5, "left23": 4.7, "right23": 4.7, "left3": 2.2, "mid3": 2.2, "right3": 2.2}.get(item["area"] or "", 7.2)
        _iw = int(item["img_w"] or 0)
        print_w = round(_aw * (min(_iw, 60) if (item["img_pos"] or "") in ("inset-left", "inset-right") and _iw else 45 if (item["img_pos"] or "") in ("inset-left", "inset-right") else (_iw or 100)) / 100, 2)
        return render_template("section.html", print_w=print_w, where=page_spans(rows, bid).get(sid), photos=photo_choices(), b=b, s=s, item=item, can=can_edit_section(sid), giving=giving_data(item["body"]), mass_header=mass_data(item["body"])[0], mass=mass_data(item["body"])[1] + [{}],
                               times=times_rows(item["body"]) + [{}],
                               staff=staff_rows(item["body"]) + [{}],
                               events=calendar_entries(item["body"]) + [dict(date="", time="", event="", lectors="", sacristans="", ministers="")] * 2,
                               cover=cover_prefilled(item["body"], b["sunday"]),
                               intent_header=intention_data(item["body"])[0],
                               groups=[dict(d, items=d["items"] + [{}]) for d in intention_data(item["body"])[1]] + [{"day": "", "date": "", "items": [{}, {}]}])

    COPY_COLS = ("bg_color", "img_align", "body", "image", "alt", "decorative", "caption", "source", "crop_zoom", "crop_x", "crop_y", "crop_h", "crop_t", "crop_b", "crop_l", "crop_r", "title", "title_gap", "item_gap", "img_pos", "para_gap", "hide_title", "space_before", "bullet_style", "gallery", "img_w", "img_shadow")

    @app.route("/bulletins/<int:bid>/section/<int:sid>/copy-next", methods=["POST"])
    @need()
    def copy_next(bid, sid):
        b = get_bulletin(bid)
        if not can_edit_section(sid):
            abort(403)
        conn = db()
        nxt = (date.fromisoformat(b["sunday"]) + timedelta(days=7)).isoformat()
        conn.execute("INSERT OR IGNORE INTO bulletins(sunday) VALUES(?)", (nxt,))
        conn.commit()
        nb = conn.execute("SELECT * FROM bulletins WHERE sunday=?", (nxt,)).fetchone()
        if nb["status"] == "published":
            flash(f"The {nxt} bulletin is already published.", "error")
            return redirect(url_for("section", bid=bid, sid=sid))
        ensure_items(nb["id"], nxt)
        src = conn.execute("SELECT * FROM items WHERE bulletin_id=? AND section_id=?", (bid, sid)).fetchone()
        sec = conn.execute("SELECT * FROM sections WHERE id=?", (sid,)).fetchone()
        if sec["bulletin_id"] == bid:
            slug = "%s-to%d" % (sec["slug"], nb["id"])
            dest = conn.execute("SELECT id FROM sections WHERE slug=?", (slug,)).fetchone()
            if dest:
                dest_id = dest["id"]
            else:
                dest_id = conn.execute("INSERT INTO sections(slug,title,weekly,pages,position,hint,bulletin_id) VALUES(?,?,?,?,?,?,?)",
                                       (slug, sec["title"], sec["weekly"], sec["pages"], sec["position"], sec["hint"], nb["id"])).lastrowid
                conn.execute("INSERT INTO items(bulletin_id,section_id,position,pages,cols) VALUES(?,?,?,1,1)",
                             (nb["id"], dest_id, src["position"]))
            conn.execute("UPDATE items SET position=?, pages=?, cols=?, " + ",".join(f"{c}=?" for c in COPY_COLS) +
                         ", status='approved', carried=0, updated_by=? WHERE bulletin_id=? AND section_id=?",
                         [src["position"], src["pages"], src["cols"]] + [src[c] for c in COPY_COLS] + [g.user["username"], nb["id"], dest_id])
            conn.commit()
            flash(f"Copied to the {nxt} bulletin.", "ok")
            return redirect(url_for("section", bid=bid, sid=sid))
        is_cover = conn.execute("SELECT slug FROM sections WHERE id=?", (sid,)).fetchone()["slug"] == "cover"
        conn.execute("UPDATE items SET " + ",".join(f"{c}=?" for c in COPY_COLS) +
                     ", status='approved', carried=0, updated_by=? WHERE bulletin_id=? AND section_id=?",
                     [(week_free(src[c]) if c == "body" and is_cover else src[c]) for c in COPY_COLS] + [g.user["username"], nb["id"], sid])
        conn.commit()
        flash(f"Copied to the {nxt} bulletin.", "ok")
        return redirect(url_for("section", bid=bid, sid=sid))

    @app.route("/bulletins/<int:bid>/section/<int:sid>/keep", methods=["POST"])
    @need("admin", "editor")
    def keep_section(bid, sid):
        get_bulletin(bid)
        conn = db()
        s = conn.execute("SELECT * FROM sections WHERE id=?", (sid,)).fetchone() or abort(404)
        conn.execute("UPDATE sections SET weekly=? WHERE id=?", (0 if s["weekly"] else 1, sid))
        conn.commit()
        flash("This section will now be kept for all bulletin editions." if s["weekly"]
              else "This section will now start blank each week.", "ok")
        return redirect(url_for("section", bid=bid, sid=sid))

    @app.route("/bulletins/<int:bid>/section/<int:sid>/delete", methods=["POST"])
    @need()
    def delete_content(bid, sid):
        b = get_bulletin(bid)
        if not can_edit_section(sid) or b["status"] == "published":
            abort(403)
        conn = db()
        conn.execute("UPDATE items SET body='', image='', alt='', decorative=0, caption='', source='', gallery='', img_w=0, img_shadow='', carried=0, updated_by=? "
                     "WHERE bulletin_id=? AND section_id=?", (g.user["username"], bid, sid))
        conn.commit()
        flash("Section content deleted.", "ok")
        return redirect(url_for("section", bid=bid, sid=sid))

    def footer_address(bid):
        row = db().execute("""SELECT i.body FROM items i JOIN sections s ON s.id=i.section_id
            WHERE i.bulletin_id=? AND s.slug='cover'""", (bid,)).fetchone()
        cv = cover_data(row["body"] if row else "")
        a = cv["addr"]
        line = ", ".join(x for x in (a["street"], a["city"]) if x)
        phone = a["phone"]
        if not line and not phone:
            line, phone = "31300 Overseas Hwy, Big Pine Key, FL 33043", "305-872-2537"
        h = cv["hours"][0] if cv["hours"] else None
        g.cover_dt = (cv["date"], cv["title"])
        hours = " ".join(x for x in (h["days"], h["time"]) if x) if h else ""
        return line, phone, hours or "Monday-Friday 9:30am-3:30pm"

    def open_bulletin(bid):
        b = get_bulletin(bid)
        if b["status"] == "published":
            abort(403)
        return b

    @app.route("/bulletins/<int:bid>/add-section", methods=["POST"])
    @need("admin", "editor")
    def add_section(bid):
        open_bulletin(bid)
        conn = db()
        title = request.form.get("title", "").strip() or "New section"
        last = conn.execute("SELECT COALESCE(MAX(position),0) m FROM items WHERE bulletin_id=?", (bid,)).fetchone()["m"]
        cur = conn.execute("""INSERT INTO sections(slug,title,weekly,pages,position,hint,bulletin_id)
            VALUES(?,?,1,1,?,?,?)""", (f"custom-{uuid.uuid4().hex[:8]}", title, last + 1,
                                      "Added for this week only", bid))
        conn.execute("INSERT INTO items(bulletin_id,section_id,position,pages,cols) VALUES(?,?,?,1,1)",
                     (bid, cur.lastrowid, last + 1))
        conn.commit()
        return redirect(url_for("section", bid=bid, sid=cur.lastrowid))

    @app.route("/bulletins/<int:bid>/history")
    @need("admin", "editor")
    def history_page(bid):
        b = get_bulletin(bid)
        rows = db().execute("SELECT id,label,created,by FROM history WHERE bulletin_id=? ORDER BY id DESC", (bid,)).fetchall()
        return render_template("history.html", b=b, rows=rows)

    @app.route("/bulletins/<int:bid>/history/<int:hid>/restore", methods=["POST"])
    @need("admin", "editor")
    def history_restore(bid, hid):
        open_bulletin(bid)
        conn = db()
        row = conn.execute("SELECT snapshot FROM history WHERE id=? AND bulletin_id=?", (hid, bid)).fetchone()
        if not row:
            abort(404)
        snapshot_bulletin(bid, "Before going back")
        snap = json.loads(row["snapshot"])
        keep = {"id", "sunday", "status"}
        cols = [k for k in snap["bulletin"] if k not in keep]
        if cols:
            conn.execute("UPDATE bulletins SET " + ",".join(f"{k}=?" for k in cols) + " WHERE id=?",
                         [snap["bulletin"][k] for k in cols] + [bid])
        conn.execute("DELETE FROM items WHERE bulletin_id=?", (bid,))
        have = {r["id"] for r in conn.execute("SELECT id FROM sections WHERE bulletin_id=?", (bid,))}
        want = {r["id"] for r in snap["sections"]}
        for sid in have - want:
            conn.execute("DELETE FROM sections WHERE id=?", (sid,))
        for sec in snap["sections"]:
            if sec["id"] not in have:
                conn.execute(f"INSERT INTO sections({','.join(sec)}) VALUES({','.join('?' * len(sec))})", list(sec.values()))
        for it in snap["items"]:
            conn.execute(f"INSERT INTO items({','.join(it)}) VALUES({','.join('?' * len(it))})", list(it.values()))
        conn.commit()
        flash("Went back to how the bulletin looked at that point.", "ok")
        return redirect(url_for("edit", bid=bid))

    @app.route("/bulletins/<int:bid>/reset-layout", methods=["POST"])
    @need("admin", "editor")
    def reset_layout(bid):
        open_bulletin(bid)
        conn = db()
        try:
            sid = int(request.form.get("sid") or 0)
        except ValueError:
            sid = 0
        sql = ("UPDATE items SET position=(SELECT position FROM sections WHERE id=items.section_id), "
               "pages=(SELECT pages FROM sections WHERE id=items.section_id), cols=1, area='', page_no=0, font_pt=0, "
               "font_style='', min_h=0, title_gap=0, item_gap=0, para_gap=0, img_pos='', hide_title=0, space_before=0, bullet_style='', "
               "crop_zoom=1, crop_x=50, crop_y=50, crop_h=0, crop_t=0, crop_b=0, crop_l=0, crop_r=0, img_align='', bg_color='' WHERE bulletin_id=?")
        args = [bid]
        if sid:
            sql += " AND section_id=?"
            args.append(sid)
        conn.execute(sql, args)
        conn.commit()
        flash("Layout reset to the standard settings. Your text and photos were not changed.", "ok")
        return redirect(url_for("section", bid=bid, sid=sid) if sid else url_for("edit", bid=bid))

    @app.route("/bulletins/<int:bid>/layout/<int:sid>", methods=["POST"])
    @need("admin", "editor")
    def configure(bid, sid):
        open_bulletin(bid)
        conn = db()
        rows = conn.execute("SELECT section_id FROM items WHERE bulletin_id=? ORDER BY position, section_id",
                            (bid,)).fetchall()
        order = [r["section_id"] for r in rows]
        if sid not in order:
            abort(404)
        move = request.form.get("move")
        i = order.index(sid)
        if move == "up" and i > 0:
            order[i - 1], order[i] = order[i], order[i - 1]
        elif move == "down" and i < len(order) - 1:
            order[i + 1], order[i] = order[i], order[i + 1]
        try:
            want = int(request.form.get("order") or 0)
        except ValueError:
            want = 0
        if want and not move:
            order.remove(sid)
            order.insert(max(0, min(len(order), want - 1)), sid)
        for n, section_id in enumerate(order):
            conn.execute("UPDATE items SET position=? WHERE bulletin_id=? AND section_id=?", (n, bid, section_id))
        if request.form.get("pages"):
            conn.execute("UPDATE items SET pages=?, cols=? WHERE bulletin_id=? AND section_id=?",
                         (max(1, min(8, int(request.form["pages"]))),
                          int(request.form["cols"]) if request.form.get("cols") in ("2", "3") else 1, bid, sid))
        if "area" in request.form:
            area = request.form["area"] if request.form["area"] in ("full", "left", "right", "below", "above", "left3", "mid3", "right3", "left23", "right23") else ""
            try:
                page_no = max(0, min(12, int(request.form.get("page_no") or 0)))
            except ValueError:
                page_no = 0
            try:
                font_pt = int(request.form.get("font_pt") or 0)
            except ValueError:
                font_pt = 0
            font_pt = font_pt if font_pt == 0 or 8 <= font_pt <= 16 else 0
            if request.form.get("toc_ui"):
                conn.execute("UPDATE items SET text_center=?, text_nudge=? WHERE bulletin_id=? AND section_id=?", (1 if request.form.get("text_center") else 0, gap_value(request.form.get("text_nudge")) or 0, bid, sid))
                conn.execute("UPDATE items SET vcenter=? WHERE bulletin_id=? AND section_id=?", (1 if request.form.get("vcenter") else 0, bid, sid))
            conn.execute("UPDATE items SET no_toc=?, space_before=?, para_gap=?, hide_title=? WHERE bulletin_id=? AND section_id=?",
                         ((0 if request.form.get("in_toc") else 1) if request.form.get("toc_ui") else (1 if request.form.get("no_toc") else 0), gap_value(request.form.get("space_before")), gap_value(request.form.get("para_gap")),
                          1 if request.form.get("hide_title") else 0, bid, sid))
            conn.execute("UPDATE items SET area=?, page_no=?, font_pt=?, font_style=?, title_gap=?, item_gap=? WHERE bulletin_id=? AND section_id=?",
                         (area, page_no, font_pt, request.form.get("font_style") if request.form.get("font_style") in all_fonts() else "",
                          gap_value(request.form.get("title_gap")), gap_value(request.form.get("item_gap")), bid, sid))
        if move == "remove":
            keep_page_blank(bid, sid)
            sec = conn.execute("SELECT bulletin_id FROM sections WHERE id=?", (sid,)).fetchone()
            if sec["bulletin_id"] == bid:
                conn.execute("DELETE FROM items WHERE bulletin_id=? AND section_id=?", (bid, sid))
                conn.execute("DELETE FROM sections WHERE id=?", (sid,))
            else:
                it = conn.execute("SELECT * FROM items WHERE bulletin_id=? AND section_id=?", (bid, sid)).fetchone()
                if it:
                    conn.execute("INSERT OR REPLACE INTO hidden_sections(bulletin_id, section_id, snap) VALUES(?,?,?)",
                                 (bid, sid, json.dumps(dict(it))))
                    conn.execute("DELETE FROM items WHERE bulletin_id=? AND section_id=?", (bid, sid))
        conn.commit()
        if not move and "area" in request.form:
            flash("Settings saved.", "ok")
        return redirect(url_for("edit", bid=bid))

    @app.route("/bulletins/<int:bid>/section/<int:sid>/restore", methods=["POST"])
    @need("admin", "editor")
    def restore_section(bid, sid):
        open_bulletin(bid)
        conn = db()
        h = conn.execute("SELECT snap FROM hidden_sections WHERE bulletin_id=? AND section_id=?", (bid, sid)).fetchone() or abort(404)
        snap = json.loads(h["snap"])
        snap.pop("id", None)
        cols = {r[1] for r in conn.execute("PRAGMA table_info(items)")}
        snap = {k: v for k, v in snap.items() if k in cols}
        snap["position"] = conn.execute("SELECT COALESCE(MAX(position),0)+1 m FROM items WHERE bulletin_id=?", (bid,)).fetchone()["m"]
        conn.execute("INSERT INTO items(" + ",".join(snap) + ") VALUES(" + ",".join("?" * len(snap)) + ")", list(snap.values()))
        conn.execute("DELETE FROM hidden_sections WHERE bulletin_id=? AND section_id=?", (bid, sid))
        conn.commit()
        flash("Section added back at the end of the list.", "ok")
        return redirect(url_for("edit", bid=bid))

    @app.route("/bulletins/<int:bid>/resize/<int:sid>", methods=["POST"])
    @need("admin")
    def resize(bid, sid):
        """Drag-to-resize: snap the dragged width to a layout width and store a minimum height (inches)."""
        open_bulletin(bid)
        conn = db()
        item = conn.execute("""SELECT i.area, s.slug FROM items i JOIN sections s ON s.id=i.section_id
                               WHERE i.bulletin_id=? AND i.section_id=?""", (bid, sid)).fetchone()
        if item is None:
            abort(404)
        try:
            frac = float(request.form.get("frac") or 0)
            height = float(request.form.get("height") or 0)
        except ValueError:
            abort(400)
        height = 0 if height < 0.5 else min(9.0, round(height, 2))
        cur = item["area"] or ("left" if item["slug"] == "mass-intentions" else "full")
        right = cur in ("right", "right3", "right23")
        if item["slug"] == "cover":
            area = cur
        elif frac >= 0.85:
            area = "full"
        elif frac >= 0.6:
            area = "right23" if right else "left23"
        elif frac >= 0.42:
            area = "right" if right else "left"
        else:
            area = "mid3" if cur == "mid3" else ("right3" if right else "left3")
        conn.execute("UPDATE items SET area=?, min_h=? WHERE bulletin_id=? AND section_id=?", (area, height, bid, sid))
        conn.commit()
        return jsonify(area=area, min_h=height)

    # ---------- AI assistant ----------
    def find_section(bid, name):
        name = (name or "").strip().lower()
        best = None
        for r in db().execute("""SELECT s.id, s.title FROM sections s JOIN items i ON i.section_id=s.id
                WHERE i.bulletin_id=?""", (bid,)):
            t = r["title"].lower()
            if t == name or (name and (name in t or t in name)):
                if best is None or t == name:
                    best = r
        return best

    def rule_based(bid, text):
        """Offline fallback understanding a few plain-English layout requests."""
        t = text.strip()
        m = re.search(r"(?:set|change|make) (?:the )?(?:header|banner|tagline|theme)(?: text| line)? (?:to|as|say) [\"']?(.+?)[\"']?$", t, re.I)
        if m:
            return [{"action": "set_header", "text": m.group(1)}]
        m = re.search(r"add (?:a |an )?(?:new )?section (?:called |named |titled )?[\"']?(.+?)[\"']?$", t, re.I)
        if m:
            return [{"action": "add_section", "title": m.group(1)}]
        m = re.search(r"move (.+?) (up|down)", t, re.I)
        if m:
            return [{"action": "move", "section": m.group(1), "direction": m.group(2).lower()}]
        m = re.search(r"(?:make|set) (.+?) (?:to |be )?(\d) (column|page)s?", t, re.I)
        if m:
            key = "cols" if m.group(3).lower() == "column" else "pages"
            return [{"action": "layout", "section": m.group(1), key: int(m.group(2))}]
        m = re.search(r"(?:remove|delete) (?:the )?(?:section )?(.+)$", t, re.I)
        if m:
            return [{"action": "remove", "section": m.group(1)}]
        return []

    def llm_actions(bid, text):
        url = os.environ.get("BULLETIN_AI_URL")
        key = os.environ.get("BULLETIN_AI_KEY")
        if not (url and key):
            return None
        titles = [r["title"] for r in db().execute(
            "SELECT s.title FROM sections s JOIN items i ON i.section_id=s.id WHERE i.bulletin_id=?", (bid,))]
        system = ("You edit a parish bulletin layout. Reply with ONLY a JSON array of actions. Allowed: "
                  '{"action":"add_section","title":str,"text":str?}, {"action":"move","section":str,"direction":"up|down"}, '
                  '{"action":"layout","section":str,"pages":1-8?,"cols":1|2|3?}, {"action":"set_text","section":str,"text":str}, '
                  '{"action":"remove","section":str}, {"action":"set_header","text":str} (the tagline in the banner header). Never invent facts, names, dates or Mass intentions. '
                  "Existing sections: " + json.dumps(titles))
        payload = json.dumps({"model": os.environ.get("BULLETIN_AI_MODEL", "gpt-4o-mini"),
                              "messages": [{"role": "system", "content": system},
                                           {"role": "user", "content": text}]}).encode()
        req = urllib.request.Request(url, payload, {"Content-Type": "application/json",
                                                    "Authorization": f"Bearer {key}"})
        try:
            with urllib.request.urlopen(req, timeout=30) as resp:
                content = json.load(resp)["choices"][0]["message"]["content"]
            data = json.loads(re.search(r"\[.*\]", content, re.S).group(0))
            return data if isinstance(data, list) else []
        except Exception:
            return None

    def describe(bid, a):
        sec = find_section(bid, a.get("section")) if a.get("section") else None
        kind = a.get("action")
        if kind == "set_header":
            return f'Set the banner header tagline to "{str(a.get("text", ""))[:120]}"'
        if kind == "add_section":
            return f'Add section "{a.get("title", "New section")}"'
        if not sec:
            return None
        if kind == "move" and a.get("direction") in ("up", "down"):
            return f'Move "{sec["title"]}" {a["direction"]}'
        if kind == "layout":
            return f'Set "{sec["title"]}" to ' + ", ".join(f"{k} = {a[k]}" for k in ("pages", "cols") if k in a)
        if kind == "set_text":
            return f'Replace text of "{sec["title"]}" (needs your review)'
        if kind == "remove":
            return f'Remove "{sec["title"]}" (only sections added for this week can be removed)'
        return None

    def apply_action(bid, a):
        conn = db()
        kind = a.get("action")
        if kind == "set_header":
            conn.execute("UPDATE bulletins SET theme=? WHERE id=?", (str(a.get("text", ""))[:120], bid))
            return
        if kind == "add_section":
            title = str(a.get("title", "New section"))[:80]
            last = conn.execute("SELECT COALESCE(MAX(position),0) m FROM items WHERE bulletin_id=?", (bid,)).fetchone()["m"]
            cur = conn.execute("""INSERT INTO sections(slug,title,weekly,pages,position,hint,bulletin_id)
                VALUES(?,?,1,1,?,?,?)""", (f"custom-{uuid.uuid4().hex[:8]}", title, last + 1, "Added for this week only", bid))
            conn.execute("INSERT INTO items(bulletin_id,section_id,body,position,pages,cols,updated_by) VALUES(?,?,?,?,1,1,?)",
                         (bid, cur.lastrowid, str(a.get("text", "")), last + 1, "AI assistant"))
            return
        sec = find_section(bid, a.get("section"))
        if not sec:
            return
        sid = sec["id"]
        if kind == "move":
            order = [r["section_id"] for r in conn.execute(
                "SELECT section_id FROM items WHERE bulletin_id=? ORDER BY position, section_id", (bid,))]
            i = order.index(sid)
            j = i - 1 if a.get("direction") == "up" else i + 1
            if 0 <= j < len(order):
                order[i], order[j] = order[j], order[i]
            for n, x in enumerate(order):
                conn.execute("UPDATE items SET position=? WHERE bulletin_id=? AND section_id=?", (n, bid, x))
        elif kind == "layout":
            if "pages" in a:
                conn.execute("UPDATE items SET pages=? WHERE bulletin_id=? AND section_id=?",
                             (max(1, min(8, int(a["pages"]))), bid, sid))
            if "cols" in a:
                conn.execute("UPDATE items SET cols=? WHERE bulletin_id=? AND section_id=?",
                             (int(a["cols"]) if int(a["cols"]) in (2, 3) else 1, bid, sid))
        elif kind == "set_text":
            conn.execute("UPDATE items SET body=?, status='approved', carried=0, updated_by=? WHERE bulletin_id=? AND section_id=?",
                         (str(a.get("text", "")), "AI assistant", bid, sid))
        elif kind == "remove":
            keep_page_blank(bid, sid)
            row = conn.execute("SELECT bulletin_id FROM sections WHERE id=?", (sid,)).fetchone()
            if row["bulletin_id"] == bid:
                conn.execute("DELETE FROM items WHERE bulletin_id=? AND section_id=?", (bid, sid))
                conn.execute("DELETE FROM sections WHERE id=?", (sid,))

    @app.route("/bulletins/<int:bid>/assistant", methods=["POST"])
    @need("admin", "editor")
    def assistant(bid):
        open_bulletin(bid)
        if request.form.get("confirm"):
            for a in session.pop(f"ai{bid}", []):
                apply_action(bid, a)
            db().commit()
            flash("Assistant changes applied.", "ok")
            return redirect(url_for("edit", bid=bid))
        text = request.form.get("request", "").strip()[:1000]
        actions = llm_actions(bid, text)
        if actions is None:
            actions = rule_based(bid, text)
        actions = [a for a in actions if isinstance(a, dict) and describe(bid, a)][:10]
        if not actions:
            flash("I could not turn that into a change. Try: set header to Join us for Fall Festival, add section Pumpkin Patch, move Calendar down, "
                  "make Food Pantry 2 columns, or set a section to 2 pages.", "error")
            return redirect(url_for("edit", bid=bid))
        session[f"ai{bid}"] = actions
        return render_template("assistant.html", b=get_bulletin(bid),
                               steps=[describe(bid, a) for a in actions], request_text=text)

    @app.route("/bulletins/<int:bid>/approve/<int:sid>", methods=["POST"])
    @need("admin", "editor")
    def approve(bid, sid):
        db().execute("UPDATE items SET status='approved' WHERE bulletin_id=? AND section_id=?", (bid, sid))
        db().commit()
        return redirect(url_for("edit", bid=bid))

    @app.route("/bulletins/<int:bid>/footer", methods=["POST"])
    @need("admin", "editor")
    def set_footer(bid):
        open_bulletin(bid)
        text = " ".join(request.form.get("footer", "").split())[:160]
        db().execute("UPDATE bulletins SET footer=? WHERE id=?", (text, bid))
        db().commit()
        flash("Footer saved.", "ok")
        return redirect(url_for("edit", bid=bid))

    @app.route("/bulletins/<int:bid>/suggestions", methods=["GET", "POST"])
    @need()
    def suggestions(bid):
        b = get_bulletin(bid)
        days = liturgy.week(date.fromisoformat(b["sunday"]))
        targets = db().execute("""SELECT s.id, s.title FROM sections s JOIN items i ON i.section_id=s.id
            WHERE i.bulletin_id=? AND s.slug NOT IN (%s) ORDER BY i.position""" % ",".join("?" * len(PLAIN_SLUGS_EXCLUDED)),
                               (bid,) + PLAIN_SLUGS_EXCLUDED).fetchall()
        cals = db().execute("""SELECT s.id, s.title FROM sections s JOIN items i ON i.section_id=s.id
            WHERE i.bulletin_id=? AND s.slug IN ('calendar','youth-calendar') ORDER BY i.position""", (bid,)).fetchall()
        pasted = [" ".join(l.split()) for l in request.form.get("pasted", "").splitlines() if l.strip()][:200]
        return render_template("suggestions.html", b=b, days=days, targets=targets, cals=cals, pasted=pasted,
                               can=g.user["role"] in ("admin", "editor") and b["status"] != "published")

    @app.route("/bulletins/<int:bid>/suggestions/apply", methods=["POST"])
    @need("admin", "editor")
    def apply_suggestion(bid):
        open_bulletin(bid)
        lines = [" ".join(t.split())[:400] for t in request.form.getlist("text") if t.strip()]
        text = "\n".join(lines)
        target = request.form.get("target", "")
        if target.startswith("cal:") and lines:
            sid = target[4:]
            sec = db().execute("SELECT s.id, s.title FROM sections s JOIN items i ON i.section_id=s.id WHERE i.bulletin_id=? "
                               "AND s.id=? AND s.slug IN ('calendar','youth-calendar')", (bid, sid if sid.isdigit() else 0)).fetchone()
            if not sec:
                flash("Choose a calendar section.", "error")
                return redirect(url_for("suggestions", bid=bid))
            when = long_date(request.form.get("when", ""))
            tm = request.form.get("when_time", "").replace("|", "/").strip()[:30]
            add = "\n".join(f"{when} | {tm} | {l.replace('|', '/')}" for l in lines)
            ensure_items(bid, get_bulletin(bid)["sunday"])
            db().execute("UPDATE items SET body=CASE WHEN TRIM(body)='' THEN ? ELSE body || char(10) || ? END, carried=0, updated_by=? "
                         "WHERE bulletin_id=? AND section_id=?", (add, add, g.user["username"], bid, sec["id"]))
            db().commit()
            flash(f"Added {len(lines)} event(s) to {sec['title']}. Review them there.", "ok")
            return redirect(url_for("suggestions", bid=bid))
        if not text:
            flash("Nothing to add.", "error")
        elif target == "title":
            db().execute("UPDATE bulletins SET theme=? WHERE id=?", (" \u00b7 ".join(lines)[:120], bid))
            db().commit()
            flash("Added as the header title.", "ok")
        else:
            sec = db().execute("SELECT s.id, s.slug, s.title FROM sections s JOIN items i ON i.section_id=s.id "
                               "WHERE i.bulletin_id=? AND s.id=?", (bid, target if target.isdigit() else 0)).fetchone()
            if not sec or sec["slug"] in PLAIN_SLUGS_EXCLUDED:
                flash("Choose a text section to add to.", "error")
            else:
                ensure_items(bid, get_bulletin(bid)["sunday"])
                db().execute("UPDATE items SET body=CASE WHEN TRIM(body)='' THEN ? ELSE body || char(10) || char(10) || ? END, "
                             "carried=0, updated_by=? WHERE bulletin_id=? AND section_id=?",
                             (text, text, g.user["username"], bid, sec["id"]))
                db().commit()
                flash(f"Added to {sec['title']}. Review and edit it there.", "ok")
        return redirect(url_for("suggestions", bid=bid))

    YOUTH_TITLES = ["Special Event", "Edge Night", "Rock the Universe", "Ash Wednesday", "End of Year Party"]

    def youth_titles():
        return YOUTH_TITLES + [r["title"] for r in db().execute("SELECT title FROM event_titles ORDER BY title") if r["title"] not in YOUTH_TITLES]

    app.jinja_env.globals["youth_titles"] = youth_titles

    @app.route("/bulletins/<int:bid>/section/<int:sid>/event-titles", methods=["POST"])
    @need("admin", "editor", "contributor")
    def event_titles(bid, sid):
        title = " ".join((request.form.get("title") or "").split())[:80]
        if request.form.get("action") == "remove":
            db().execute("DELETE FROM event_titles WHERE title=?", (title,))
        elif title and title not in YOUTH_TITLES:
            db().execute("INSERT OR IGNORE INTO event_titles(title) VALUES(?)", (title,))
        db().commit()
        flash("Event list updated. Save your section first if you have unsaved changes.", "ok")
        return redirect(url_for("section", bid=bid, sid=sid))

    @app.route("/bulletins/<int:bid>/toc-settings", methods=["POST"])
    @need("admin", "editor")
    def toc_settings(bid):
        open_bulletin(bid)
        try:
            start = max(1, min(12, int(request.form.get("toc_start") or 2)))
            pt = int(request.form.get("toc_pt") or 0)
        except ValueError:
            start, pt = 2, 0
        pt = pt if pt == 0 or 8 <= pt <= 16 else 0
        style = request.form.get("toc_style")
        db().execute("UPDATE bulletins SET toc_title=?, toc_start=?, toc_show=?, toc_pt=?, toc_style=?, toc_place=?, toc_cols=?, toc_title_gap=?, toc_item_gap=?, toc_before=? WHERE id=?",
                     (request.form.get("toc_title", "").strip()[:60], start, 1 if request.form.get("toc_show") else 0,
                      pt, style if style in all_fonts() else "",
                      request.form.get("toc_place") if request.form.get("toc_place") in ("left", "right", "full") else "left",
                      request.form.get("toc_cols") if request.form.get("toc_cols") in ("1", "2", "3", "4") else 1,
                      gap_value(request.form.get("toc_title_gap")), gap_value(request.form.get("toc_item_gap")), gap_value(request.form.get("toc_before")), bid))
        db().commit()
        flash("Contents settings saved.", "ok")
        return redirect(url_for("edit", bid=bid))

    @app.route("/bulletins/<int:bid>/no-stragglers", methods=["POST"])
    @need("admin", "editor")
    def no_stragglers(bid):
        open_bulletin(bid)
        db().execute("UPDATE bulletins SET no_stragglers=? WHERE id=?", (1 if request.form.get("no_stragglers") else 0, bid))
        db().commit()
        flash("Straggler setting saved.", "ok")
        return redirect(url_for("edit", bid=bid))

    @app.route("/bulletins/<int:bid>/underlines", methods=["POST"])
    @need("admin", "editor")
    def underlines(bid):
        open_bulletin(bid)
        db().execute("UPDATE bulletins SET no_underline=? WHERE id=?", (1 if request.form.get("no_underline") else 0, bid))
        db().commit()
        flash("Underline setting saved.", "ok")
        return redirect(url_for("edit", bid=bid))

    @app.route("/bulletins/<int:bid>/extra-pages", methods=["POST"])
    @need("admin", "editor")
    def extra_pages(bid):
        open_bulletin(bid)
        cur = db().execute("SELECT extra_pages FROM bulletins WHERE id=?", (bid,)).fetchone()["extra_pages"] or 0
        new = cur + 4 if request.form.get("do") == "add" else max(0, cur - 4)
        db().execute("UPDATE bulletins SET extra_pages=? WHERE id=?", (new, bid))
        db().commit()
        flash(f"{new} extra pages added." if new else "Extra pages removed.", "ok")
        return redirect(url_for("edit", bid=bid))

    @app.route("/bulletins/<int:bid>/clear-blank-pages", methods=["POST"])
    @need("admin", "editor")
    def clear_blank_pages(bid):
        open_bulletin(bid)
        db().execute("UPDATE bulletins SET blank_at='' WHERE id=?", (bid,))
        db().commit()
        flash("The kept blank pages were removed.", "ok")
        return redirect(url_for("edit", bid=bid))

    @app.route("/bulletins/<int:bid>/total-pages", methods=["POST"])
    @need("admin", "editor")
    def total_pages(bid):
        open_bulletin(bid)
        try:
            n = max(0, min(60, int(request.form.get("total_pages") or 0)))
        except ValueError:
            n = 0
        db().execute("UPDATE bulletins SET total_pages=? WHERE id=?", (n, bid))
        db().commit()
        flash(f"The bulletin will have {n} pages." if n else "The app will fill the bulletin to a multiple of 4 pages.", "ok")
        return redirect(url_for("edit", bid=bid))

    @app.route("/bulletins/<int:bid>/ad-pages", methods=["POST"])
    @need("admin", "editor")
    def ad_page_setting(bid):
        open_bulletin(bid)
        try:
            n = max(0, min(8, int(request.form.get("ad_page_count") or 0)))
        except ValueError:
            n = 0
        db().execute("UPDATE bulletins SET ad_page_count=? WHERE id=?", (n, bid))
        db().commit()
        flash(f"Advertisements will be spread over {n} page(s)." if n else "Advertisements will use the automatic number of pages.", "ok")
        return redirect(url_for("edit", bid=bid))

    @app.route("/bulletins/<int:bid>/header-text", methods=["POST"])
    @need("admin", "editor")
    def set_header_text(bid):
        open_bulletin(bid)
        d = request.form.get("hdr_date", "").strip()
        d = d if long_date(d) else ""
        title = " ".join(request.form.get("theme", "").split())[:120]
        db().execute("UPDATE bulletins SET hdr_date=?, theme=? WHERE id=?", (d, title, bid))
        cov = db().execute("""SELECT i.id, i.body FROM items i JOIN sections s ON s.id=i.section_id
            WHERE i.bulletin_id=? AND s.slug='cover'""", (bid,)).fetchone()
        if cov:
            keep = [l for l in (cov["body"] or "").splitlines() if not l.strip().lower().startswith("@date")]
            ct = title.replace("|", "/")
            new = ([f"@date | {long_date(d)} | {ct}"] if (d or ct) else []) + keep
            db().execute("UPDATE items SET body=? WHERE id=?", ("\n".join(new), cov["id"]))
        db().commit()
        flash("Header date and title saved.", "ok")
        return redirect(url_for("edit", bid=bid))

    @app.route("/bulletins/<int:bid>/header-photo", methods=["POST"])
    @need("admin", "editor")
    def set_header_photo(bid):
        b = open_bulletin(bid)
        image, alt = b["hdr_image"], request.form.get("hdr_alt", "").strip()[:300] or b["hdr_alt"]
        f = request.files.get("image")
        lib = request.form.get("library_id")
        if request.form.get("reset"):
            image, alt = "", ""
        elif f and f.filename:
            ext = f.filename.rsplit(".", 1)[-1].lower()
            if ext not in IMAGE_EXT:
                flash("Images must be PNG, JPG, GIF or WebP.", "error")
                return redirect(url_for("edit", bid=bid))
            if not request.form.get("hdr_alt", "").strip():
                flash("Alternative text is required for the header photo.", "error")
                return redirect(url_for("edit", bid=bid))
            image = f"{uuid.uuid4().hex}_{secure_filename(f.filename)}"
            f.save(os.path.join(app.config["UPLOAD_DIR"], image))
        elif lib:
            photo = db().execute("SELECT * FROM photos WHERE id=?", (lib,)).fetchone()
            if photo:
                image, alt = photo["filename"], request.form.get("hdr_alt", "").strip()[:300] or photo["alt"] or "Parish photo"
        image2, alt2 = b["hdr2_image"] or "", request.form.get("hdr2_alt", "").strip()[:300] or b["hdr2_alt"] or ""
        f2 = request.files.get("image2")
        lib2 = request.form.get("library_id2")
        if request.form.get("reset2"):
            image2, alt2 = "", ""
        elif f2 and f2.filename:
            ext = f2.filename.rsplit(".", 1)[-1].lower()
            if ext not in IMAGE_EXT:
                flash("Images must be PNG, JPG, GIF or WebP.", "error")
                return redirect(url_for("edit", bid=bid))
            if not request.form.get("hdr2_alt", "").strip():
                flash("Alternative text is required for the second header photo.", "error")
                return redirect(url_for("edit", bid=bid))
            image2 = f"{uuid.uuid4().hex}_{secure_filename(f2.filename)}"
            f2.save(os.path.join(app.config["UPLOAD_DIR"], image2))
        elif lib2:
            photo2 = db().execute("SELECT * FROM photos WHERE id=?", (lib2,)).fetchone()
            if photo2:
                image2, alt2 = photo2["filename"], request.form.get("hdr2_alt", "").strip()[:300] or photo2["alt"] or "Parish photo"
        def num(key, lo, hi, default):
            try:
                return max(lo, min(hi, float(request.form.get(key, default))))
            except ValueError:
                return default
        db().execute("UPDATE bulletins SET hdr2_image=?, hdr2_alt=?, hdr2_zoom=?, hdr2_x=?, hdr2_y=?, hdr_split=? WHERE id=?",
                     (image2, alt2, num("hdr2_zoom", 1, 4, 1), num("hdr2_x", 0, 100, 50), num("hdr2_y", 0, 100, 50),
                      num("hdr_split", 30, 70, 62), bid))
        db().execute("UPDATE bulletins SET hdr_image=?, hdr_alt=?, hdr_h=?, hdr_zoom=?, hdr_x=?, hdr_y=?, hdr_style=? WHERE id=?",
                     (image, alt, num("hdr_h", 0.5, 5, 2.4), num("hdr_zoom", 1, 4, 1), num("hdr_x", 0, 100, 50),
                      num("hdr_y", 0, 100, 50),
                      style if (style := request.form.get("hdr_style")) in ("waves", "swell", "gentle", "straight", "allsouls") else "waves", bid))
        db().commit()
        flash("Header photo saved.", "ok")
        return redirect(url_for("edit", bid=bid))

    @app.route("/bulletins/<int:bid>/publish", methods=["POST"])
    @need("admin", "editor")
    def publish(bid):
        status = "draft" if request.form.get("unpublish") else "published"
        db().execute("UPDATE bulletins SET status=? WHERE id=?", (status, bid))
        db().commit()
        return redirect(url_for("edit", bid=bid))

    @app.route("/bulletins/<int:bid>/preview")
    @need()
    def preview(bid):
        b = get_bulletin(bid)
        pages, toc = layout(bid, include_pending=g.user["role"] != "viewer" or b["status"] == "published")
        apply_fit_memory(bid, pages)
        return sup_ordinals(render_template("bulletin.html", b=b, pages=pages, toc=toc, foot_addr=footer_address(bid)[0], foot_phone=footer_address(bid)[1], foot_hours=footer_address(bid)[2], cover_dt=g.cover_dt))

    def save_ad_file(f):
        from PIL import Image, ImageOps
        ext = f.filename.rsplit(".", 1)[-1].lower() if f.filename and "." in f.filename else ""
        base = f"{uuid.uuid4().hex}"
        if ext == "pdf":
            try:
                import fitz
                doc = fitz.open(stream=f.read(), filetype="pdf")
                pix = doc[0].get_pixmap(dpi=300)
                name = base + ".png"
                pix.save(os.path.join(app.config["UPLOAD_DIR"], name))
                return name, pix.width, pix.height
            except Exception:
                return None
        if ext not in IMAGE_EXT:
            return None
        try:
            im = ImageOps.exif_transpose(Image.open(f.stream))
            png = im.mode in ("RGBA", "LA", "P")
            name = base + (".png" if png else ".jpg")
            im.save(os.path.join(app.config["UPLOAD_DIR"], name), **({} if png else {"quality": 92}))
            return name, im.width, im.height
        except Exception:
            return None

    @app.route("/ads", methods=["GET", "POST"])
    @need()
    @need_area("ads")
    def ads():
        conn = db()
        if request.method == "POST":
            if g.user["role"] not in ("admin", "editor", "contributor"):
                abort(403)
            live = 1 if g.user["role"] in ("admin", "editor") else 0
            added, bad = 0, 0
            for f in request.files.getlist("cards"):
                if not f or not f.filename:
                    continue
                res = save_ad_file(f)
                if not res:
                    bad += 1
                    continue
                nxt = conn.execute("SELECT COALESCE(MAX(position),0)+1 FROM ads").fetchone()[0]
                title = os.path.splitext(secure_filename(f.filename))[0].replace("_", " ").replace("-", " ").strip()[:80]
                conn.execute("INSERT INTO ads(name,image,w,h,position,active) VALUES(?,?,?,?,?,?)", (title or "Advertiser", res[0], res[1], res[2], nxt, live))
                added += 1
            conn.commit()
            flash(f"Added {added} card(s)." + (" An editor must tick Print before they appear in the bulletin." if added and not live else "") + (f" {bad} file(s) could not be read (use JPG, PNG, GIF, WebP or PDF)." if bad else ""), "error" if bad and not added else "ok")
            return redirect(url_for("ads"))
        rows = conn.execute("SELECT * FROM ads ORDER BY position, id").fetchall()
        pg = ad_pages([dict(r) for r in rows if r["active"]], 0, *ad_layout())
        return render_template("ads.html", ads=rows, ad_page_count=len(pg))

    @app.route("/ads/save", methods=["POST"])
    @need("admin", "editor")
    def ads_save():
        conn = db()
        for r in conn.execute("SELECT id FROM ads").fetchall():
            i = r["id"]
            try:
                pos = int(request.form.get(f"pos_{i}", "0"))
            except ValueError:
                pos = 0
            conn.execute("UPDATE ads SET name=?, position=?, active=?, url=? WHERE id=?",
                         (request.form.get(f"name_{i}", "").strip()[:80] or "Advertiser", pos, 1 if request.form.get(f"active_{i}") else 0,
                          clean_ad_url(request.form.get(f"url_{i}", "")), i))
        conn.commit()
        flash("Advertisements saved.", "ok")
        return redirect(url_for("ads"))

    @app.route("/ads/shadow", methods=["POST"])
    @need("admin", "editor")
    def ads_shadow():
        conn = db()
        raw = shadow_from_form("ad", "") or "off"
        conn.execute("INSERT OR REPLACE INTO meta(k,v) VALUES('ad_shadow',?)", (raw,))
        try:
            gap = max(0.05, min(0.6, float(request.form.get("adgap", ad_gap_value()))))
        except ValueError:
            gap = ad_gap_value()
        conn.execute("INSERT OR REPLACE INTO meta(k,v) VALUES('ad_gap',?)", (str(gap),))
        conn.commit()
        flash("Advertisement drop shadow saved.", "ok")
        return redirect(url_for("ads"))

    @app.route("/ads/<int:aid>/delete", methods=["POST"])
    @need("admin", "editor")
    def ads_delete(aid):
        conn = db()
        a = conn.execute("SELECT * FROM ads WHERE id=?", (aid,)).fetchone() or abort(404)
        conn.execute("DELETE FROM ads WHERE id=?", (aid,))
        conn.commit()
        try:
            os.remove(os.path.join(app.config["UPLOAD_DIR"], a["image"]))
        except OSError:
            pass
        flash("Card removed.", "ok")
        return redirect(url_for("ads"))

    MAX_SIDE, THUMB_SIDE = 2400, 400

    def photo_choices():
        libs = {}
        for r in db().execute("SELECT pl.photo_id, l.name FROM photo_libs pl JOIN libraries l ON l.id=pl.lib_id ORDER BY l.name"):
            libs.setdefault(r["photo_id"], []).append(r["name"])
        return [dict(r, libs=", ".join(libs.get(r["id"], []))) for r in db().execute("SELECT * FROM photos ORDER BY title")]

    def put_in_libraries(photo_id, lib_ids):
        valid = {r["id"] for r in db().execute("SELECT id FROM libraries")}
        for i in lib_ids:
            if str(i).isdigit() and int(i) in valid:
                db().execute("INSERT OR IGNORE INTO photo_libs VALUES(?,?)", (photo_id, int(i)))

    def add_photo(data, filename, title="", alt="", decorative=0, extra_tags=(), allow_blank_alt=False, lib_ids=(), extra_meta=None):
        """Save one image: read its metadata, shrink it for print and web, make a small thumbnail, add it to the library.
        Returns (ok, message)."""
        from PIL import Image, ImageOps
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        if ext not in LIBRARY_EXT:
            return False, "not an image"
        orig_name, orig_bytes = filename.lower(), len(data)
        for r in db().execute("SELECT meta FROM photos WHERE meta LIKE ?", ("%" + json.dumps(orig_name)[1:-1] + "%",)):
            try:
                m = json.loads(r["meta"] or "{}")
            except ValueError:
                continue
            if m.get("orig_name") == orig_name and m.get("orig_bytes") == orig_bytes:
                return False, "duplicate"
        is_tiff = ext in ("tif", "tiff")
        if is_tiff:
            filename, ext = filename.rsplit(".", 1)[0] + ".jpg", "jpg"
        base = secure_filename(filename) or "photo." + ext
        name = f"{uuid.uuid4().hex}_{base}"
        path = os.path.join(app.config["UPLOAD_DIR"], name)
        with open(path, "wb") as out:
            out.write(data)
        try:
            meta = read_photo_meta(path)
            meta["orig_name"], meta["orig_bytes"] = orig_name, orig_bytes
            meta.update({k: v for k, v in (extra_meta or {}).items() if v})
            with Image.open(path) as im:
                im.load()
                animated = getattr(im, "is_animated", False) and not is_tiff
                if not animated:
                    im = ImageOps.exif_transpose(im)
                    big = max(im.size) > MAX_SIDE
                    if big:
                        im.thumbnail((MAX_SIDE, MAX_SIDE))
                    if big or ext in ("jpg", "jpeg"):
                        kw = {"quality": 85, "optimize": True} if ext in ("jpg", "jpeg") else {}
                        if ext in ("jpg", "jpeg") and im.mode not in ("RGB", "L"):
                            im = im.convert("RGB")
                        im.save(path, **kw)
                        meta["size"] = "%d x %d px" % im.size
                th = im.copy()
                th.thumbnail((THUMB_SIDE, THUMB_SIDE))
                if th.mode not in ("RGB", "RGBA", "L"):
                    th = th.convert("RGB")
                if th.mode == "RGBA" and ext in ("jpg", "jpeg"):
                    th = th.convert("RGB")
                th.save(os.path.join(app.config["UPLOAD_DIR"], "thumb_" + name), **({"quality": 80} if ext in ("jpg", "jpeg") else {}))
        except Exception:
            os.remove(path)
            return False, "could not be read as an image"
        title = title or meta["title"] or os.path.splitext(base)[0].replace("_", " ").replace("-", " ").strip().title()
        alt = alt or meta["description"]
        tags = [t.strip() for t in extra_tags if t and t.strip()]
        have = {t.lower() for t in tags}
        tags += [k for k in meta["keywords"] if k.lower() not in have]
        if not (alt or decorative):
            if not allow_blank_alt:
                os.remove(path)
                try:
                    os.remove(os.path.join(app.config["UPLOAD_DIR"], "thumb_" + name))
                except OSError:
                    pass
                return False, "needs a description"
            alt = title
            tags.append("needs-description")
        cur = db().execute("INSERT INTO photos(filename,title,alt,tags,uploaded_by,decorative,meta) VALUES(?,?,?,?,?,?,?)",
                           (name, title, alt, ", ".join(tags), g.user["username"], decorative, json.dumps(meta)))
        put_in_libraries(cur.lastrowid, lib_ids)
        db().commit()
        return True, ""

    @app.errorhandler(413)
    def too_big(e):
        flash("That upload is too big (the limit is 100 MB per upload). Add fewer photos at a time, or use smaller files.", "error")
        return redirect(request.referrer or url_for("library"))

    @app.route("/library/bulk", methods=["POST"])
    @need("admin", "editor", "contributor")
    @need_area("library")
    def library_bulk():
        files = [f for f in request.files.getlist("photos") if f and f.filename]
        extra = [t for t in request.form.get("tags", "").split(",")]
        added, skipped = 0, []
        items = []
        dupes = 0
        for f in files:
            if f.filename.lower().endswith(".zip"):
                try:
                    z = zipfile.ZipFile(io.BytesIO(f.read()))
                    for info in z.infolist()[:500]:
                        base = os.path.basename(info.filename)
                        if info.is_dir() or base.startswith(("._", ".")) or "__MACOSX" in info.filename or info.file_size > 40 * 1024 * 1024:
                            continue
                        items.append((base, z.read(info)))
                except zipfile.BadZipFile:
                    skipped.append("%s (not a valid zip)" % f.filename)
            else:
                items.append((f.filename, f.read()))
        for fname, data in items[:500]:
            ok, why = add_photo(data, fname, extra_tags=extra, allow_blank_alt=True, lib_ids=request.form.getlist("libs"))
            if ok:
                added += 1
            elif why == "duplicate":
                dupes += 1
            else:
                skipped.append("%s (%s)" % (fname, why))
        flash("Added %d photo%s. They were resized for print and web and their keywords were read automatically. Photos tagged needs-description have no description of their own; search for that tag to fix them." % (added, "" if added == 1 else "s")
              + ((" %d already in the library (same name and size) and not added again." % dupes) if dupes else "")
              + ((" Skipped: " + "; ".join(skipped[:10])) if skipped else ""), "ok" if added or dupes else "error")
        return redirect(url_for("library"))

    @app.route("/library/browse")
    @need()
    @need_area("browse")
    def library_browse():
        conn = db()
        q = request.args.get("q", "").strip()
        lib = request.args.get("lib", "")
        sql = "SELECT * FROM photos WHERE (title LIKE ? OR tags LIKE ? OR alt LIKE ? OR meta LIKE ?)"
        args = [f"%{q}%"] * 4
        if lib.isdigit():
            sql += " AND id IN (SELECT photo_id FROM photo_libs WHERE lib_id=?)"
            args.append(int(lib))
        elif lib == "none":
            sql += " AND id NOT IN (SELECT photo_id FROM photo_libs)"
        rows = conn.execute(sql + " ORDER BY title", args).fetchall()
        libs_all = conn.execute("SELECT l.*, (SELECT COUNT(*) FROM photo_libs WHERE lib_id=l.id) n FROM libraries l ORDER BY l.name").fetchall()
        return render_template("library_browse.html", photos=rows, q=q, libs=libs_all, lib=lib)

    @app.route("/library", methods=["GET", "POST"])
    @need()
    @need_area("library")
    def library():
        conn = db()
        if request.method == "POST":
            if g.user["role"] not in ("admin", "editor", "contributor"):
                abort(403)
            f = request.files.get("photo")
            pend = request.form.get("pending", "")
            pend_path = os.path.join(app.config["UPLOAD_DIR"], "ai_tmp", pend + ".png")
            if (not f or not f.filename) and re.fullmatch(r"[0-9a-f]{32}", pend) and os.path.exists(pend_path):
                with open(pend_path, "rb") as fh:
                    pdata = fh.read()
                os.remove(pend_path)
                fname = (secure_filename(request.form.get("title", "")) or "ai-picture") + ".png"
            elif not f or not f.filename:
                flash("Choose an image file.", "error")
                return redirect(url_for("library"))
            else:
                pdata, fname = f.read(), f.filename
            ok, why = add_photo(pdata, fname, request.form.get("title", "").strip(), request.form.get("alt", "").strip(),
                                1 if request.form.get("decorative") else 0, request.form.get("tags", "").split(","),
                                lib_ids=request.form.getlist("libs"),
                                extra_meta={k: request.form.get(k, "").strip()[:1500] for k in ("long_desc", "ai_tool", "prompt", "credit", "purpose")})
            if not ok and why == "duplicate":
                flash("That photo is already in the library (same name and size), so it was not added again.", "error")
                return redirect(url_for("library"))
            if not ok:
                flash("That photo was not added: %s. Choose a PNG, JPG, GIF or WebP image; if it has no description of its own, describe it in the alternative text box or tick Decorative." % why, "error")
                return redirect(url_for("library"))
            flash("Photo added to the library.", "ok")
            return redirect(url_for("library"))
        q = request.args.get("q", "").strip()
        lib = request.args.get("lib", "")
        sql = "SELECT * FROM photos WHERE (title LIKE ? OR tags LIKE ? OR alt LIKE ? OR meta LIKE ?)"
        args = [f"%{q}%"] * 4
        if lib.isdigit():
            sql += " AND id IN (SELECT photo_id FROM photo_libs WHERE lib_id=?)"
            args.append(int(lib))
        elif lib == "none":
            sql += " AND id NOT IN (SELECT photo_id FROM photo_libs)"
        rows = conn.execute(sql + " ORDER BY title", args).fetchall()
        libs_all = conn.execute("SELECT l.*, (SELECT COUNT(*) FROM photo_libs WHERE lib_id=l.id) n FROM libraries l ORDER BY l.name").fetchall()
        member = {}
        for r in conn.execute("SELECT photo_id, lib_id FROM photo_libs"):
            member.setdefault(r["photo_id"], set()).add(r["lib_id"])
        metas = {}
        for r in rows:
            try:
                metas[r["id"]] = json.loads(r["meta"] or "{}")
            except ValueError:
                metas[r["id"]] = {}
        return render_template("library.html", photos=rows, q=q, metas=metas, libs=libs_all, member=member, lib=lib)

    @app.route("/library/libraries", methods=["POST"])
    @need("admin", "editor")
    def library_libraries():
        conn = db()
        action = request.form.get("action")
        if action == "add":
            name = " ".join(request.form.get("name", "").split())[:60]
            if name:
                conn.execute("INSERT OR IGNORE INTO libraries(name) VALUES(?)", (name,))
                flash("Library \"%s\" is ready." % name, "ok")
        elif action == "rename" and request.form.get("name", "").strip():
            try:
                conn.execute("UPDATE libraries SET name=? WHERE id=?", (" ".join(request.form["name"].split())[:60], int(request.form["lib"])))
            except sqlite3.IntegrityError:
                flash("A library with that name already exists.", "error")
        elif action == "remove":
            lid = int(request.form["lib"])
            conn.execute("DELETE FROM photo_libs WHERE lib_id=?", (lid,))
            conn.execute("DELETE FROM libraries WHERE id=?", (lid,))
            flash("Library removed. Its photos are still in the photo library.", "ok")
        conn.commit()
        return redirect(url_for("library"))

    @app.route("/library/<int:pid>/libraries", methods=["POST"])
    @need("admin", "editor")
    def photo_libraries(pid):
        conn = db()
        conn.execute("SELECT 1 FROM photos WHERE id=?", (pid,)).fetchone() or abort(404)
        conn.execute("DELETE FROM photo_libs WHERE photo_id=?", (pid,))
        put_in_libraries(pid, request.form.getlist("libs"))
        conn.commit()
        return redirect(url_for("library", lib=request.form.get("back", "")))

    @app.route("/library/libraries-all", methods=["POST"])
    @need("admin", "editor")
    def photo_libraries_all():
        conn = db()
        for pid in {int(x) for x in request.form.getlist("pids") if x.isdigit()}:
            if conn.execute("SELECT 1 FROM photos WHERE id=?", (pid,)).fetchone():
                conn.execute("DELETE FROM photo_libs WHERE photo_id=?", (pid,))
                put_in_libraries(pid, request.form.getlist(f"libs_{pid}"))
        conn.commit()
        flash("All library changes saved.", "ok")
        return redirect(url_for("library", lib=request.form.get("back", "")))

    @app.route("/library/<int:pid>/delete", methods=["POST"])
    @need("admin", "editor")
    def library_delete(pid):
        conn = db()
        p = conn.execute("SELECT * FROM photos WHERE id=?", (pid,)).fetchone() or abort(404)
        in_use = conn.execute("SELECT 1 FROM items WHERE image=?", (p["filename"],)).fetchone()
        if in_use:
            flash("That photo is used in a bulletin and cannot be deleted.", "error")
        else:
            conn.execute("DELETE FROM photo_libs WHERE photo_id=?", (pid,))
            conn.execute("DELETE FROM photos WHERE id=?", (pid,))
            conn.commit()
            try:
                os.remove(os.path.join(app.config["UPLOAD_DIR"], p["filename"]))
            except OSError:
                pass
            try:
                os.remove(os.path.join(app.config["UPLOAD_DIR"], "thumb_" + p["filename"]))
            except OSError:
                pass
        return redirect(url_for("library"))

    PDF_BASE = "http://pdf.local/"

    lite_cache = {}

    sized_dir = os.path.join(os.path.dirname(os.path.realpath(app.config["UPLOAD_DIR"])), os.path.basename(os.path.realpath(app.config["UPLOAD_DIR"])) + "_sized")

    def lite_image(path, limit=700):
        """A smaller copy of a photo (kept on disk); layout is set by CSS size, so the page looks the same."""
        key = (path, os.path.getmtime(path), limit)
        if key in lite_cache:
            return lite_cache[key]
        res = None
        try:
            from PIL import Image
            with Image.open(path) as im:
                fmt = im.format or "PNG"
                if max(im.size) > limit and fmt in ("JPEG", "PNG"):
                    mime = "image/jpeg" if fmt == "JPEG" else "image/png"
                    tag = hashlib.sha1(f"{path}|{key[1]}".encode()).hexdigest()[:16]
                    f = os.path.join(sized_dir, f"{tag}-{limit}")
                    if os.path.exists(f):
                        res = (open(f, "rb").read(), mime)
                    else:
                        im.load()
                        im.thumbnail((limit, limit))
                        buf = io.BytesIO()
                        im.save(buf, fmt, **({"quality": 70 if limit <= 700 else 85} if fmt == "JPEG" else {}))
                        res = (buf.getvalue(), mime)
                        try:
                            os.makedirs(sized_dir, exist_ok=True)
                            with open(f + ".tmp", "wb") as fh:
                                fh.write(res[0])
                            os.replace(f + ".tmp", f)
                        except OSError:
                            pass
        except Exception:
            res = None
        lite_cache[key] = res
        while len(lite_cache) > 300:
            lite_cache.pop(next(iter(lite_cache)))
        return res

    def pdf_fetcher(url, limit=2000):
        from weasyprint import default_url_fetcher
        for prefix, folder in ((PDF_BASE + "uploads/", app.config["UPLOAD_DIR"]),
                               (PDF_BASE + "static/", app.static_folder)):
            if url.startswith(prefix):
                path = os.path.realpath(os.path.join(folder, url[len(prefix):]))
                if not path.startswith(os.path.realpath(folder) + os.sep):
                    raise ValueError("blocked path")
                import mimetypes
                if prefix.endswith("uploads/"):
                    small = lite_image(path, limit)
                    if small:
                        return {"file_obj": io.BytesIO(small[0]), "mime_type": small[1]}
                return {"file_obj": open(path, "rb"),
                        "mime_type": mimetypes.guess_type(path)[0] or "application/octet-stream"}
        if url.startswith("data:"):
            return default_url_fetcher(url)
        raise ValueError("external resources are not allowed in PDFs")

    pdf_cache = {}
    page_cache = {}
    fit_memory = {}
    fit_file = sized_dir + "_fit2.json"  # new name: drops sizes saved by the briefly-buggy measuring code
    try:
        with open(fit_file) as fh:
            for k, v in json.load(fh).items():
                a, b_ = k.split(":")
                fit_memory[(int(a), int(b_))] = (v[0], v[1])
    except (OSError, ValueError):
        pass

    def save_fit_memory():
        try:
            tmp = fit_file + ".tmp"
            with open(tmp, "w") as fh:
                json.dump({f"{a}:{b_}": list(v) for (a, b_), v in fit_memory.items()}, fh)
            os.replace(tmp, fit_file)
        except OSError:
            pass

    def apply_fit_memory(bid, pages):
        """Start from the photo and text-fit sizes the last PDF build settled on, so the on-screen preview and later builds match it."""
        for pg in pages:
            mem = fit_memory.get((bid, pg["number"]))
            if mem:
                if pg.get("photo_h") and mem[0]:
                    pg["photo_h"] = mem[0]
                pg["fit"] = max(pg.get("fit", 0), mem[1])

    building = {}
    bg_timers = {}

    @app.teardown_request
    def release_building(exc):
        ck_ = g.pop("building_ck", None)
        if ck_:
            ev_ = building.pop(ck_, None)
            if ev_:
                ev_.set()

    def bg_preview(bid, sess):
        """Build the Exact PDF preview in the background so it is cached by the time it is opened."""
        try:
            c = app.test_client()
            with c.session_transaction() as st:
                st.update(sess)
            c.get(f"/bulletins/{bid}/download/web?inline=1")
        except Exception:
            app.logger.exception("background preview failed")

    @app.after_request
    def schedule_preview(resp):
        try:
            if (request.method == "POST" and resp.status_code < 400 and os.environ.get("BULLETIN_BG_PREVIEW") == "1"
                    and g.get("user") and request.view_args and request.view_args.get("bid")
                    and "/download" not in request.path):
                import threading
                bid = request.view_args["bid"]
                sess = dict(session)
                old = bg_timers.pop(bid, None)
                if old:
                    old.cancel()
                t = threading.Timer(8.0, bg_preview, (bid, sess))
                t.daemon = True
                bg_timers[bid] = t
                t.start()
        except Exception:
            pass
        return resp

    def boxes_with_parent(page_box):
        """Yield (box, parent) for every box on a page, including those inside absolutely positioned boxes (WeasyPrint hides them from descendants() and has no .parent)."""
        stack = [(page_box, None)]
        while stack:
            box, par = stack.pop()
            yield box, par
            for ch in reversed(list(getattr(box, "children", None) or [])):
                stack.append((ch, box))

    def page_overflow(doc):
        """Return {page index: px of content below the usable page body} for pages whose content is cut off."""
        out = {}
        for i, page in enumerate(doc.pages):
            for box in page._page_box.descendants():
                el = getattr(box, "element", None)
                if el is None or box.element_tag != "div" or "pg-body" not in (el.get("class") or "").split():
                    continue
                limit = box.content_box_y() + box.height
                low = max((d.position_y + d.margin_height() for d in box.descendants() if getattr(d, "height", None) is not None), default=0)
                if low > limit + 1:
                    out[i] = low - limit
                break
        return out

    @app.route("/bulletins/<int:bid>/download/<kind>")
    @need()
    def download(bid, kind):
        """kind=web: small, tagged (PDF/UA) file for the website. kind=print: 300 dpi tagged file for printing."""
        if kind not in ("web", "print", "booklet"):
            abort(404)
        b = get_bulletin(bid)
        if g.user["role"] in ("viewer", "webmaster") and b["status"] != "published":
            abort(403)
        if g.user["role"] == "webmaster" and kind != "web":
            abort(403)
        try:
            from weasyprint import HTML
        except (ImportError, OSError):
            flash("PDF export is not available on this server (WeasyPrint and its system libraries are not installed).", "error")
            return redirect(url_for("edit", bid=bid))
        bk = (bid, kind, bool(request.args.get("inline")))
        wait_ev = building.get(bk)
        if wait_ev is not None:
            wait_ev.wait(150)
        t_req = time.time()
        pages, toc = layout(bid, include_pending=g.user["role"] != "viewer" or b["status"] == "published")
        t_layout = time.time() - t_req
        apply_fit_memory(bid, pages)
        passes = []
        def build(pgs):
            h = sup_ordinals(render_template("bulletin.html", b=b, pages=pgs, n_pages=len(pages), toc=toc, pdf=True, foot_addr=footer_address(bid)[0], foot_phone=footer_address(bid)[1], foot_hours=footer_address(bid)[2], cover_dt=g.cover_dt))
            # WeasyPrint cannot resolve calc(var(...)): without this the photo-size limit is dropped and shared-page photos overflow and vanish.
            h = h.replace("</head>", "<style>.sheet .shared figure img, .sheet .shared .sec.img-only figure img { max-height: var(--ph) !important; }</style></head>", 1)
            return h.replace('url(/uploads/', f'url({PDF_BASE}uploads/').replace('src="/uploads/', f'src="{PDF_BASE}uploads/').replace(
                'href="/static/', f'href="{PDF_BASE}static/').replace('src="/static/', f'src="{PDF_BASE}static/')
        opts = ({"dpi": 96, "jpeg_quality": 70} if kind == "web" else {"dpi": 300, "jpeg_quality": 95})
        html = build(pages)
        doc = None
        lite = bool(request.args.get("inline"))
        limit = 700 if lite else (1200 if kind == "web" else 2000)
        fetch = lambda u: pdf_fetcher(u, limit)
        t0 = time.time()
        ck = hashlib.sha1((kind + ("lite" if lite else "") + html).encode("utf-8", "ignore")).hexdigest()
        name = f"St-Peter-Bulletin-{b['sunday']}-{kind}.pdf"
        disp = "inline" if request.args.get("inline") else "attachment"
        if ck in pdf_cache:
            return app.response_class(pdf_cache[ck], mimetype="application/pdf", headers={"Content-Disposition": f'{disp}; filename="{name}"'})
        import threading
        if bk not in building:
            building[bk] = threading.Event()
            g.building_ck = bk
        if lite and os.environ.get("PER_PAGE_PREVIEW") == "1":
            try:
                from pypdf import PdfReader, PdfWriter
                writer = PdfWriter()
                for pg in pages:
                    one = build([pg])
                    pk = hashlib.sha1(("page" + one).encode("utf-8", "ignore")).hexdigest()
                    if pk not in page_cache:
                        pdoc = None
                        for _ in range(4):
                            pdoc = HTML(string=one, base_url=PDF_BASE, url_fetcher=fetch).render()
                            changed = False
                            if pg.get("photo_h"):
                                for box, par in boxes_with_parent(pdoc.pages[0]._page_box):
                                    el = getattr(box, "element", None)
                                    if el is not None and getattr(box, "element_tag", None) == "div" and el.get("id") == f"pg{pg['number']}" and par is not None:
                                        slack = (par.content_box_y() + par.height - (box.position_y + box.margin_height())) / 96.0
                                        n_ph = sum(1 for r in [pg["row"]] + list(pg.get("more") or []) if r["image"]) or 1
                                        if slack > 0.12 or slack < 0:
                                            new_h = round(max(1.0, min(6.5, pg["photo_h"] + (slack - 0.08) / n_ph)), 2)
                                            if new_h != pg["photo_h"]:
                                                pg["photo_h"] = new_h
                                                changed = True
                                        break
                            if page_overflow(pdoc) and pg.get("fit", 0) < 3 and not changed:
                                pg["fit"] = pg.get("fit", 0) + 1
                                changed = True
                            if not changed:
                                break
                            one = build([pg])
                            pdoc = None
                        if pdoc is None:
                            pdoc = HTML(string=one, base_url=PDF_BASE, url_fetcher=fetch).render()
                        page_cache[pk] = pdoc.write_pdf(optimize_images=True, **opts)
                        while len(page_cache) > 200:
                            page_cache.pop(next(iter(page_cache)))
                    for rp in PdfReader(io.BytesIO(page_cache[pk])).pages:
                        writer.add_page(rp)
                out = io.BytesIO()
                writer.write(out)
                pdf = out.getvalue()
                pdf_cache[ck] = pdf
                while len(pdf_cache) > 3:
                    pdf_cache.pop(next(iter(pdf_cache)))
                return app.response_class(pdf, mimetype="application/pdf", headers={"Content-Disposition": f'{disp}; filename="{name}"', "X-Build-Seconds": f"{time.time() - t0:.1f}"})
            except ImportError:
                pass
        # One measuring loop: size shared-page photos to the real spare room, and tighten type on any page that still overflows.
        try:
            shared = {pg["number"]: pg for pg in pages if pg.get("photo_h")}
            for _ in range(4):
                if doc is None:
                    t_p = time.time()
                    doc = HTML(string=html, base_url=PDF_BASE, url_fetcher=fetch).render()
                    passes.append(time.time() - t_p)
                changed = False
                for page in doc.pages:
                    for box, par in boxes_with_parent(page._page_box):
                        el = getattr(box, "element", None)
                        m = re.fullmatch(r"pg(\d+)", el.get("id") or "") if el is not None and getattr(box, "element_tag", None) == "div" else None
                        if not m or int(m.group(1)) not in shared or par is None:
                            continue
                        slack = (par.content_box_y() + par.height - (box.position_y + box.margin_height())) / 96.0
                        pg = shared[int(m.group(1))]
                        n_ph = sum(1 for r in [pg["row"]] + list(pg.get("more") or []) if r["image"]) or 1
                        if slack > 0.12 or slack < 0:
                            new_h = round(max(1.0, min(6.5, pg["photo_h"] + (slack - 0.08) / n_ph)), 2)
                            if new_h != pg["photo_h"]:
                                pg["photo_h"] = new_h
                                changed = True
                        break
                for k in page_overflow(doc):
                    pg = pages[k] if k < len(pages) else None
                    if pg is not None and pg.get("fit", 0) < 3 and not (pg.get("photo_h") and pg["photo_h"] > 1.0 and changed):
                        pg["fit"] = pg.get("fit", 0) + 1
                        changed = True
                if not changed:
                    break
                html = build(pages)
                doc = None
        except Exception:
            doc = None
        # Reuse the last measured render when photo sizes were already right, saving a full layout pass.
        if doc is None:
            t_p = time.time()
            doc = HTML(string=html, base_url=PDF_BASE, url_fetcher=fetch).render()
            passes.append(time.time() - t_p)
        t_w = time.time()
        if kind == "booklet":
            try:
                from bulletin_designer import booklet as _bk
                pdf = _bk.make_booklet(doc.write_pdf(optimize_images=True, **opts))
            except ImportError:
                flash("Booklet export needs the 'pypdf' package (pip install pypdf).", "error")
                return redirect(url_for("edit", bid=bid))
        else:
            pdf = doc.write_pdf(pdf_variant="pdf/ua-1", optimize_images=True, **opts)
        t_w = time.time() - t_w
        for pg in pages:
            fit_memory[(bid, pg["number"])] = (pg.get("photo_h"), pg.get("fit", 0))
        save_fit_memory()
        pdf_cache[ck] = pdf
        pdf_cache[hashlib.sha1((kind + ("lite" if lite else "") + html).encode("utf-8", "ignore")).hexdigest()] = pdf
        while len(pdf_cache) > 4:
            pdf_cache.pop(next(iter(pdf_cache)))
        timing = f"layout {t_layout:.1f}s; renders {','.join(f'{x:.1f}' for x in passes)}s; write {t_w:.1f}s; total {time.time() - t_req:.1f}s"
        app.logger.info("PDF %s", timing)
        return app.response_class(pdf, mimetype="application/pdf",
                                  headers={"Content-Disposition": f'{disp}; filename="{name}"', "X-Build-Seconds": f"{time.time() - t0:.1f}", "X-Timing": timing})

    @app.route("/uploads/<path:name>")
    @need()
    def uploads(name):
        return send_from_directory(app.config["UPLOAD_DIR"], name)

    def help_match(q):
        words = {w for w in re.findall(r"[a-z0-9]+", q.lower()) if len(w) > 2}
        scored = []
        for tid, title, kw, steps in help_content.TOPICS:
            hay = set(re.findall(r"[a-z0-9]+", (kw + " " + title).lower()))
            n = len(words & hay)
            if n:
                scored.append((n, tid))
        scored.sort(reverse=True)
        return [t for _, t in scored[:3]]

    @app.route("/bulletins/<int:bid>/check")
    @need()
    def bulletin_check(bid):
        get_bulletin(bid)
        rows = db().execute("""SELECT s.id sid, s.title, i.body, i.image, i.status, i.pages, i.font_pt, i.item_gap, i.para_gap, i.title_gap, i.min_h, i.hide_title, i.carried
            FROM items i JOIN sections s ON s.id=i.section_id WHERE i.bulletin_id=? ORDER BY i.position""", (bid,)).fetchall()
        issues = []
        for r in rows:
            name = r["title"]
            filled = bool((r["body"] or "").strip() or r["image"])
            if r["status"] == "pending":
                issues.append((r["sid"], name, "Waiting for approval, so it will not print yet.", "An editor should open it and approve it."))
            if filled and (r["pages"] or 1) > 1:
                issues.append((r["sid"], name, "Set to take %d pages. Its text prints only on the first page." % r["pages"], "Set Pages back to 1 unless you really want blank follow-on pages."))
            if r["font_pt"] and r["font_pt"] < 12:
                issues.append((r["sid"], name, "Text size is %s pt, which is hard to read." % r["font_pt"], "Use 12 pt or larger."))
            if max(r["item_gap"] or 0, r["para_gap"] or 0, r["title_gap"] or 0) > 0.4:
                issues.append((r["sid"], name, "A spacing value is very large.", "Use Reset layout to standard on this section."))
            if r["carried"] and filled:
                issues.append((r["sid"], name, "Text was carried over from last week.", "Check dates, names and amounts are updated."))
        return render_template("check.html", b=get_bulletin(bid), issues=issues)

    @app.route("/help", methods=["GET", "POST"])
    @need()
    def help_page():
        conn = db()
        q, found = "", []
        if request.method == "POST":
            q = request.form.get("question", "").strip()[:500]
            if q:
                found = help_match(q)
                conn.execute("INSERT INTO questions(asked_by,question,created) VALUES(?,?,?)",
                             (g.user["username"], q, datetime.now().strftime("%Y-%m-%d %H:%M")))
                conn.commit()
        mine = conn.execute("SELECT * FROM questions WHERE asked_by=? AND answer!='' ORDER BY id DESC LIMIT 10",
                            (g.user["username"],)).fetchall()
        open_q = conn.execute("SELECT * FROM questions WHERE answer='' ORDER BY id DESC LIMIT 50").fetchall() \
            if g.user["role"] in ("admin", "editor") else []
        return render_template("help.html", updated=help_content.LAST_UPDATED, topics=help_content.TOPICS, q=q, found=found, mine=mine, open_q=open_q)

    @app.route("/help/answer/<int:qid>", methods=["POST"])
    @need("admin", "editor")
    def help_answer(qid):
        ans = request.form.get("answer", "").strip()[:2000]
        if ans:
            db().execute("UPDATE questions SET answer=?, answered_by=? WHERE id=?", (ans, g.user["display_name"] or g.user["username"], qid))
            db().commit()
            flash("Answer saved.", "ok")
        return redirect(url_for("help_page") + "#questions")

    def all_fonts():
        d = dict(FONT_STYLES)
        try:
            for r in db().execute("SELECT id, name FROM custom_fonts ORDER BY name"):
                d["cf%d" % r["id"]] = r["name"] + " (uploaded)"
        except sqlite3.Error:
            pass
        return d

    @app.context_processor
    def font_css():
        from markupsafe import Markup
        css = ""
        try:
            rows = db().execute("SELECT id, name, filename FROM custom_fonts").fetchall()
        except (sqlite3.Error, RuntimeError):
            rows = []
        for r in rows:
            nm = "".join(ch for ch in r["name"] if ch.isalnum() or ch in " -")
            css += ('@font-face { font-family: "%s"; src: url(/uploads/fonts/%s); }\n.fs-cf%d, .fs-cf%d * { font-family: "%s", Georgia, serif; }\n'
                    '.toc-box.fs-cf%d .toc li { font-family: inherit; }\n' % (nm, r["filename"], r["id"], r["id"], nm, r["id"]))
        return {"font_css": Markup(css)}

    @app.route("/design-studio")
    @need()
    def design_studio():
        libs = db().execute("SELECT id, name FROM libraries ORDER BY name").fetchall()
        return render_template("design_studio.html", libs=libs, ai_on=bool(os.environ.get("OPENAI_API_KEY")))

    AI_SIZES = {"1536x1024": "1536x1024", "1024x1536": "1024x1536", "1024x1024": "1024x1024",
                "1152x1536": "1024x1536", "1536x1152": "1536x1024", "1536x768": "1536x1024", "1920x384": "1536x1024"}

    @app.route("/design-studio/generate", methods=["POST"])
    @need("admin", "editor", "contributor")
    def design_generate():
        key = os.environ.get("OPENAI_API_KEY")
        if not key:
            return jsonify(error="AI picture making is not turned on."), 400
        prompt = (request.form.get("prompt") or "").strip()[:3000]
        if len(prompt) < 20:
            return jsonify(error="Fill in the prompt first."), 400
        size = AI_SIZES.get(request.form.get("size", ""), "1536x1024")
        limit = int(os.environ.get("AI_IMAGE_DAILY_LIMIT", "10"))
        day = date.today().isoformat()
        conn = db()
        row = conn.execute("SELECT n FROM ai_usage WHERE day=?", (day,)).fetchone()
        if row and row["n"] >= limit:
            return jsonify(error="The daily limit of %d AI pictures has been reached. Try again tomorrow." % limit), 429
        body = json.dumps({"model": os.environ.get("OPENAI_IMAGE_MODEL", "gpt-image-1"), "prompt": prompt, "size": size, "n": 1}).encode()
        req = urllib.request.Request("https://api.openai.com/v1/images/generations", data=body,
                                     headers={"Authorization": "Bearer " + key, "Content-Type": "application/json"})
        try:
            with urllib.request.urlopen(req, timeout=120) as resp:
                data = json.loads(resp.read())
            raw = base64.b64decode(data["data"][0]["b64_json"])
        except Exception:
            return jsonify(error="The AI service could not make the picture. Check the key, the account balance, or try again."), 502
        conn.execute("INSERT INTO ai_usage(day, n) VALUES(?, 1) ON CONFLICT(day) DO UPDATE SET n=n+1", (day,))
        conn.commit()
        tmp = os.path.join(app.config["UPLOAD_DIR"], "ai_tmp")
        os.makedirs(tmp, exist_ok=True)
        token = secrets.token_hex(16)
        with open(os.path.join(tmp, token + ".png"), "wb") as fh:
            fh.write(raw)
        return jsonify(token=token, image="data:image/png;base64," + base64.b64encode(raw).decode(), left=max(0, limit - (row["n"] if row else 0) - 1))

    @app.route("/fonts", methods=["GET", "POST"])
    @need("admin")
    def fonts():
        conn = db()
        if request.method == "POST":
            if request.form.get("action") == "remove":
                r = conn.execute("SELECT * FROM custom_fonts WHERE id=?", (request.form.get("fid"),)).fetchone()
                if r:
                    conn.execute("DELETE FROM custom_fonts WHERE id=?", (r["id"],))
                    conn.execute("UPDATE items SET font_style='' WHERE font_style=?", ("cf%d" % r["id"],))
                    conn.execute("UPDATE bulletins SET toc_style='' WHERE toc_style=?", ("cf%d" % r["id"],))
                    conn.commit()
                    try:
                        os.remove(os.path.join(app.config["UPLOAD_DIR"], "fonts", r["filename"]))
                    except OSError:
                        pass
                    flash("Font removed. Sections that used it went back to the standard font.", "ok")
                return redirect(url_for("fonts"))
            f = request.files.get("font")
            name = " ".join("".join(ch for ch in request.form.get("name", "") if ch.isalnum() or ch in " -").split())[:40]
            ext = f.filename.rsplit(".", 1)[-1].lower() if f and f.filename and "." in f.filename else ""
            data = f.read() if f else b""
            if ext not in ("ttf", "otf", "woff", "woff2") or not data:
                flash("Choose a font file ending in .ttf, .otf, .woff or .woff2.", "error")
            elif len(data) > 5 * 1024 * 1024:
                flash("That font file is too big (limit 5 MB).", "error")
            elif not name:
                flash("Type a name for the font.", "error")
            elif conn.execute("SELECT 1 FROM custom_fonts WHERE name=?", (name,)).fetchone():
                flash("A font with that name is already added.", "error")
            else:
                os.makedirs(os.path.join(app.config["UPLOAD_DIR"], "fonts"), exist_ok=True)
                fn = uuid.uuid4().hex + "." + ext
                with open(os.path.join(app.config["UPLOAD_DIR"], "fonts", fn), "wb") as out:
                    out.write(data)
                conn.execute("INSERT INTO custom_fonts(name, filename) VALUES(?,?)", (name, fn))
                conn.commit()
                flash("Font added. It now appears in every Font style menu, and is built into the PDF.", "ok")
            return redirect(url_for("fonts"))
        rows = conn.execute("SELECT * FROM custom_fonts ORDER BY name").fetchall()
        return render_template("fonts.html", rows=rows, builtin=FONT_STYLES)

    @app.route("/notes", methods=["GET", "POST"])
    @need("admin", "editor")
    def notes():
        """Editors leave notes about things the design tool should do better; the admin works through them."""
        conn = db()
        if request.method == "POST":
            act = request.form.get("action", "add")
            if act == "add":
                text = request.form.get("note", "").strip()[:2000]
                if text:
                    conn.execute("INSERT INTO admin_notes(from_user,page,note,created) VALUES(?,?,?,?)",
                                 (g.user["display_name"] or g.user["username"], request.form.get("page", "").strip()[:200], text,
                                  datetime.now().strftime("%Y-%m-%d %H:%M")))
                    conn.commit()
                    flash("Thank you. Your note was saved for the admin.", "ok")
            elif g.user["role"] == "admin":
                nid = int(request.form.get("nid", 0))
                conn.execute("UPDATE admin_notes SET status=? WHERE id=?", ("done" if act == "done" else "open", nid))
                conn.commit()
            else:
                abort(403)
            return redirect(url_for("notes"))
        rows = conn.execute("SELECT * FROM admin_notes ORDER BY status='done', id DESC LIMIT 200").fetchall()
        if g.user["role"] != "admin":
            me = g.user["display_name"] or g.user["username"]
            rows = [r for r in rows if r["from_user"] == me]
        return render_template("notes.html", rows=rows, page=request.args.get("page", ""))

    @app.route("/slides-guide")
    @need()
    def slides_guide():
        return render_template("slides_guide.html")

    slides_root = os.path.join(os.path.dirname(os.path.realpath(app.config["UPLOAD_DIR"])), "slides_files")
    SLIDE_KINDS = (("standard", "Standard slides (kept every week)"), ("prelude", "Prelude (before Mass, may be secular)"), ("entrance", "Entrance song"), ("gifts", "Presentation of the Gifts song"), ("communion", "Communion song"), ("praise", "Song of Praise (optional, not every week)"), ("songs", "Other songs"), ("readings", "Readings"),
                   ("announcements", "Announcements"), ("notation", "Music notation"), ("going_forth", "Going Forth (after Mass, may be secular)"),
                   ("special", "Special slides (optional, this week only)"),
                   ("ref_prelude", "Reference photos for Prelude"), ("ref_entrance", "Reference photos for Entrance"), ("ref_gifts", "Reference photos for Presentation of the Gifts"), ("ref_communion", "Reference photos for Communion"), ("ref_praise", "Reference photos for Song of Praise"), ("ref_going_forth", "Reference photos for Going Forth"), ("ref_songs", "Reference photos for Other songs"))

    def slides_dir(kind, mass_date=""):
        return os.path.join(slides_root, "standard") if kind == "standard" else os.path.join(slides_root, "weeks", mass_date, kind)

    def slides_list(kind, mass_date=""):
        d = slides_dir(kind, mass_date)
        return sorted(os.listdir(d)) if os.path.isdir(d) else []

    if routes_choir:
        routes_choir.register(app, need, need_area)
        app.config["CHOIR_ON"] = True

    SONG_EXTS = ("txt", "png", "jpg", "jpeg")

    def song_library():
        d = os.path.join(slides_root, "songs")
        found = {}
        for n in sorted(os.listdir(d)) if os.path.isdir(d) else []:
            base, ext = os.path.splitext(n)
            if ext[1:].lower() in SONG_EXTS:
                found.setdefault(base, []).append(n)
        return found

    def song_files(lib, song):
        for key in (music_outline.slug(song["title"]), music_outline.slug(f"{song['book']}-{song['number']}") if song["number"] else ""):
            if key and key in lib:
                return key, lib[key]
        return "", []

    def outline_path(mass_date):
        return os.path.join(slides_root, "weeks", mass_date, "outline.json")

    def load_outline(mass_date):
        try:
            with open(outline_path(mass_date)) as f:
                return json.load(f)
        except (OSError, ValueError):
            return None

    def apply_outline(mass_date, songs):
        """Fill the week's song sections from the lyrics library; re-running replaces only what an earlier run added."""
        previous = load_outline(mass_date) or {}
        for kind, name in previous.get("auto", []):
            path = os.path.join(slides_dir(kind, mass_date), secure_filename(name))
            if kind in dict(SLIDE_KINDS) and os.path.isfile(path):
                os.remove(path)
        lib, auto, report = song_library(), [], []
        for song in songs:
            key, names = song_files(lib, song)
            status = "found" if names else ("mass part" if song["setting"] else "missing")
            if names and song["kind"]:
                for n in names:
                    ext = n.rsplit(".", 1)[-1].lower()
                    kind = song["kind"] if ext == "txt" else "notation"
                    dest = slides_dir(kind, mass_date)
                    os.makedirs(dest, exist_ok=True)
                    out = secure_filename(f"{song['label']}-{n}") or n
                    if os.path.exists(os.path.join(dest, out)):
                        continue
                    shutil.copyfile(os.path.join(slides_root, "songs", n), os.path.join(dest, out))
                    auto.append([kind, out])
            if status == "missing" and song["kind"] in ("psalm", "acclamation") and routes_choir:
                try:
                    pages = routes_choir.choir.book_page_images(app, song, date.fromisoformat(mass_date))
                except Exception:
                    pages = []
                dest = slides_dir("notation", mass_date)
                for i, png in enumerate(pages, 1):
                    os.makedirs(dest, exist_ok=True)
                    out = secure_filename(f"{song['label']}-book-page-{i}.png")
                    with open(os.path.join(dest, out), "wb") as f:
                        f.write(png)
                    auto.append(["notation", out])
                status = "words in readings, music from book" if pages else "in readings"
            report.append(dict(song, status=status))
        os.makedirs(os.path.dirname(outline_path(mass_date)), exist_ok=True)
        with open(outline_path(mass_date), "w") as f:
            json.dump({"songs": report, "auto": auto}, f)
        return previous, report

    @app.route("/slides/outline", methods=["POST"])
    @need()
    @need_area("slides")
    def slides_outline():
        mass_date = request.form.get("date") or ""
        text = (request.form.get("outline_text") or "").strip()
        up = request.files.get("outline_pdf")
        if up and up.filename:
            try:
                text = music_outline.extract_pdf_text(up.stream)
            except Exception:
                flash("That PDF could not be read. Try pasting the outline text instead.", "error")
                return redirect(url_for("slides", date=mass_date))
        parsed = music_outline.parse_outline(text) if text else {"songs": [], "date": ""}
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", mass_date):
            mass_date = parsed["date"]
        if not mass_date:
            flash("Choose the Mass date, then upload the outline.", "error")
            return redirect(url_for("slides"))
        if not parsed["songs"]:
            flash("No songs were found in that outline.", "error")
            return redirect(url_for("slides", date=mass_date))
        previous, report = apply_outline(mass_date, parsed["songs"])
        old = {(x["label"], x["title"].lower()) for x in previous.get("songs", [])}
        changed = [x["title"] for x in report if (x["label"], x["title"].lower()) not in old] if old else []
        missing = [x["title"] for x in report if x["status"] == "missing"]
        flash(f"Outline read: {len(report)} songs. Lyrics added for {sum(1 for x in report if x['status'] == 'found')}."
              + (f" Changed since last upload: {', '.join(changed)}." if changed else "")
              + (f" Not in the song library yet: {', '.join(missing)}." if missing else ""), "error" if missing else "")
        return redirect(url_for("slides", date=mass_date))

    @app.route("/slides/songs", methods=["POST"])
    @need()
    @need_area("slides")
    def slides_songs():
        mass_date = request.form.get("date") or ""
        title, code = (request.form.get("title") or "").strip(), (request.form.get("code") or "").strip()
        key = music_outline.slug(title) or music_outline.slug(code)
        if request.form.get("action") == "delete":
            name = secure_filename(request.form.get("name", ""))
            path = os.path.join(slides_root, "songs", name)
            if name and os.path.isfile(path):
                os.remove(path)
            return redirect(url_for("slides", date=mass_date))
        if not key:
            flash("Enter the song title.", "error")
            return redirect(url_for("slides", date=mass_date))
        os.makedirs(os.path.join(slides_root, "songs"), exist_ok=True)
        n = 0
        lyrics = (request.form.get("lyrics") or "").strip()
        if lyrics:
            with open(os.path.join(slides_root, "songs", key + ".txt"), "w", encoding="utf-8") as f:
                f.write(lyrics + "\n")
            n += 1
        for f in request.files.getlist("files"):
            ext = (f.filename or "").rsplit(".", 1)[-1].lower()
            if ext in SONG_EXTS:
                f.save(os.path.join(slides_root, "songs", f"{key}.{ext}"))
                n += 1
        if code and n:
            ckey = music_outline.slug(code)
            for fn in os.listdir(os.path.join(slides_root, "songs")):
                if fn.startswith(key + ".") and ckey != key:
                    shutil.copyfile(os.path.join(slides_root, "songs", fn), os.path.join(slides_root, "songs", ckey + "." + fn.rsplit(".", 1)[-1]))
        flash(f"Saved {title or code} to the song library." if n else "Paste the lyrics or choose a file.", "" if n else "error")
        return redirect(url_for("slides", date=mass_date))

    def week_status(mass_date):
        try:
            with open(os.path.join(slides_root, "weeks", mass_date, "status.json")) as f:
                if json.load(f).get("ready"):
                    return "Ready to use"
        except (OSError, ValueError):
            pass
        has_files = any(slides_list(k, mass_date) for k, _ in SLIDE_KINDS if k != "standard")
        return "In progress" if has_files else "Not started"

    @app.route("/slides/library")
    @need()
    @need_area("slides")
    def slides_library():
        weeks_dir = os.path.join(slides_root, "weeks")
        dates = sorted(os.listdir(weeks_dir), reverse=True) if os.path.isdir(weeks_dir) else []
        today = date.today()
        upcoming = []
        for k in range(0, 8):
            d = today + timedelta(days=(6 - today.weekday()) % 7 + 7 * k)
            upcoming.append(d.isoformat())
        rows = []
        for iso in sorted(set(upcoming) | {d for d in dates if re.fullmatch(r"\d{4}-\d{2}-\d{2}", d)}, reverse=True):
            li = liturgy.info(date.fromisoformat(iso))
            rows.append({"iso": iso, "name": li["name"] or f"Sunday in {li['season']}", "status": week_status(iso),
                         "count": sum(len(slides_list(k, iso)) for k, _ in SLIDE_KINDS if k != "standard" and not k.startswith("ref_"))})
        return render_template("slides_library.html", rows=rows, standard=slides_list("standard"), songs=song_library())

    @app.route("/slides/status", methods=["POST"])
    @need()
    @need_area("slides")
    def slides_status():
        mass_date = request.form.get("date", "")
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", mass_date):
            abort(404)
        os.makedirs(os.path.join(slides_root, "weeks", mass_date), exist_ok=True)
        with open(os.path.join(slides_root, "weeks", mass_date, "status.json"), "w") as f:
            json.dump({"ready": request.form.get("ready") == "1"}, f)
        return redirect(request.form.get("back") == "week" and url_for("slides", date=mass_date) or url_for("slides_library"))

    @app.route("/slides", methods=["GET", "POST"])
    @need()
    @need_area("slides")
    def slides():
        mass_date = request.values.get("date") or ""
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", mass_date):
            mass_date = ""
        if request.method == "POST":
            kind = request.form.get("kind")
            action = request.form.get("action")
            if action == "save_graphics" and mass_date:
                os.makedirs(os.path.join(slides_root, "weeks", mass_date), exist_ok=True)
                with open(os.path.join(slides_root, "weeks", mass_date, "graphics.json"), "w") as f:
                    json.dump([x for x in request.form.getlist("photo") if x.isdigit()][:40], f)
                flash("Saved your ticked graphics for this week.")
                return redirect(url_for("slides", date=mass_date))
            if kind not in dict(SLIDE_KINDS) or (kind != "standard" and not mass_date):
                flash("Choose the Mass date first.", "error")
                return redirect(url_for("slides", date=mass_date))
            d = slides_dir(kind, mass_date)
            if action == "delete":
                name = secure_filename(request.form.get("name", ""))
                if name and os.path.isfile(os.path.join(d, name)):
                    os.remove(os.path.join(d, name))
                    flash(f"Removed {name}.")
            else:
                os.makedirs(d, exist_ok=True)
                n = 0
                for f in request.files.getlist("files"):
                    name = secure_filename(f.filename or "")
                    if name:
                        f.save(os.path.join(d, name))
                        n += 1
                flash(f"Uploaded {n} file(s)." if n else "Choose at least one file.", "" if n else "error")
            return redirect(url_for("slides", date=mass_date))
        dates = sorted(os.listdir(os.path.join(slides_root, "weeks")), reverse=True) if os.path.isdir(os.path.join(slides_root, "weeks")) else []
        return render_template("slides.html", kinds=SLIDE_KINDS, mass_date=mass_date, dates=dates,
                               files={k: slides_list(k, mass_date) for k, _ in SLIDE_KINDS if k == "standard" or mass_date},
                               concept=week_concept(mass_date) if mass_date else None, layouts=CONCEPT_LAYOUTS, accents=CONCEPT_ACCENTS,
                               outline=load_outline(mass_date) if mass_date else None, status=week_status(mass_date) if mass_date else "", library=song_library(),
                               photos=db().execute("SELECT id, filename, title FROM photos ORDER BY id DESC LIMIT 60").fetchall() if mass_date else [], selected=saved_graphics(mass_date) if mass_date else [])

    @app.route("/slides/year")
    @need()
    @need_area("slides")
    def slides_year():
        """The whole church year at a glance: each Sunday's season, color and imagery, and whether its files are in."""
        start = request.args.get("start", "")
        try:
            d = date.fromisoformat(start)
        except ValueError:
            t = date.today()
            d = t
            while d.weekday() != 6 or liturgy.info(d)["name"] != "First Sunday of Advent":
                d -= timedelta(days=1)
        d += timedelta(days=(6 - d.weekday()) % 7)
        rows = []
        for n in range(53):
            sd = d + timedelta(days=7 * n)
            li = liturgy.info(sd)
            if n and li["name"] == "First Sunday of Advent":
                break
            iso = sd.isoformat()
            rows.append({"iso": iso, "info": li, "have": {k: len(slides_list(k, iso)) for k, _ in SLIDE_KINDS if k != "standard"}})
        return render_template("slides_year.html", rows=rows, kinds=[k for k in SLIDE_KINDS if k[0] != "standard"],
                               prev=(d - timedelta(days=364)).isoformat(), nxt=(d + timedelta(days=364)).isoformat())

    @app.route("/slides/file/<kind>/<name>")
    @need()
    @need_area("slides")
    def slides_file(kind, name):
        mass_date = request.args.get("date", "")
        if kind not in dict(SLIDE_KINDS) or (kind != "standard" and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", mass_date)):
            abort(404)
        return send_from_directory(slides_dir(kind, mass_date), secure_filename(name), as_attachment=not request.args.get("inline"))

    CONCEPT_SCENES = ("sunrise over a calm turquoise bay", "a fishing boat returning at golden hour", "mangrove roots reflecting evening light",
                      "palm silhouettes against a glowing sunset", "a lighthouse above gentle waves", "clouds opening over the Keys horizon",
                      "a quiet pier at first light", "shallow clear water with sunlit ripples", "a sailboat crossing a rose-gold sky",
                      "seabirds gliding over the seven-mile bridge", "moonlight on still water", "a coral-pink dusk over the flats")
    CONCEPT_LAYOUTS = ("full-bleed photo with a soft dark overlay and centered text", "left text panel over a blurred scene on the right",
                       "framed cream panel over the scene", "gradient band across the lower third with the text above the horizon",
                       "centered text with a thin gold border", "two-tone split: scene above, text panel below")
    CONCEPT_ACCENTS = ("gold leaf edge", "wave line along the bottom", "small cross motif in a corner", "brush-stroke title banner",
                       "soft light rays", "rope border")

    def week_concept(mass_date):
        """A different but repeatable design concept for each week, tied to the season's colors."""
        d = date.fromisoformat(mass_date)
        n = d.toordinal() // 7
        c = {"scene": CONCEPT_SCENES[n % len(CONCEPT_SCENES)], "layout": CONCEPT_LAYOUTS[(n * 5 + 1) % len(CONCEPT_LAYOUTS)],
             "accent": CONCEPT_ACCENTS[(n * 7 + 2) % len(CONCEPT_ACCENTS)], "color": "", "scale": 100}
        c.update({k: v for k, v in load_design(mass_date).items() if v})
        return c

    def load_design(mass_date):
        try:
            with open(os.path.join(slides_root, "weeks", mass_date, "design.json")) as f:
                d = json.load(f)
        except (OSError, ValueError):
            return {}
        out = {}
        if d.get("layout") in CONCEPT_LAYOUTS:
            out["layout"] = d["layout"]
        if d.get("accent") in CONCEPT_ACCENTS:
            out["accent"] = d["accent"]
        if re.fullmatch(r"[0-9A-Fa-f]{6}", d.get("color", "")):
            out["color"] = d["color"]
        if str(d.get("scale", "")).isdigit() and 70 <= int(d["scale"]) <= 140:
            out["scale"] = int(d["scale"])
        return out

    @app.route("/slides/design", methods=["POST"])
    @need()
    @need_area("slides")
    def slides_design():
        mass_date = request.form.get("date", "")
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", mass_date):
            abort(404)
        os.makedirs(os.path.join(slides_root, "weeks", mass_date), exist_ok=True)
        path = os.path.join(slides_root, "weeks", mass_date, "design.json")
        if request.form.get("reset"):
            if os.path.isfile(path):
                os.remove(path)
            flash("Design is back to this week's automatic choice.")
        else:
            with open(path, "w") as f:
                json.dump({k: request.form.get(k, "").lstrip("#") if k == "color" else request.form.get(k, "") for k in ("layout", "accent", "color", "scale")}, f)
            flash("Design saved. Preview or download the slides to see it.")
        return redirect(url_for("slides", date=mass_date))

    SECTION_ACCENT = {"prelude": "8FB8C9", "entrance": "E0B040", "gifts": "E07A5F", "communion": "F4E9D0", "praise": "3FB8AF",
                      "going_forth": "F28C28", "songs": "E0B040"}
    CREDIT_RE = re.compile(r"^\s*(©|\(c\)|copyright|reprinted|used with permission|text:|music:|words:)", re.I)

    def section_ref(kind, mass_date):
        names = [n for n in slides_list("ref_" + kind, mass_date) if n.rsplit(".", 1)[-1].lower() in ("png", "jpg", "jpeg")]
        if not names:
            return None, None
        return (os.path.join(slides_dir("ref_" + kind, mass_date), names[0]),
                url_for("slides_file", kind="ref_" + kind, name=names[0], date=mass_date, inline=1))

    SONG_KINDS = ("prelude", "entrance", "songs", "gifts", "communion", "praise", "going_forth")
    SEASON_BG = {"Advent": "3B2470", "Christmas": "8A6D12", "Ordinary Time": "14527A", "Lent": "4F3A6B", "Easter": "7A6A3A", "Pentecost": "8C1D26"}

    def saved_graphics(mass_date):
        try:
            with open(os.path.join(slides_root, "weeks", mass_date, "graphics.json")) as f:
                return [str(x) for x in json.load(f)]
        except (OSError, ValueError):
            return []

    def chosen_graphics():
        return (request.args.getlist("photo") or saved_graphics(request.args.get("date", "")))[:40]

    def plan_slides(mass_date, photo_ids=()):
        """The ordered list of slides, shared by the preview and the PowerPoint so they always match."""
        li = liturgy.info(date.fromisoformat(mass_date))
        season = li["season"]
        feast = bool(li["name"]) and li["name"] not in SEASON_BG

        def skip(name):
            n = name.lower()
            if "alleluia" in n and season == "Lent":
                return True
            return "gloria" in n and season in ("Advent", "Lent") and not feast

        BLACK_BETWEEN = ("prelude", "entrance", "songs", "readings", "notation", "gifts", "communion", "praise", "going_forth")
        seen_songs = set()
        plan = [{"type": "text", "title": mass_date, "text": li["name"] or f"Sunday in {season}", "size": 60}]
        for kind in ("prelude", "standard", "entrance", "songs", "readings", "notation", "gifts", "communion", "praise", "announcements", "special", "going_forth"):
            for name in slides_list(kind, mass_date):
                if kind == "standard" and skip(name):
                    continue
                if kind in SONG_KINDS or kind == "special":
                    stem, _, ext0 = name.lower().rpartition(".")
                    key = (re.sub(r"[^a-z0-9]+", "", stem), "img" if ext0 in ("png", "jpg", "jpeg") else ext0)
                    if key in seen_songs:
                        continue
                    seen_songs.add(key)
                path = os.path.join(slides_dir(kind, mass_date), name)
                ext = name.rsplit(".", 1)[-1].lower()
                if kind in BLACK_BETWEEN and ext in ("png", "jpg", "jpeg", "txt", "docx") and len(plan) > 1 and plan[-1]["type"] != "blank":
                    plan.append({"type": "blank"})
                if ext in ("png", "jpg", "jpeg"):
                    plan.append({"type": "image", "path": path, "src": url_for("slides_file", kind=kind, name=name, date=mass_date, inline=1), "label": name})
                elif ext in ("txt", "docx"):
                    label = os.path.splitext(name)[0].replace("_", " ")
                    section = dict(SLIDE_KINDS).get(kind, "").split(" (")[0]
                    try:
                        raw = docx_text(path) if ext == "docx" else music_outline.read_text(path)
                    except Exception:
                        continue
                    groups = music_outline.reading_sections(raw) if kind == "readings" else None
                    if not groups:
                        parts = [x.strip() for x in re.split(r"\n\s*\n", raw) if x.strip()]
                        if ext == "docx" and kind == "announcements":
                            parts = music_outline.announcements_from(parts)
                        parts = music_outline.repeat_refrain(parts)
                        if kind in ("readings", "announcements"):
                            parts = [piece for part in parts for piece in music_outline.split_long(part)]
                        groups = [(label, parts)]
                    song = kind in SONG_KINDS
                    bgp, bgs = section_ref(kind, mass_date) if song else (None, None)
                    for g, (gtitle, parts) in enumerate(groups):
                        if g and plan[-1]["type"] != "blank":
                            plan.append({"type": "blank"})
                        for k, part in enumerate(parts):
                            lines = part.splitlines()
                            credit = " ".join(x.strip() for x in lines if CREDIT_RE.match(x))
                            body = "\n".join(x for x in lines if not CREDIT_RE.match(x)).strip() or part
                            item = {"type": "text", "title": None, "subtitle": None, "text": music_outline.no_orphans(body), "size": 48 if len(body) > 140 else 56,
                                    "style": "song" if song else "plain", "credit": credit if song else "", "bg_path": bgp, "bg_src": bgs}
                            if k == 0 and song:
                                item["title"], item["subtitle"] = section, gtitle
                            elif k == 0 and kind in ("readings", "standard"):
                                item["title"] = gtitle
                            plan.append(item)
        for pid in photo_ids:
            row = db().execute("SELECT filename FROM photos WHERE id=?", (pid,)).fetchone()
            path = row and os.path.join(app.config["UPLOAD_DIR"], os.path.basename(row["filename"]))
            if path and os.path.isfile(path):
                plan.append({"type": "image", "path": path, "src": url_for("uploads", name=row["filename"]), "label": "bulletin graphic"})
        return plan, SEASON_BG.get(season, "14527A")

    def build_pptx(plan, bg_hex, concept=None):
        """Editable 16:9 deck: images become full-slide pictures, text one slide per paragraph."""
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        from pptx.enum.dml import MSO_LINE
        from PIL import Image
        bg = RGBColor.from_string(bg_hex)
        scale = (concept or {}).get("scale", 100) / 100
        for it in plan:
            if it.get("size"):
                it["size"] = max(20, int(it["size"] * scale))
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        blank = prs.slide_layouts[6]
        for item in plan:
            sl = prs.slides.add_slide(blank)
            sl.background.fill.solid()
            if item["type"] == "blank":
                sl.background.fill.fore_color.rgb = RGBColor(0, 0, 0)
                continue
            if item["type"] == "image":
                sl.background.fill.fore_color.rgb = RGBColor(0xFA, 0xF5, 0xE6)
                try:
                    with Image.open(item["path"]) as im:
                        w, h = im.size
                    sw, sh = prs.slide_width, prs.slide_height
                    scale = min(sw / w, sh / h)
                    pw, ph = int(w * scale), int(h * scale)
                    sl.shapes.add_picture(item["path"], int((sw - pw) / 2), int((sh - ph) / 2), pw, ph)
                except Exception:
                    pass
                continue
            sl.background.fill.fore_color.rgb = bg
            sw, sh = prs.slide_width, prs.slide_height
            song = item.get("style") == "song"
            ink = RGBColor(0x3B, 0x0F, 0x0F) if song else RGBColor(255, 255, 255)
            if item.get("bg_path"):
                try:
                    sl.shapes.add_picture(item["bg_path"], 0, 0, sw, sh)
                except Exception:
                    pass

            def shape(kind, x, y, w, h, rgb, alpha=None):
                sp = sl.shapes.add_shape(kind, x, y, w, h)
                sp.fill.solid(); sp.fill.fore_color.rgb = rgb; sp.line.fill.background()
                if alpha:
                    from pptx.oxml.ns import qn
                    from lxml import etree
                    etree.SubElement(sp.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr")), qn("a:alpha")).set("val", str(alpha))
                return sp

            def put(tf, text, size, rgb, bold=True, italic=False):
                for k, line in enumerate(text.splitlines() or [""]):
                    para = tf.paragraphs[0] if k == 0 and not tf.paragraphs[0].runs else tf.add_paragraph()
                    para.alignment = PP_ALIGN.CENTER
                    r = para.add_run(); r.text = line
                    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic; r.font.color.rgb = rgb

            top, bottom = 0.5, 7.0
            if concept:
                lay, acc = concept["layout"], concept["accent"]
                gold = RGBColor(0xE0, 0xB0, 0x40)
                if lay.startswith("full-bleed"):
                    shape(1, 0, 0, sw, sh, RGBColor(0, 0, 0), alpha=35000)
                elif lay.startswith("left text"):
                    shape(1, 0, 0, Inches(2.2), sh, RGBColor(0, 0, 0), alpha=30000)
                elif lay.startswith("framed"):
                    shape(1, Inches(0.5), Inches(0.5), sw - Inches(1.0), sh - Inches(1.0), RGBColor(0xFA, 0xF1, 0xDC), alpha=22000)
                elif lay.startswith("gradient"):
                    shape(1, 0, sh - Inches(2.6), sw, Inches(2.6), RGBColor(0, 0, 0), alpha=40000)
                elif lay.startswith("centered"):
                    b = shape(1, Inches(0.25), Inches(0.25), sw - Inches(0.5), sh - Inches(0.5), bg); b.fill.background(); b.line.color.rgb = gold; b.line.width = Pt(3)
                else:
                    shape(1, 0, 0, sw, Inches(3.2), RGBColor(255, 255, 255), alpha=18000)
                if acc.startswith("gold leaf"):
                    shape(1, 0, 0, Inches(0.18), sh, gold)
                    shape(1, sw - Inches(0.18), 0, Inches(0.18), sh, gold)
                elif acc.startswith("wave"):
                    shape(1, 0, sh - Inches(0.25), sw, Inches(0.25), RGBColor(0x3F, 0xB8, 0xAF), alpha=85000)
                    shape(1, 0, sh - Inches(0.1), sw, Inches(0.1), RGBColor(255, 255, 255), alpha=70000)
                elif acc.startswith("small cross"):
                    cx = sl.shapes.add_textbox(sw - Inches(1.1), Inches(0.15), Inches(0.9), Inches(0.9))
                    put(cx.text_frame, "\u271d", 40, gold)
                elif acc.startswith("brush"):
                    shape(1, 0, Inches(0.12), sw, Inches(0.12), RGBColor(0x8A, 0x16, 0x1E))
                elif acc.startswith("soft light"):
                    shape(7, sw - Inches(4), 0, Inches(4), Inches(3), RGBColor(0xFF, 0xF1, 0xB8), alpha=18000)
                else:
                    r = shape(1, Inches(0.15), Inches(0.15), sw - Inches(0.3), sh - Inches(0.3), bg); r.fill.background()
                    r.line.color.rgb = RGBColor(0xC8, 0xA9, 0x6E); r.line.width = Pt(5); r.line.dash_style = MSO_LINE.DASH
            if item.get("credit"):
                shape(1, 0, sh - Inches(0.95), sw, Inches(0.95), RGBColor(0x6B, 0x14, 0x14))
                cb = sl.shapes.add_textbox(Inches(0.5), sh - Inches(0.9), sw - Inches(1.0), Inches(0.85))
                cb.text_frame.word_wrap = True; cb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                put(cb.text_frame, item["credit"], 14, RGBColor(255, 255, 255), bold=False)
                bottom = 6.4
            if item["title"]:
                bh = 1.4 if item.get("subtitle") else 0.95
                box = shape(1, Inches(2.2), Inches(0.3), Inches(8.9), Inches(bh), RGBColor(0x8A, 0x16, 0x1E))
                box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                put(box.text_frame, item["title"], 38, RGBColor(0xFF, 0xE9, 0xA8), italic=bool(song))
                if item.get("subtitle"):
                    put(box.text_frame, item["subtitle"], 24, RGBColor(0xFF, 0xE9, 0xA8), italic=True)
                top = 0.3 + bh + 0.2
            if song:
                shape(5, Inches(0.9), Inches(top), Inches(11.5), Inches(bottom - top), RGBColor(0xFA, 0xF1, 0xDC), alpha=88000)
            tb = sl.shapes.add_textbox(Inches(1.1), Inches(top), Inches(11.1), Inches(bottom - top))
            tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            put(tf, item["text"], item["size"], ink)
        buf = io.BytesIO(); prs.save(buf); buf.seek(0)
        return buf

    @app.route("/slides/preview")
    @need()
    @need_area("slides")
    def slides_preview():
        mass_date = request.args.get("date", "")
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", mass_date):
            abort(404)
        plan, bg = plan_slides(mass_date, chosen_graphics())
        return render_template("slides_preview.html", plan=plan, bg=bg, mass_date=mass_date)

    @app.route("/slides/pptx")
    @need()
    @need_area("slides")
    def slides_pptx():
        mass_date = request.args.get("date", "")
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", mass_date):
            abort(404)
        try:
            import pptx  # noqa: F401
        except ImportError:
            flash("The PowerPoint tool is not installed on the server yet. Ask the administrator to run: pip install -r requirements.txt", "error")
            return redirect(url_for("slides", date=mass_date))
        plan, bg = plan_slides(mass_date, chosen_graphics())
        concept = week_concept(mass_date)
        return send_file(build_pptx(plan, concept["color"] or bg, concept), as_attachment=True, download_name=f"Mass-slides-{mass_date}.pptx",
                         mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation")

    @app.route("/slides/package")
    @need()
    @need_area("slides")
    def slides_package():
        """One ZIP with the master prompt, the standard slides and this week's uploads, ready to hand to the AI tool."""
        mass_date = request.args.get("date", "")
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", mass_date):
            abort(404)
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            z.write(os.path.join(app.root_path, "static", "slides_master_prompt.txt"), "00_Master_Prompt.txt")
            z.writestr("00_README.txt", f"Mass date: {mass_date}\nGive this whole folder to the AI tool and ask for Stages 1-3 only (inventory, theme, previews). "
                       "Folders: standard (slides used every week), prelude, entrance, gifts, communion, praise, going_forth, songs, readings, announcements, notation.\n")
            li = liturgy.info(date.fromisoformat(mass_date))
            z.writestr("00_Week_Info.txt", f"Date: {mass_date}\nCelebration: {li['name'] or 'Sunday in ' + li['season']}\nSeason: {li['season']}\n"
                       f"Color: {li['color']}\nImagery ideas: {'; '.join(li['ideas'])}\nReadings: {li['readings_url']}\n"
                       f"NEW CONCEPT FOR THIS WEEK (do not reuse last week's look)\nScene: {week_concept(mass_date)['scene']}\n"
                       f"Layout: {week_concept(mass_date)['layout']}\nAccent: {week_concept(mass_date)['accent']}\n"
                       "Blend with the season colors above. Keep the standard slides' structure, the red title banner and readability.\n"
                       "Special slides (if any) are in the 'special' folder: place them where the file names or notes say.\n")
            for kind, _ in SLIDE_KINDS:
                for name in slides_list(kind, mass_date):
                    z.write(os.path.join(slides_dir(kind, mass_date), name), f"{kind}/{name}")
            for pid in chosen_graphics():
                row = db().execute("SELECT filename, title FROM photos WHERE id=?", (pid,)).fetchone()
                path = row and os.path.join(app.config["UPLOAD_DIR"], os.path.basename(row["filename"]))
                if path and os.path.isfile(path):
                    z.write(path, f"graphics/{secure_filename(row['title'] or '') or 'photo'}_{pid}{os.path.splitext(path)[1]}")
        buf.seek(0)
        return send_file(buf, mimetype="application/zip", as_attachment=True, download_name=f"Mass-slides-package-{mass_date}.zip")

    @app.route("/style-guide")
    def style_guide():
        return render_template("style_guide.html", font_styles=all_fonts())

    # ---------- admin ----------
    @app.route("/admin/users", methods=["GET", "POST"])
    @need("admin", "editor")
    def users():
        conn = db()
        mine = ROLES.index(g.user["role"])
        allowed = ROLES[mine:]

        def may_manage(uid):
            row = conn.execute("SELECT role FROM users WHERE id=?", (uid,)).fetchone()
            return bool(row) and row["role"] in allowed

        if request.method == "POST":
            action = request.form.get("action")
            if action in ("invite", "getlink", "setemail", "role", "deactivate", "activate", "delete") and not may_manage(int(request.form.get("uid", 0))):
                abort(403)
            if action == "create" and request.form.get("role") not in allowed:
                abort(403)
            if action == "role" and request.form.get("role") not in allowed:
                abort(403)
            if action == "create":
                role = request.form.get("role")
                email = request.form.get("username", "").strip().lower()
                if role in ROLES and EMAIL_RE.match(email):
                    try:
                        cur = conn.execute("""INSERT INTO users(username,email,display_name,pw_hash,role,must_reset)
                            VALUES(?,?,?,?,?,1)""", (email, email, request.form.get("display_name", "").strip() or email,
                                                  generate_password_hash(secrets.token_urlsafe(24)), role))
                        if role == "contributor":
                            for sid in request.form.getlist("sections"):
                                conn.execute("INSERT OR IGNORE INTO user_sections VALUES(?,?)",
                                             (cur.lastrowid, int(sid)))
                        if role not in ("admin", "editor"):
                            save_areas(conn, cur.lastrowid)
                        conn.commit()
                        sent, link = invite_user(conn.execute("SELECT * FROM users WHERE id=?", (cur.lastrowid,)).fetchone())
                        flash("Invitation emailed to %s." % email if sent else
                              "Email was not sent (" + (app.config.get("MAIL_ERROR") or "unknown reason") + "). Copy this link and send it to %s yourself: %s" % (email, link),
                              "ok" if sent else "error")
                    except sqlite3.IntegrityError:
                        flash("That email already has an account.", "error")
                else:
                    flash("Enter a valid email address and an access level.", "error")
            elif action == "invite":
                em = request.form.get("email", "").strip().lower()
                if EMAIL_RE.match(em):
                    try:
                        conn.execute("UPDATE users SET email=? WHERE id=?", (em, int(request.form["uid"])))
                        conn.commit()
                    except sqlite3.IntegrityError:
                        flash("Another account already uses that email.", "error")
                u = conn.execute("SELECT * FROM users WHERE id=?", (int(request.form["uid"]),)).fetchone()
                if u and u["email"]:
                    sent, link = invite_user(u)
                    flash("Invitation emailed to %s." % u["email"] if sent else
                          "Email was not sent (" + (app.config.get("MAIL_ERROR") or "unknown reason") + "). Copy this link and send it yourself: %s" % link, "ok" if sent else "error")
                else:
                    flash("That user has no email address. Add one first.", "error")
            elif action == "getlink":
                u = conn.execute("SELECT * FROM users WHERE id=?", (int(request.form["uid"]),)).fetchone()
                if u and u["must_reset"]:
                    link = url_for("register", token=make_token(u["id"], "invite", 24 * 7), _external=True)
                    flash("Registration link for %s (good for 7 days; a new link replaces any earlier one): %s" % (u["display_name"], link), "ok")
                elif u:
                    link = url_for("reset_with_token", token=make_token(u["id"], "reset", 2), _external=True)
                    flash("Password-reset link for %s (good for 2 hours): %s" % (u["display_name"], link), "ok")
            elif action in ("deactivate", "activate", "delete"):
                uid = int(request.form["uid"])
                target = conn.execute("SELECT * FROM users WHERE id=?", (uid,)).fetchone()
                others = conn.execute("SELECT COUNT(*) FROM users WHERE role='admin' AND active<>0 AND id<>?", (uid,)).fetchone()[0]
                if uid == g.user["id"]:
                    flash("You cannot deactivate or delete your own account.", "error")
                elif action != "activate" and target["role"] == "admin" and not others:
                    flash("There must be at least one active administrator.", "error")
                elif action == "activate":
                    conn.execute("UPDATE users SET active=1 WHERE id=?", (uid,))
                    flash("%s can sign in again." % target["display_name"], "ok")
                elif action == "deactivate":
                    conn.execute("UPDATE users SET active=0 WHERE id=?", (uid,))
                    flash("%s is deactivated and cannot sign in. You can turn them back on any time." % target["display_name"], "ok")
                else:
                    conn.execute("DELETE FROM user_sections WHERE user_id=?", (uid,))
                    conn.execute("DELETE FROM tokens WHERE user_id=?", (uid,))
                    conn.execute("DELETE FROM users WHERE id=?", (uid,))
                    flash("%s was deleted." % target["display_name"], "ok")
            elif action == "setemail":
                em = request.form.get("email", "").strip().lower()
                if EMAIL_RE.match(em):
                    try:
                        conn.execute("UPDATE users SET email=? WHERE id=?", (em, int(request.form["uid"])))
                    except sqlite3.IntegrityError:
                        pass
                else:
                    flash("Enter a valid email address.", "error")
            elif action == "role":
                uid = int(request.form["uid"])
                if request.form.get("role") in ROLES and uid != g.user["id"]:
                    conn.execute("UPDATE users SET role=? WHERE id=?", (request.form["role"], uid))
                conn.execute("DELETE FROM user_sections WHERE user_id=?", (uid,))
                for sid in request.form.getlist("sections"):
                    conn.execute("INSERT OR IGNORE INTO user_sections VALUES(?,?)", (uid, int(sid)))
                save_areas(conn, uid)
            conn.commit()
            return redirect(url_for("users"))
        rows = conn.execute("SELECT * FROM users ORDER BY username").fetchall()
        now = datetime.utcnow()
        pend = {r["user_id"] for r in conn.execute("SELECT user_id FROM tokens WHERE kind='invite'")}
        info = {}
        for r in rows:
            seen = r["last_seen"] or ""
            online = bool(seen) and seen >= (now - timedelta(minutes=10)).isoformat()
            if r["active"] == 0:
                state = "Deactivated"
            elif r["must_reset"]:
                state = "Invited, not registered yet" if r["id"] in pend else "Must change password"
            else:
                state = "Registered"
            info[r["id"]] = {"active": r["active"] != 0, "online": online, "state": state, "verified": not r["must_reset"],
                             "last": eastern(r["last_login"])}
        secs = conn.execute("SELECT * FROM sections WHERE bulletin_id IS NULL ORDER BY position").fetchall()
        assigned = {r["id"]: {x["section_id"] for x in conn.execute(
            "SELECT section_id FROM user_sections WHERE user_id=?", (r["id"],))} for r in rows}
        user_areas = {r["id"]: {x["area"] for x in conn.execute("SELECT area FROM user_areas WHERE user_id=?", (r["id"],))} for r in rows}
        return render_template("users.html", areas=AREAS, user_areas=user_areas, users=rows, roles=ROLES, help=ROLE_HELP,
                               secs=secs, assigned=assigned, info=info, allowed=allowed, can_manage=lambda role: role in allowed)

    @app.route("/admin/backup")
    @need("admin")
    def backup():
        """One zip with the database and every upload; restore it on any other server to move the site."""
        buf = io.BytesIO()
        with tempfile.TemporaryDirectory() as tmp:
            snap = os.path.join(tmp, "bulletin.db")
            dest = sqlite3.connect(snap)
            db().backup(dest)
            dest.close()
            with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
                z.write(snap, "bulletin.db")
                for n in os.listdir(app.config["UPLOAD_DIR"]):
                    p = os.path.join(app.config["UPLOAD_DIR"], n)
                    if os.path.isfile(p):
                        z.write(p, "uploads/" + n)
        return app.response_class(buf.getvalue(), mimetype="application/zip", headers={
            "Content-Disposition": f'attachment; filename="bulletin-backup-{date.today().isoformat()}.zip"'})

    @app.route("/admin/restore", methods=["POST"])
    @need("admin")
    def restore():
        f = request.files.get("backup")
        try:
            z = zipfile.ZipFile(io.BytesIO(f.read())) if f and f.filename else None
            if z is None or "bulletin.db" not in z.namelist():
                raise ValueError
            names = [n for n in z.namelist() if n.startswith("uploads/") and not n.endswith("/")]
            if any(".." in n or n.startswith("/") for n in names):
                raise ValueError
            with tempfile.TemporaryDirectory() as tmp:
                snap = os.path.join(tmp, "bulletin.db")
                with open(snap, "wb") as out:
                    out.write(z.read("bulletin.db"))
                check = sqlite3.connect(snap)
                check.execute("SELECT 1 FROM users LIMIT 1")
                check.execute("SELECT 1 FROM bulletins LIMIT 1")
                db().commit()
                check.backup(db())
                check.close()
            for n in names:
                target = os.path.join(app.config["UPLOAD_DIR"], os.path.basename(n))
                with open(target, "wb") as out:
                    out.write(z.read(n))
        except (ValueError, zipfile.BadZipFile, sqlite3.Error):
            flash("That is not a valid bulletin backup file.", "error")
            return redirect(url_for("users"))
        session.clear()
        flash("Backup restored. Please sign in again.", "ok")
        return redirect(url_for("login"))

    @app.route("/admin/sections", methods=["GET", "POST"])
    @need("admin")
    def sections_admin():
        conn = db()
        if request.method == "POST":
            sid = int(request.form["sid"])
            pages = max(1, min(8, int(request.form.get("pages", 1))))
            conn.execute("UPDATE sections SET title=?, pages=?, weekly=?, position=? WHERE id=?",
                         (request.form.get("title", "").strip() or "Untitled", pages,
                          1 if request.form.get("weekly") else 0, int(request.form.get("position", 0)), sid))
            conn.commit()
            return redirect(url_for("sections_admin"))
        return render_template("sections.html",
                               secs=conn.execute("SELECT * FROM sections WHERE bulletin_id IS NULL ORDER BY position").fetchall())

    return app


if __name__ == "__main__":
    create_app().run(debug=False, port=5050)
